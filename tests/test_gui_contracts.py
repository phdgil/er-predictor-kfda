import json
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from core.contracts import (
    ERBA_BINDING_CLASSIFICATION_EXCEL_CONTRACT_ID,
    ERBA_CLASSIFICATION_DIAGNOSTIC_COLUMNS,
    ERBA_CLASSIFICATION_METADATA_COLUMNS,
    ERBA_CLASSIFICATION_PREDICTION_COLUMNS,
    ERBA_CLASSIFICATION_PUBCHEM_COLUMNS,
    ERBAArtifactSpec,
    ERBARoutePreflight,
    ERBAResult,
    ERBAStatusCode,
    ERBASubtype,
    ERBATask,
)
from gui.erba_tab import (
    BATCH_RUNNING_RESULT,
    PREDICTED_AD_COUNT_NOTE,
    ERBA_BATCH_AD_COLUMNS,
    ERBABatchExportResult,
    ERBAReleasedCatalog,
    ErbaTab,
    allocate_output_path,
    batch_prediction_availability,
    batch_progress_text,
    batch_pubchem_status,
    batch_run_status,
    batch_success_dialog,
    batch_destination_display as erba_batch_destination_display,
    build_primary_predictions,
    canonicalize_batch_input,
    classification_prediction_row,
    export_erba_batch,
    load_erba_catalog,
    load_shared_example_input,
    metadata_row,
    save_erba_batch_graphs,
)
from gui.main_window import (
    ERTA_LEGACY_VALID_COUNT_NOTE,
    MainWindow,
    allocate_erta_output_path,
)


class GuiContractTests(unittest.TestCase):
    def test_main_window_keeps_erta_first_and_nests_legacy_notebook(self):
        source = Path("gui/main_window.py").read_text(encoding="utf-8")
        self.assertIn('self.notebook.add(self.erta_tab, text="ERTA")', source)
        self.assertIn('self.notebook.add(self.eralpha_tab, text="ERalpha")', source)
        self.assertNotIn('text="ERbeta"', source)
        self.assertNotIn("self.erbeta_tab", source)
        self.assertIn("self.notebook.select(self.erta_tab)", source)
        self.assertIn("self.erta_notebook = ttk.Notebook(self.erta_tab)", source)
        self.assertIn('self.title("ER Predictor")', source)
        erba_source = Path("gui/erba_tab.py").read_text(encoding="utf-8")
        self.assertIn("self.single_task_var", erba_source)
        self.assertIn("self.batch_task_var", erba_source)
        self.assertNotIn("self.task_var =", erba_source)
        self.assertIn('text="Single prediction"', erba_source)
        self.assertIn('text="Batch prediction"', erba_source)
        self.assertIn('text="Input"', erba_source)
        self.assertIn('text="Prediction result"', erba_source)
        self.assertIn("Applicability domain: Not evaluated", erba_source)
        self.assertIn("Non-binding probability", erba_source)
        self.assertIn("Binding probability", erba_source)
        self.assertIn('text="Model"', erba_source)
        self.assertIn("command=self.browse_model", erba_source)
        self.assertIn("command=self.load_model_clicked", erba_source)
        self.assertIn("command=self.browse_ad_reference", erba_source)
        self.assertIn("command=self.fit_ad_clicked", erba_source)
        self.assertNotIn('text="Task"', erba_source)
        self.assertNotIn('text="Receptor"', erba_source)
        self.assertNotIn("single_route_reason", erba_source)
        self.assertNotIn("batch_route_reason", erba_source)
        self.assertNotIn("Negative probability", erba_source)
        self.assertNotIn("Positive probability", erba_source)
        self.assertIn('"detail_key"', erba_source)
        self.assertNotIn("values = asdict(result)", erba_source)

    def test_runtime_install_root_is_carried_to_both_batch_workflows(self):
        app_source = Path("app.py").read_text(encoding="utf-8")
        main_source = Path("gui/main_window.py").read_text(encoding="utf-8")
        erba_source = Path("gui/erba_tab.py").read_text(encoding="utf-8")
        self.assertIn("install_root=str(paths.install_root)", app_source)
        self.assertIn("self.install_root = install_root", main_source)
        self.assertIn("install_root=self.install_root", main_source)
        self.assertIn("self.install_root = Path(install_root)", erba_source)
        self.assertIn(
            "forbidden_roots=(self.project_root, self.install_root)",
            main_source,
        )
        self.assertIn(
            "forbidden_roots=(self.project_root, self.install_root)",
            erba_source,
        )

    def test_parity_widget_contract_has_identical_complete_key_sets(self):
        keys = (
            "endpoint_tab",
            "options_button",
            "options_frame",
            "model_entry",
            "model_browse_button",
            "model_reload_button",
            "ad_entry",
            "ad_browse_button",
            "ad_reload_button",
            "mode_notebook",
            "single_tab",
            "batch_tab",
            "cas_label",
            "cas_entry",
            "pubchem_button",
            "smiles_label",
            "smiles_entry",
            "example_button",
            "single_predict_button",
            "draw_structure_button",
            "batch_input_entry",
            "batch_browse_button",
            "batch_destination_entry",
            "batch_example_button",
            "batch_predict_button",
            "batch_progress",
            "batch_status",
            "batch_result",
        )
        for source_path in ("gui/main_window.py", "gui/erba_tab.py"):
            source = Path(source_path).read_text(encoding="utf-8")
            for key in keys:
                self.assertEqual(
                    source.count(f'"{key}":'),
                    1,
                    f"{source_path} must expose exactly one {key!r} parity handle",
                )
            self.assertIn('text="Example input"', source)
            self.assertIn('text="Download template"', source)
            self.assertIn('uniform="single"', source)

    def test_shared_batch_presentation_contract_is_endpoint_neutral(self):
        self.assertEqual(
            BATCH_RUNNING_RESULT,
            "Batch prediction is running.\n\n"
            "Completion details will appear after the workbook is saved.",
        )
        self.assertEqual(
            batch_progress_text("Ready to read input", 0, 0, 0),
            "0% - 0/0 - Reading input workbook",
        )
        self.assertEqual(
            batch_progress_text("Predicting", 3, 10, 61),
            "61% - 3/10 - Preprocessing and prediction",
        )
        self.assertEqual(
            batch_progress_text("Complete", 10, 10, 100),
            "100% - 10/10 - Completed",
        )
        self.assertEqual(
            batch_pubchem_status(5, 25, "107-13-1"),
            "Fetching SMILES from PubChem: 5 / 25 (107-13-1)",
        )
        self.assertEqual(batch_run_status("started"), "Batch prediction started.")
        self.assertEqual(
            batch_run_status("completed", "C:/published/result.xlsx"),
            "Batch prediction completed: C:/published/result.xlsx",
        )
        self.assertEqual(
            batch_run_status("failed", "write failed"),
            "Batch prediction failed: write failed",
        )

    def test_published_all_unavailable_dialog_is_blue_contract_without_false_claim(self):
        title, message = batch_success_dialog(
            destination="C:/published/result.xlsx",
            total_count=2,
            predicted_count=0,
            not_predicted_count=2,
            outcome_counts=(("Binding", 0), ("Non-binding", 0)),
            graph_paths=(),
            graph_directory=None,
            detail_lines=("Graph details: Optional graph files were not generated.",),
        )

        self.assertEqual(title, "Batch prediction done")
        self.assertIn("Total rows: 2", message)
        self.assertIn("Predicted: 0", message)
        self.assertIn("Not predicted: 2", message)
        self.assertIn("No rows could be predicted.", message)
        self.assertIn("Graph files: 0", message)
        self.assertNotIn("All rows were predicted", message)
        self.assertEqual(
            batch_prediction_availability(2, 1, 1),
            "1 row(s) could not be predicted. "
            "Row-level reasons are saved in the workbook.",
        )

    def test_catalog_releases_only_approved_routes_and_keeps_beta_regression_disabled(self):
        payload = {
            "schema_version": 2,
            "regression_parity_approved": True,
            "routes": [
                {"task": "classification", "subtype": "er_alpha", "release_status": "released",
                 "relative_path": "alpha.joblib", "size_bytes": 10, "sha256": "a" * 64, "model_id": "alpha"},
                {"task": "ic50_regression", "subtype": "er_beta", "release_status": "released",
                 "relative_path": "beta.joblib", "size_bytes": 10, "sha256": "b" * 64, "model_id": "beta"},
                {"task": "classification", "subtype": "er_beta", "release_status": "pending",
                 "relative_path": "pending.joblib", "size_bytes": 10, "sha256": "c" * 64, "model_id": "pending"},
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "catalog.v2.json")
            path.write_text(json.dumps(payload), encoding="utf-8")
            catalog, parity, _ = load_erba_catalog(path)
        self.assertTrue(parity)
        self.assertIn((ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA), catalog)
        self.assertEqual(
            [
                spec.model_id
                for spec in catalog.choices_for(
                    ERBATask.CLASSIFICATION,
                    ERBASubtype.ER_ALPHA,
                )
            ],
            ["alpha"],
        )
        self.assertNotIn((ERBATask.IC50_REGRESSION, ERBASubtype.ER_BETA), catalog)
        self.assertNotIn((ERBATask.CLASSIFICATION, ERBASubtype.ER_BETA), catalog)

    def test_catalog_retains_multiple_released_model_choices_for_one_route(self):
        payload = {
            "schema_version": 2,
            "regression_parity_approved": False,
            "routes": [
                {
                    "task": "classification",
                    "subtype": "er_alpha",
                    "release_status": "released",
                    "relative_path": "alpha-v7.joblib",
                    "size_bytes": 10,
                    "sha256": "a" * 64,
                    "model_id": "alpha-v7",
                },
                {
                    "task": "classification",
                    "subtype": "er_alpha",
                    "release_status": "released",
                    "relative_path": "alpha-v8.joblib",
                    "size_bytes": 20,
                    "sha256": "b" * 64,
                    "model_id": "alpha-v8",
                },
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "catalog.v2.json")
            path.write_text(json.dumps(payload), encoding="utf-8")
            catalog, _, _ = load_erba_catalog(path)

        route = (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
        self.assertIsInstance(catalog, ERBAReleasedCatalog)
        self.assertEqual(catalog[route].model_id, "alpha-v7")
        self.assertEqual(
            [spec.model_id for spec in catalog.choices_for(*route)],
            ["alpha-v7", "alpha-v8"],
        )
        self.assertEqual(catalog.model_for(*route, "alpha-v8").size_bytes, 20)

    def test_output_names_are_reserved_without_overwriting(self):
        with tempfile.TemporaryDirectory() as directory:
            first = allocate_output_path(directory, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
            second = allocate_output_path(directory, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
            self.assertTrue(first.exists())
            self.assertTrue(second.exists())
            self.assertNotEqual(first, second)

    def test_both_endpoints_use_and_independently_restore_the_shared_example(self):
        workbook = Path("templates/test.xlsx").resolve()
        with patch(
            "gui.erba_tab.resolve_shared_example_input",
            return_value=workbook,
        ):
            example = load_shared_example_input(Path.cwd())
        source = pd.read_excel(
            example.workbook,
            sheet_name=0,
            nrows=1,
            dtype=str,
            keep_default_na=False,
        )
        cas_header = next(
            header
            for header in ("CAS", "CARSRN", "CASRN")
            if header in source.columns
        )
        self.assertEqual(example.cas, source.loc[0, cas_header])
        expected_smiles = (
            source.loc[0, "SMILES"] if "SMILES" in source.columns else ""
        )
        self.assertEqual(example.smiles, expected_smiles)
        self.assertTrue(example.cas)

        class Variable:
            def __init__(self, value=""):
                self.value = value

            def get(self):
                return self.value

            def set(self, value):
                self.value = value

        erta = MainWindow.__new__(MainWindow)
        eralpha = ErbaTab.__new__(ErbaTab)
        for endpoint in (erta, eralpha):
            endpoint.example_cas = example.cas
            endpoint.example_smiles = example.smiles
            endpoint.example_workbook = example.workbook
            endpoint.cas_var = Variable("changed-cas")
            endpoint.smiles_var = Variable("changed-smiles")
            endpoint.batch_input_var = Variable("changed-input")
            endpoint.batch_input_display_var = Variable("changed.xlsx")
            endpoint.batch_destination_var = Variable("changed-destination")
        erta.set_status = lambda _text: None
        eralpha.single_status_var = Variable()

        erta.load_example_input()
        eralpha.load_example_input()

        self.assertEqual(
            (erta.cas_var.get(), erta.smiles_var.get()),
            (example.cas, example.smiles),
        )
        self.assertEqual(
            (eralpha.cas_var.get(), eralpha.smiles_var.get()),
            (example.cas, example.smiles),
        )
        for endpoint in (erta, eralpha):
            self.assertEqual(
                endpoint.batch_input_var.get(),
                str(example.workbook),
            )
            self.assertEqual(
                endpoint.batch_input_display_var.get(),
                example.workbook.name,
            )
            self.assertEqual(
                endpoint.batch_destination_var.get(),
                str(example.workbook.resolve().parent),
            )
        erta.cas_var.set("ERTA-only")
        self.assertEqual(eralpha.cas_var.get(), example.cas)
        main_source = Path("gui/main_window.py").read_text(encoding="utf-8")
        eralpha_source = Path("gui/erba_tab.py").read_text(encoding="utf-8")
        self.assertIn(
            "self.batch_input_var = tk.StringVar(value=example_path)",
            main_source,
        )
        self.assertIn(
            "self.batch_input_var = tk.StringVar(value=example_path)",
            eralpha_source,
        )

    def test_source_run_derives_the_user_example_sibling_without_a_fixed_path(self):
        with tempfile.TemporaryDirectory() as directory:
            parent = Path(directory)
            project_root = parent / "ER_Predictor_Code"
            project_root.mkdir()
            workbook = parent / "ER_Predictor" / "test.xlsx"
            workbook.parent.mkdir()
            pd.DataFrame({"CARSRN": ["50-00-0"]}).to_excel(
                workbook,
                index=False,
            )

            with patch(
                "gui.erba_tab.resolve_shared_example_input",
                return_value=workbook,
            ) as resolve:
                example = load_shared_example_input(project_root)

        resolve.assert_called_once_with(workbook)
        self.assertEqual(example.workbook, workbook)
        self.assertEqual(example.cas, "50-00-0")
        self.assertEqual(example.smiles, "")

    def test_unavailable_shared_example_is_recoverable_without_fallback_chemistry(self):
        with patch(
            "gui.erba_tab.resolve_shared_example_input",
            side_effect=RuntimeError("Examples/test.xlsx is not writable"),
        ):
            example = load_shared_example_input(Path.cwd())

        self.assertIsNone(example.workbook)
        self.assertEqual((example.cas, example.smiles), ("", ""))
        self.assertIn("Examples/test.xlsx is not writable", example.error)

        class Variable:
            def __init__(self, value=""):
                self.value = value

            def get(self):
                return self.value

            def set(self, value):
                self.value = value

        erta = MainWindow.__new__(MainWindow)
        erta.example_workbook = None
        erta.example_error = example.error
        erta.cas_var = Variable("manual-cas")
        erta.smiles_var = Variable("manual-smiles")
        erta_status = []
        erta.set_status = erta_status.append

        eralpha = ErbaTab.__new__(ErbaTab)
        eralpha.example_workbook = None
        eralpha.example_error = example.error
        eralpha.cas_var = Variable("manual-cas")
        eralpha.smiles_var = Variable("manual-smiles")
        eralpha.single_status_var = Variable()

        erta.load_example_input()
        eralpha.load_example_input()

        for endpoint in (erta, eralpha):
            self.assertEqual(endpoint.cas_var.get(), "manual-cas")
            self.assertEqual(endpoint.smiles_var.get(), "manual-smiles")
        self.assertIn(example.error, erta_status[-1])
        self.assertIn(example.error, eralpha.single_status_var.get())



class _Variable:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class _StartupPredictor:
    def __init__(self, error=None):
        self.error = error
        self.calls = []

    def load_model(self, path):
        self.calls.append(path)
        if self.error:
            raise self.error


class _StartupADCalculator:
    def __init__(self, error=None):
        self.error = error
        self.calls = []

    def fit_from_excel_cached(self, path):
        self.calls.append(path)
        if self.error:
            raise self.error


class _StartupControl:
    def __init__(self):
        self.state = "normal"

    def configure(self, **kwargs):
        self.state = kwargs.get("state", self.state)


class StartupLoadingContractTests(unittest.TestCase):
    def _window(self, model_path, ad_path, predictor=None, ad_calculator=None):
        window = MainWindow.__new__(MainWindow)
        window.model_path_var = _Variable(model_path)
        window.ad_ref_path_var = _Variable(ad_path)
        window.loaded_model_path_var = _Variable()
        window.loaded_ad_ref_path_var = _Variable()
        window.predictor = predictor or _StartupPredictor()
        window.ad_calculator = ad_calculator or _StartupADCalculator()
        window.statuses = []
        window.set_status = window.statuses.append
        window.run_threaded = lambda job: job()
        window.ui = lambda callback, *args: callback(*args)
        window._batch_active = False
        window._erta_reload_in_flight = False
        window.run_batch_button = _StartupControl()
        window.model_entry = _StartupControl()
        window.model_browse_button = _StartupControl()
        window.model_reload_button = _StartupControl()
        window.ad_entry = _StartupControl()
        window.ad_browse_button = _StartupControl()
        window.ad_reload_button = _StartupControl()
        return window

    def test_startup_reports_blank_or_missing_model_path_and_does_not_claim_ready(self):
        with tempfile.TemporaryDirectory() as directory:
            ad_path = str(Path(directory, "reference.xlsx"))
            missing_model_path = str(Path(directory, "missing.keras"))
            Path(ad_path).touch()
            for model_path, failure in (
                ("", "model path is blank"),
                (missing_model_path, f"model path does not exist: {missing_model_path}"),
            ):
                with self.subTest(model_path=model_path):
                    window = self._window(model_path, ad_path)

                    window.load_defaults_on_startup()

                    self.assertEqual(window.predictor.calls, [])
                    self.assertEqual(window.ad_calculator.calls, [ad_path])
                    self.assertEqual(window.statuses[-1], f"Startup failed: {failure}")
                    self.assertNotIn("Ready", window.statuses)

    def test_startup_reports_blank_or_missing_ad_reference_path_and_does_not_claim_ready(self):
        with tempfile.TemporaryDirectory() as directory:
            model_path = str(Path(directory, "model.keras"))
            missing_ad_path = str(Path(directory, "missing.xlsx"))
            Path(model_path).touch()
            for ad_path, failure in (
                ("", "AD reference path is blank"),
                (missing_ad_path, f"AD reference path does not exist: {missing_ad_path}"),
            ):
                with self.subTest(ad_path=ad_path):
                    window = self._window(model_path, ad_path)

                    window.load_defaults_on_startup()

                    self.assertEqual(window.predictor.calls, [model_path])
                    self.assertEqual(window.ad_calculator.calls, [])
                    self.assertEqual(window.statuses[-1], f"Startup failed: {failure}")
                    self.assertNotIn("Ready", window.statuses)

    def test_startup_reports_model_load_exception_and_still_loads_ad(self):
        with tempfile.TemporaryDirectory() as directory:
            model_path = str(Path(directory, "model.keras"))
            ad_path = str(Path(directory, "reference.xlsx"))
            Path(model_path).touch()
            Path(ad_path).touch()
            window = self._window(
                model_path,
                ad_path,
                predictor=_StartupPredictor(RuntimeError("corrupt model")),
            )

            window.load_defaults_on_startup()

        self.assertEqual(window.predictor.calls, [model_path])
        self.assertEqual(window.ad_calculator.calls, [ad_path])
        self.assertEqual(window.statuses[-1], "Startup failed: model load failed: corrupt model")
        self.assertNotIn("Ready", window.statuses)

    def test_startup_reports_ad_fit_exception_and_does_not_claim_ready(self):
        with tempfile.TemporaryDirectory() as directory:
            model_path = str(Path(directory, "model.keras"))
            ad_path = str(Path(directory, "reference.xlsx"))
            Path(model_path).touch()
            Path(ad_path).touch()
            window = self._window(
                model_path,
                ad_path,
                ad_calculator=_StartupADCalculator(RuntimeError("invalid reference")),
            )

            window.load_defaults_on_startup()

        self.assertEqual(window.predictor.calls, [model_path])
        self.assertEqual(window.ad_calculator.calls, [ad_path])
        self.assertEqual(window.statuses[-1], "Startup failed: AD fit failed: invalid reference")
        self.assertNotIn("Ready", window.statuses)

    def test_startup_loads_both_resources_before_setting_ready(self):
        with tempfile.TemporaryDirectory() as directory:
            model_path = str(Path(directory, "model.keras"))
            ad_path = str(Path(directory, "reference.xlsx"))
            Path(model_path).touch()
            Path(ad_path).touch()
            window = self._window(model_path, ad_path)

            window.load_defaults_on_startup()

        self.assertEqual(window.predictor.calls, [model_path])
        self.assertEqual(window.ad_calculator.calls, [ad_path])
        self.assertEqual(
            window.statuses,
            ["Loading default model...", "Loading default AD reference...", "Ready"],
        )


class EralphaModelSelectionContractTests(unittest.TestCase):
    class Variable:
        def __init__(self, value=""):
            self.value = value

        def get(self):
            return self.value

        def set(self, value):
            self.value = value

    def _catalog_tab(self):
        tab = ErbaTab.__new__(ErbaTab)
        tab.project_root = Path.cwd()
        tab.fixed_subtype = ERBASubtype.ER_ALPHA
        tab.catalog_path = (
            tab.project_root / "models" / "erba" / "catalog.v2.json"
        )
        (
            tab.catalog,
            tab.regression_parity_approved,
            tab.catalog_payload,
        ) = load_erba_catalog(tab.catalog_path)
        return tab

    def test_approved_released_model_browse_selects_catalogued_bytes(self):
        tab = self._catalog_tab()
        route = (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
        spec = tab.catalog[route]
        approved = (tab.catalog_path.parent / spec.relative_path).resolve()
        tab.model_path_var = self.Variable()
        tab.single_status_var = self.Variable()
        tab.batch_status_var = self.Variable()

        with patch(
            "gui.erba_tab.filedialog.askopenfilename",
            return_value=str(approved),
        ):
            tab.browse_model()

        self.assertEqual(tab.model_path_var.get(), str(approved))
        self.assertIn("Reload model", tab.single_status_var.get())
        selected = tab._released_model_for_path(approved)
        self.assertEqual(selected[:3], (*route, spec))

    def test_reload_activates_the_selected_released_model_and_binds_its_ad_route(self):
        tab = self._catalog_tab()
        route = (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
        spec = tab.catalog[route]
        approved = (tab.catalog_path.parent / spec.relative_path).resolve()
        selected_text = str(approved.relative_to(Path.cwd())).replace("\\", "/")
        tab.model_path_var = self.Variable(selected_text)
        tab.ad_ref_path_var = self.Variable()
        tab.loaded_model_path_var = self.Variable()
        tab.loaded_ad_ref_path_var = self.Variable()
        tab.single_status_var = self.Variable()
        tab.batch_status_var = self.Variable()
        tab._model_generation = 0
        tab._model_reload_id = 0
        tab._route_changed = lambda *_args: None
        tab.after = lambda _delay, callback: callback()

        class Predictor:
            def preflight(self, task, subtype):
                return ERBARoutePreflight(task, subtype, ERBAStatusCode.OK)

        predictor = Predictor()

        class ImmediateThread:
            def __init__(self, target, daemon):
                self.target = target

            def start(self):
                self.target()

        with patch.object(
            tab,
            "_predictor_for_spec",
            return_value=predictor,
        ), patch("gui.erba_tab.threading.Thread", ImmediateThread), patch(
            "gui.erba_tab.messagebox.showinfo"
        ):
            tab.load_model_clicked()

        self.assertIs(tab.predictor, predictor)
        self.assertEqual(tab.selected_model_spec, spec)
        self.assertEqual(tab.loaded_model_path_var.get(), str(approved))
        self.assertEqual(
            Path(tab.ad_ref_path_var.get()).name,
            "classification_er_alpha_reference.xlsx",
        )
        self.assertEqual(tab._model_generation, 1)
        self.assertEqual(tab._model_reload_id, 1)

    def test_late_model_reload_completion_and_failure_cannot_replace_newer_success(self):
        tab = self._catalog_tab()
        route = (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
        spec = tab.catalog[route]
        approved = (tab.catalog_path.parent / spec.relative_path).resolve()
        selected_text = str(approved.relative_to(Path.cwd())).replace("\\", "/")
        tab.model_path_var = self.Variable(selected_text)
        tab.ad_ref_path_var = self.Variable()
        tab.loaded_model_path_var = self.Variable()
        tab.loaded_ad_ref_path_var = self.Variable()
        tab.single_status_var = self.Variable()
        tab.batch_status_var = self.Variable()
        tab._model_reload_id = 2
        tab._model_generation = 0
        tab._route_changed = lambda *_args: None
        older_predictor, newer_predictor = object(), object()

        with patch("gui.erba_tab.messagebox.showinfo") as show_info, patch(
            "gui.erba_tab.messagebox.showerror"
        ) as show_error:
            tab._finish_model_reload(
                2,
                selected_text,
                *route,
                spec,
                approved,
                newer_predictor,
            )
            status_after_newer = tab.single_status_var.get()
            tab._finish_model_reload(
                1,
                selected_text,
                *route,
                spec,
                approved,
                older_predictor,
            )
            tab._model_reload_failed(1, selected_text, "obsolete failure")

        self.assertIs(tab.predictor, newer_predictor)
        self.assertEqual(tab._model_generation, 1)
        self.assertEqual(tab.single_status_var.get(), status_after_newer)
        show_info.assert_called_once()
        show_error.assert_not_called()

    def test_unapproved_model_browse_is_rejected_before_predictor_construction(self):
        tab = self._catalog_tab()
        original = "approved-selection-remains"
        tab.model_path_var = self.Variable(original)
        tab.single_status_var = self.Variable()
        tab.batch_status_var = self.Variable()
        with tempfile.TemporaryDirectory() as directory:
            unapproved = Path(directory, "unapproved.joblib")
            unapproved.write_bytes(b"not an approved model")
            with patch(
                "gui.erba_tab.filedialog.askopenfilename",
                return_value=str(unapproved),
            ), patch("gui.erba_tab.messagebox.showerror") as show_error, patch(
                "gui.erba_tab.ERBAPredictor"
            ) as predictor_constructor:
                tab.browse_model()

        predictor_constructor.assert_not_called()
        show_error.assert_called_once()
        self.assertEqual(tab.model_path_var.get(), original)
        self.assertIn("Arbitrary joblib files are not allowed", str(show_error.call_args))

    def test_ad_browse_and_reload_use_only_the_approved_eralpha_reference(self):
        tab = self._catalog_tab()
        approved = (
            tab.project_root
            / "models"
            / "erba"
            / "ad"
            / "classification_er_alpha_reference.xlsx"
        ).resolve()
        tab.ad_ref_path_var = self.Variable()
        tab.loaded_ad_ref_path_var = self.Variable()
        tab.single_status_var = self.Variable()
        tab.batch_status_var = self.Variable()
        tab._model_generation = 0
        tab._ad_reload_id = 0
        calls = []

        class Calculator:
            last_cache_status = "cache rebuilt"

            def ensure_fitted_from_excel_cached(self, path, **kwargs):
                calls.append((path, kwargs))

        calculator = Calculator()
        previous_ad = SimpleNamespace()
        fresh_ad = SimpleNamespace(
            calculator_for=lambda task, subtype: (
                calls.append((task, subtype)) or calculator
            )
        )
        tab.erba_ad = previous_ad
        tab._request_ad_calculators = {99: previous_ad}
        tab.after = lambda _delay, callback: callback()

        class ImmediateThread:
            def __init__(self, target, daemon):
                self.target = target

            def start(self):
                self.target()

        with patch(
            "gui.erba_tab.filedialog.askopenfilename",
            return_value=str(approved),
        ):
            tab.browse_ad_reference()
        self.assertEqual(tab.ad_ref_path_var.get(), str(approved))
        tab.ad_ref_path_var.set(
            str(approved.relative_to(Path.cwd())).replace("\\", "/")
        )

        with patch(
            "gui.erba_tab.ERBAApplicabilityDomain",
            return_value=fresh_ad,
        ), patch("gui.erba_tab.threading.Thread", ImmediateThread), patch(
            "gui.erba_tab.messagebox.showinfo"
        ):
            tab.fit_ad_clicked()

        self.assertEqual(
            calls[0],
            (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA),
        )
        self.assertEqual(calls[1][0], str(approved))
        self.assertTrue(calls[1][1]["force_refit"])
        self.assertEqual(tab.loaded_ad_ref_path_var.get(), str(approved))
        self.assertEqual(tab._ad_reload_id, 1)
        self.assertIs(tab.erba_ad, fresh_ad)
        self.assertIsNot(tab.erba_ad, previous_ad)
        self.assertIs(tab._request_ad_calculators[99], previous_ad)

        with tempfile.TemporaryDirectory() as directory:
            unapproved = Path(directory, "other-reference.xlsx")
            unapproved.touch()
            with patch(
                "gui.erba_tab.filedialog.askopenfilename",
                return_value=str(unapproved),
            ), patch("gui.erba_tab.messagebox.showerror") as show_error:
                tab.browse_ad_reference()
        show_error.assert_called_once()
        self.assertEqual(tab.ad_ref_path_var.get(), str(approved))

    def test_late_ad_reload_completion_and_failure_cannot_replace_newer_manager(self):
        tab = self._catalog_tab()
        approved = (
            tab.project_root
            / "models"
            / "erba"
            / "ad"
            / "classification_er_alpha_reference.xlsx"
        ).resolve()
        selected_text = str(approved.relative_to(Path.cwd())).replace("\\", "/")
        tab.ad_ref_path_var = self.Variable(selected_text)
        tab.loaded_ad_ref_path_var = self.Variable()
        tab.single_status_var = self.Variable()
        tab.batch_status_var = self.Variable()
        tab._model_generation = 4
        tab._ad_reload_id = 2
        previous_ad, older_ad, newer_ad = object(), object(), object()
        tab.erba_ad = previous_ad
        calculator = SimpleNamespace(last_cache_status="fresh")

        with patch("gui.erba_tab.messagebox.showinfo") as show_info, patch(
            "gui.erba_tab.messagebox.showerror"
        ) as show_error:
            tab._finish_ad_reload(
                1,
                selected_text,
                approved,
                older_ad,
                calculator,
                4,
            )
            self.assertIs(tab.erba_ad, previous_ad)
            tab._finish_ad_reload(
                2,
                selected_text,
                approved,
                newer_ad,
                calculator,
                4,
            )
            status_after_newer = tab.single_status_var.get()
            tab._ad_reload_failed(
                1,
                selected_text,
                4,
                "obsolete failure",
            )

        self.assertIs(tab.erba_ad, newer_ad)
        self.assertEqual(tab.single_status_var.get(), status_after_newer)
        show_info.assert_called_once()
        show_error.assert_not_called()

    def test_model_change_snapshots_model_predictor_and_ad_without_state_leakage(self):
        task, subtype = ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA
        first = ERBAArtifactSpec("first.joblib", 1, "a" * 64, "first")
        second = ERBAArtifactSpec("second.joblib", 1, "b" * 64, "second")
        catalog = ERBAReleasedCatalog()
        catalog.add(task, subtype, first)
        catalog.add(task, subtype, second)

        tab = ErbaTab.__new__(ErbaTab)
        tab.fixed_subtype = subtype
        tab.catalog = catalog
        tab.catalog_error = ""
        tab.regression_parity_approved = False
        tab._request_id = 0
        tab._started_at = {}
        tab._request_generations = {}
        tab._request_predictors = {}
        tab._request_specs = {}
        tab._request_ad_calculators = {}
        tab._model_generation = 0
        tab._active_single_request_id = None
        tab._active_batch_request_id = None
        tab.single_status_var = self.Variable()
        tab.batch_status_var = self.Variable()
        first_predictor, second_predictor = object(), object()
        first_ad, second_ad = object(), object()
        tab.selected_model_spec = first
        tab.predictor = first_predictor
        tab.erba_ad = first_ad

        first_snapshot = tab._snapshot("single")
        tab.selected_model_spec = second
        tab.predictor = second_predictor
        tab.erba_ad = second_ad
        tab._model_generation += 1
        second_snapshot = tab._snapshot("batch")
        tab._active_batch_request_id = second_snapshot[0]

        self.assertEqual(first_snapshot[1:], (task, subtype, "first"))
        self.assertEqual(second_snapshot[1:], (task, subtype, "second"))
        self.assertIs(tab._request_predictors[first_snapshot[0]], first_predictor)
        self.assertIs(tab._request_predictors[second_snapshot[0]], second_predictor)
        self.assertIs(tab._request_ad_calculators[first_snapshot[0]], first_ad)
        self.assertIs(tab._request_ad_calculators[second_snapshot[0]], second_ad)
        self.assertFalse(tab._snapshot_is_current("single", *first_snapshot))
        self.assertTrue(tab._snapshot_is_current("batch", *second_snapshot))


class _Predictor:
    def __init__(self, status=ERBAStatusCode.OK):
        self.calls = []
        self.status = status

    def preflight(self, task, subtype):
        return ERBARoutePreflight(task, subtype, ERBAStatusCode.OK)

    def predict(self, request):
        self.calls.append(request)
        return ERBAResult(
            row_index=request.row_index, task=request.task, subtype=request.subtype,
            raw_smiles=request.smiles, model_smiles=request.smiles or None,
            status_code=self.status, status_message=self.status.value,
            non_binding_probability=0.2, binding_probability=0.8,
            binding_label="binding", model_id="model-v1", model_sha256="a" * 64,
        )


class _BatchAD:
    def __init__(self, *, fail_for=(), in_domain=True):
        self.fail_for = set(fail_for)
        self.in_domain = in_domain
        self.calls = []
        self.calculator = SimpleNamespace(fitted=True)

    def evaluate(self, task, subtype, smiles):
        self.calls.append((task, subtype, smiles))
        if smiles in self.fail_for:
            raise RuntimeError("route AD unavailable")
        result = SimpleNamespace(
            fitted=True,
            in_domain=self.in_domain,
            distance=1.25,
            threshold=2.5,
            distance_in_domain=True,
            max_similarity=0.75,
            similarity_threshold=0.6,
            similarity_in_domain=True,
            pc1=0.5,
            pc2=-0.25,
            message="In-domain" if self.in_domain else "Out-of-domain",
        )
        fingerprint = np.asarray([[1.0, 0.0, 1.0]], dtype=float)
        return self.calculator, result, fingerprint


class ErbaExcelContractTests(unittest.TestCase):
    spec = ERBAArtifactSpec("model.joblib", 1, "a" * 64, "model-v1")
    provenance = {
        "metadata": {
            "protocol_sha256": "a" * 64, "source_manifest_sha256": "b" * 64,
            "historical_exposure_manifest_sha256": "c" * 64,
            "split_manifest_sha256": "d" * 64, "nested_cv_sha256": "e" * 64,
            "internal_resplit_sha256": "f" * 64, "report_sha256": "0" * 64,
            "caveat_sha256": "1" * 64,
        }
    }

    def _input(self, directory, frame):
        path = Path(directory, "input.xlsx")
        frame.to_excel(path, index=False)
        return path

    def test_classification_columns_use_numeric_binding_class_and_parse_validity(self):
        non_binding = classification_prediction_row(
            ERBAResult(
                row_index=0,
                task=ERBATask.CLASSIFICATION,
                subtype=ERBASubtype.ER_ALPHA,
                raw_smiles="CCO",
                model_smiles="CCO",
                status_code=ERBAStatusCode.OK,
                status_message="",
                non_binding_probability=0.7,
                binding_probability=0.3,
                binding_label="non_binding",
            ),
            "64-17-5",
        )
        self.assertEqual(non_binding["Prediction"], 0)
        self.assertEqual(non_binding["Prediction_label"], "Non-binding")
        self.assertIs(non_binding["Mol_valid"], True)

        unsupported_but_parseable = classification_prediction_row(
            ERBAResult(
                row_index=1,
                task=ERBATask.CLASSIFICATION,
                subtype=ERBASubtype.ER_ALPHA,
                raw_smiles="[Na+]",
                model_smiles=None,
                status_code=ERBAStatusCode.NO_CARBON,
                status_message="no carbon",
                non_binding_probability=0.8,
                binding_probability=0.2,
                binding_label="non_binding",
            ),
            "",
        )
        self.assertIs(unsupported_but_parseable["Mol_valid"], True)
        self.assertEqual(unsupported_but_parseable["Probability_Negative_0"], "")
        self.assertEqual(unsupported_but_parseable["Prediction"], "")
        self.assertEqual(unsupported_but_parseable["Prediction_label"], "")

    def test_canonical_aliases_reject_duplicates_and_preserve_source_fields(self):
        frame = pd.DataFrame([["one", "50-00-0", "CCO", "keep"]], columns=["id", "cas", "smiles", "Note"])
        canonical, passthrough = canonicalize_batch_input(frame)
        self.assertEqual(canonical.iloc[0].tolist(), ["one", "50-00-0", "CCO"])
        self.assertEqual(passthrough.columns.tolist(), ["id", "cas", "smiles", "Note"])
        with self.assertRaisesRegex(ValueError, "Duplicate canonical CAS"):
            canonicalize_batch_input(pd.DataFrame([["a", "b"]], columns=["CAS", "cas"]))
        with tempfile.TemporaryDirectory() as directory:
            input_path = self._input(directory, pd.DataFrame([["a", "b"]], columns=["CAS", "cas"]))
            with self.assertRaisesRegex(ValueError, "Duplicate canonical CAS"):
                export_erba_batch(input_path, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA,
                                  _Predictor(), _BatchAD(), self.spec, self.provenance)

    def test_mixed_batch_contract_direct_precedence_failures_and_metadata(self):
        frame = pd.DataFrame(
            [
                ["direct", "12-34-5", "CCO", "first", "spoof-1", "spoof-ad-1"],
                ["lookup", "50-00-0", "", "second", "spoof-2", "spoof-ad-2"],
                ["bad", "12-34-5", "", "third", "spoof-3", "spoof-ad-3"],
            ],
            columns=[
                "identifier",
                "cas no.",
                "SMILES",
                "Analyst note",
                " Binding_Probability ",
                "AD",
            ],
        )
        predictor = _Predictor()
        erba_ad = _BatchAD()
        with tempfile.TemporaryDirectory() as directory, patch(
            "gui.erba_tab.cas_to_smiles",
            return_value={
                "CanonicalSMILES": "CCC",
                "PubChem_CID": 123,
                "PubChem_status": "Found",
            },
        ) as lookup, patch(
            "gui.erba_tab.save_erba_batch_graphs",
            return_value=((), None),
        ):
            export_result = export_erba_batch(
                self._input(directory, frame), ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA, predictor, erba_ad, self.spec, self.provenance,
            )
            with pd.ExcelFile(export_result.destination) as workbook:
                sheet_names = workbook.sheet_names
            published_workbook = load_workbook(export_result.destination, data_only=True)
            active_sheet = published_workbook.active.title
            prediction_sheet = published_workbook["Predictions"]
            headers = [cell.value for cell in prediction_sheet[1]]
            probability_cell = prediction_sheet.cell(
                2,
                headers.index("Probability_Positive_1") + 1,
            ).value
            prediction_cell = prediction_sheet.cell(
                2,
                headers.index("Prediction") + 1,
            ).value
            mol_valid_cell = prediction_sheet.cell(
                2,
                headers.index("Mol_valid") + 1,
            ).value
            distance_cell = prediction_sheet.cell(
                2,
                headers.index("AD_MeanDistance") + 1,
            ).value
            published_workbook.close()
            predictions = pd.read_excel(
                export_result.destination,
                sheet_name="Predictions",
                dtype=str,
                keep_default_na=False,
            )
            supplied = pd.read_excel(
                export_result.destination,
                sheet_name="Input",
                dtype=str,
                keep_default_na=False,
            )
            diagnostics = pd.read_excel(
                export_result.destination,
                sheet_name="Diagnostics",
                dtype=str,
                keep_default_na=False,
            )
            metadata = pd.read_excel(
                export_result.destination,
                sheet_name="Metadata",
                dtype=str,
                keep_default_na=False,
            )
        self.assertEqual(export_result.count, 3)
        self.assertEqual(
            sheet_names,
            ["Predictions", "Guide", "Diagnostics", "Input", "Metadata"],
        )
        self.assertEqual(active_sheet, "Predictions")
        self.assertIsInstance(probability_cell, (int, float))
        self.assertEqual(prediction_cell, 1)
        self.assertIs(type(prediction_cell), int)
        self.assertIs(mol_valid_cell, True)
        self.assertIsInstance(distance_cell, (int, float))
        self.assertEqual([call.smiles for call in predictor.calls], ["CCO", "CCC", ""])
        lookup.assert_called_once_with("50-00-0")
        self.assertEqual(supplied.columns.tolist(), frame.columns.tolist())
        self.assertEqual(supplied["Analyst note"].tolist(), ["first", "second", "third"])
        self.assertEqual(
            supplied[" Binding_Probability "].tolist(),
            ["spoof-1", "spoof-2", "spoof-3"],
        )
        self.assertEqual(
            predictions.columns.tolist(),
            [
                "identifier",
                "Analyst note",
                *ERBA_CLASSIFICATION_PREDICTION_COLUMNS,
                *ERBA_BATCH_AD_COLUMNS,
                *ERBA_CLASSIFICATION_PUBCHEM_COLUMNS,
            ],
        )
        self.assertEqual(
            predictions.loc[0, list(ERBA_CLASSIFICATION_PREDICTION_COLUMNS)].to_dict(),
            {
                "CAS": "12-34-5",
                "SMILES": "CCO",
                "Canonical_SMILES": "CCO",
                "Mol_valid": "True",
                "Probability_Negative_0": "0.2",
                "Probability_Positive_1": "0.8",
                "Prediction": "1",
                "Prediction_label": "Binding",
            },
        )
        self.assertEqual(
            predictions.loc[0, list(ERBA_BATCH_AD_COLUMNS)].to_dict(),
            {
                "AD": "In-domain",
                "AD_MeanDistance": "1.25",
                "AD_DistanceThreshold": "2.5",
                "AD_Distance_InDomain": "True",
                "AD_SimilarityMax": "0.75",
                "AD_SimilarityThreshold": "0.6",
                "AD_Similarity_InDomain": "True",
                "AD_PC1": "0.5",
                "AD_PC2": "-0.25",
            },
        )
        self.assertEqual(predictions.loc[1, "PubChem_CID"], "123")
        self.assertEqual(predictions.loc[1, "PubChem_status"], "Found")
        self.assertEqual(predictions.loc[0, "PubChem_CID"], "")
        self.assertEqual(predictions.loc[0, "PubChem_status"], "")
        self.assertEqual(predictions.loc[2, "Mol_valid"], "False")
        self.assertEqual(predictions.loc[2, "Probability_Positive_1"], "")
        self.assertEqual(predictions.loc[2, "Prediction"], "")
        self.assertEqual(predictions.loc[2, "Prediction_label"], "")
        self.assertTrue(all(predictions.loc[2, column] == "" for column in ERBA_BATCH_AD_COLUMNS))
        self.assertEqual(
            diagnostics.columns.tolist(),
            list(ERBA_CLASSIFICATION_DIAGNOSTIC_COLUMNS),
        )
        self.assertEqual(diagnostics.loc[0, "SMILES_Provenance"], "direct_input")
        self.assertEqual(diagnostics.loc[1, "SMILES_Provenance"], "pubchem_lookup")
        self.assertEqual(diagnostics.loc[2, "SMILES_Provenance"], "unavailable")
        self.assertEqual(diagnostics["row_index"].tolist(), ["0", "1", "2"])
        self.assertEqual(
            diagnostics["Row_ID"].tolist(),
            predictions["identifier"].tolist(),
        )
        self.assertEqual(diagnostics.loc[2, "Status_Code"], "prediction_failed")
        self.assertEqual(diagnostics.loc[2, "Result_Status"], "Not predicted")
        self.assertEqual(diagnostics.loc[2, "Reason_Category"], "Invalid CAS")
        self.assertTrue(diagnostics.loc[2, "Reason_Description"])
        self.assertTrue(diagnostics.loc[2, "Recommended_Action"])
        self.assertEqual(
            erba_ad.calls,
            [
                (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA, "CCO"),
                (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA, "CCC"),
            ],
        )
        self.assertEqual(export_result.binding_count, 2)
        self.assertEqual(export_result.not_predicted_count, 1)
        self.assertEqual(export_result.ad_in_domain_count, 2)
        self.assertEqual(metadata.loc[0, "Performance_Evidence_Scope"], "internal_historically_exposed")
        self.assertEqual(metadata.loc[0, "Protocol_SHA256"], "a" * 64)
        self.assertEqual(metadata.loc[0, "Model_ID"], "model-v1")
        self.assertEqual(metadata.loc[0, "Model_SHA256"], "a" * 64)
        self.assertEqual(metadata.loc[0, "Source_Manifest_SHA256"], "b" * 64)
        self.assertEqual(
            metadata.loc[0, "Historical_Exposure_Manifest_SHA256"],
            "c" * 64,
        )
        self.assertEqual(metadata.loc[0, "Split_Manifest_SHA256"], "d" * 64)
        self.assertEqual(metadata.loc[0, "Nested_CV_SHA256"], "e" * 64)
        self.assertEqual(metadata.loc[0, "Internal_Resplit_SHA256"], "f" * 64)
        self.assertEqual(metadata.loc[0, "Report_SHA256"], "0" * 64)
        self.assertEqual(metadata.loc[0, "Caveat_SHA256"], "1" * 64)
        self.assertIn("historical", metadata.loc[0, "Evidence_Caveat"].lower())
        self.assertEqual(
            metadata.columns.tolist(),
            list(ERBA_CLASSIFICATION_METADATA_COLUMNS),
        )
        self.assertEqual(metadata.loc[0, "Workflow"], "erba")
        self.assertEqual(metadata.loc[0, "Task"], "classification")
        self.assertEqual(metadata.loc[0, "Subtype"], "er_alpha")
        self.assertEqual(
            metadata.loc[0, "Decision_rule"],
            "binding_probability>=0.5",
        )
        self.assertEqual(
            metadata.loc[0, "Preprocessing_Policy_ID"],
            "erba_binding_classification_parent_v2",
        )
        self.assertEqual(
            metadata.loc[0, "Excel_Contract_ID"],
            ERBA_BINDING_CLASSIFICATION_EXCEL_CONTRACT_ID,
        )
        self.assertEqual(
            ERBA_BINDING_CLASSIFICATION_EXCEL_CONTRACT_ID,
            "erba.binding.classification.excel.v3",
        )

    def test_every_export_sheet_uses_the_same_plain_style_as_erta_pandas_output(self):
        def color_signature(color):
            if color is None:
                return None
            return (
                color.type,
                color.rgb if color.type == "rgb" else None,
                color.indexed if color.type == "indexed" else None,
                color.theme if color.type == "theme" else None,
                color.tint,
            )

        def side_signature(side):
            return side.style, color_signature(side.color)

        def style_signature(cell):
            return (
                (
                    cell.font.name,
                    cell.font.sz,
                    cell.font.bold,
                    cell.font.italic,
                    cell.font.underline,
                    color_signature(cell.font.color),
                ),
                (
                    cell.fill.fill_type,
                    color_signature(cell.fill.fgColor),
                    color_signature(cell.fill.bgColor),
                ),
                (
                    cell.alignment.horizontal,
                    cell.alignment.vertical,
                    cell.alignment.wrap_text,
                    cell.alignment.shrink_to_fit,
                    cell.alignment.text_rotation,
                ),
                (
                    side_signature(cell.border.left),
                    side_signature(cell.border.right),
                    side_signature(cell.border.top),
                    side_signature(cell.border.bottom),
                ),
                cell.number_format,
            )

        frame = pd.DataFrame([["row", "CCO"]], columns=["Row_ID", "SMILES"])
        with tempfile.TemporaryDirectory() as directory, patch(
            "gui.erba_tab.save_erba_batch_graphs",
            return_value=((), None),
        ):
            export_result = export_erba_batch(
                self._input(directory, frame),
                ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA,
                _Predictor(ERBAStatusCode.INVALID_SMILES),
                _BatchAD(),
                self.spec,
                self.provenance,
            )
            erta_output = Path(directory, "ERTA_pandas_prediction.xlsx")
            pd.DataFrame({"Result": ["value"]}).to_excel(
                erta_output,
                index=False,
            )
            actual = load_workbook(export_result.destination)
            expected = load_workbook(erta_output)
            try:
                expected_header = style_signature(expected.active["A1"])
                expected_data = style_signature(expected.active["A2"])
                for sheet in actual.worksheets:
                    self.assertIsNone(sheet.freeze_panes)
                    self.assertIsNone(sheet.auto_filter.ref)
                    self.assertEqual(list(sheet.column_dimensions), [])
                    self.assertTrue(
                        all(
                            style_signature(cell) == expected_header
                            for cell in sheet[1]
                        )
                    )
                    self.assertTrue(
                        all(
                            style_signature(cell) == expected_data
                            for row in sheet.iter_rows(min_row=2)
                            for cell in row
                        )
                    )
            finally:
                actual.close()
                expected.close()

    def test_batch_progress_and_guide_explain_all_unavailable_predictions(self):
        class _FdaPredictor(_Predictor):
            def predict(self, request):
                if request.smiles == "O":
                    self.status = ERBAStatusCode.NO_CARBON
                elif request.smiles == "[Na+]":
                    self.status = ERBAStatusCode.METAL_RETAINED
                else:
                    self.status = ERBAStatusCode.OK
                return super().predict(request)

        # Fixture-level equivalent of the FDA category totals; no supplied workbook is altered.
        frame = pd.DataFrame(
            ([["predicted", "CCO"]] * 448)
            + [["inorganic", "O"]] * 18
            + [["metal", "[Na+]"]]
            + [["bad-cas", ""]] * 2
            + [["pubchem", ""]] * 35,
            columns=["Row_ID", "SMILES"],
        )
        frame.loc[467:468, "CAS"] = "12-34-5"
        frame.loc[469:, "CAS"] = "50-00-0"
        progress = []
        statuses = []
        with tempfile.TemporaryDirectory() as directory, patch(
            "gui.erba_tab.cas_to_smiles", side_effect=ValueError("404")
        ), patch("gui.erba_tab.save_erba_batch_graphs", return_value=((), None)):
            export_result = export_erba_batch(
                self._input(directory, frame), ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA, _FdaPredictor(), _BatchAD(),
                self.spec, self.provenance,
                progress_callback=lambda stage, current, total, percent: progress.append(
                    (stage, current, total, percent)
                ),
                status_callback=statuses.append,
            )
            predictions = pd.read_excel(
                export_result.destination,
                sheet_name="Predictions",
                dtype=str,
                keep_default_na=False,
            )
            diagnostics = pd.read_excel(
                export_result.destination,
                sheet_name="Diagnostics",
                dtype=str,
                keep_default_na=False,
            )
            guide = pd.read_excel(
                export_result.destination,
                sheet_name="Guide",
                dtype=str,
                keep_default_na=False,
            )
        self.assertEqual(export_result.count, 504)
        self.assertNotIn("Result_Status", predictions)
        self.assertEqual((diagnostics["Result_Status"] == "Predicted").sum(), 448)
        self.assertEqual((diagnostics["Result_Status"] == "Not predicted").sum(), 56)
        self.assertEqual(
            diagnostics.loc[diagnostics["Result_Status"] == "Not predicted", "Reason_Category"]
            .value_counts()
            .to_dict(),
            {
                "PubChem lookup unavailable": 35,
                "No carbon / inorganic": 18,
                "Invalid CAS": 2,
                "Unsupported metal-containing structure": 1,
            },
        )
        self.assertEqual(predictions.loc[448, "Mol_valid"], "True")
        self.assertEqual(predictions.loc[466, "Mol_valid"], "True")
        self.assertEqual(predictions.loc[467, "Mol_valid"], "False")
        self.assertEqual(predictions.loc[469, "Mol_valid"], "False")
        self.assertEqual(diagnostics.loc[467, "SMILES_Provenance"], "unavailable")
        self.assertEqual(
            diagnostics.loc[469, "SMILES_Provenance"],
            "pubchem_lookup_failed",
        )
        guide_details = guide.set_index("Topic")["Details"].to_dict()
        self.assertIn("not ERTA transactivation", guide_details["Endpoint semantics"])
        self.assertIn("0 = Non-binding", guide_details["Prediction"])
        self.assertIn(
            "binding_probability>=0.5",
            guide_details["Decision_rule"],
        )
        self.assertIn(
            "pubchem_lookup_failed",
            guide_details["SMILES_Provenance"],
        )
        self.assertIn("one-to-one", guide_details["Diagnostics"])
        self.assertEqual(progress[0], ("Reading input workbook", 0, 0, 0))
        self.assertEqual(progress[-1], ("Completed", 504, 504, 100))
        self.assertEqual({entry[0] for entry in progress}, {
            "Reading input workbook", "Resolving CAS/SMILES",
            "Preprocessing and prediction", "Writing workbook", "Completed",
        })
        preprocessing_currents = [
            current
            for stage, current, _total, _percent in progress
            if stage == "Preprocessing and prediction"
        ]
        self.assertEqual(
            preprocessing_currents,
            sorted(preprocessing_currents),
        )
        self.assertEqual(
            statuses[0],
            "Fetching SMILES from PubChem: 470 / 504 (50-00-0)",
        )
        self.assertEqual(
            statuses[-2],
            "Fetching SMILES from PubChem: 504 / 504 (50-00-0)",
        )
        self.assertEqual(statuses[-1], "Batch prediction started.")
        self.assertIn("Not predicted rows", guide["Topic"].tolist())
    def test_single_cas_lookup_uses_isomeric_fallback_and_clears_on_missing_smiles(self):
        class _Variable:
            def __init__(self, value=""):
                self.value = value

            def get(self):
                return self.value

            def set(self, value):
                self.value = value

        class _Widget:
            def configure(self, **kwargs):
                pass

        class _Text(_Widget):
            def __init__(self):
                self.value = ""

            def delete(self, start, end):
                self.value = ""

            def insert(self, start, value):
                self.value = value

        class _ImmediateThread:
            def __init__(self, target, daemon):
                self.target = target

            def start(self):
                self.target()

        tab = ErbaTab.__new__(ErbaTab)
        predictor = _Predictor()
        tab.predictor = predictor
        tab.smiles_var, tab.cas_var = _Variable(), _Variable("50-00-0")
        tab.single_predict_button, tab.single_status_var = _Widget(), _Variable()
        tab.single_prediction_summary_var = _Variable()
        tab.single_negative_probability_var = _Variable()
        tab.single_positive_probability_var = _Variable()
        tab.single_pic50_var, tab.single_ic50_var = _Variable(), _Variable()
        tab.single_detail_var = _Variable()
        tab.single_negative_bar, tab.single_positive_bar = _Widget(), _Widget()
        tab._draw_single_structure = lambda smiles: None
        tab.after = lambda delay, callback: callback()
        tab._finish_duration_ms = lambda request_id: 0
        tab._snapshot_is_current = lambda *args: True
        tab._route_changed = lambda workflow: None
        tab._emit = lambda *args, **kwargs: None
        tab._snapshot = lambda workflow: (1, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA, "model-v1")
        with patch("gui.erba_tab.threading.Thread", _ImmediateThread), patch(
            "gui.erba_tab.cas_to_smiles", return_value={"IsomericSMILES": "C[C@H](O)F"}
        ):
            tab.single_predict_clicked()
        self.assertEqual([request.smiles for request in predictor.calls], ["C[C@H](O)F"])
        self.assertIn("Binding prediction", tab.single_prediction_summary_var.value)
        self.assertIn("Non-binding probability", tab.single_negative_probability_var.value)
        self.assertIn("Binding probability", tab.single_positive_probability_var.value)
        tab._snapshot = lambda workflow: (2, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA, "model-v1")
        with patch("gui.erba_tab.threading.Thread", _ImmediateThread), patch(
            "gui.erba_tab.cas_to_smiles", return_value={"PubChem_CID": 712}
        ):
            tab.single_predict_clicked()
        self.assertEqual(tab.single_prediction_summary_var.value, "Prediction failed")
        self.assertIn("neither CanonicalSMILES nor IsomericSMILES", tab.single_status_var.value)

    def test_structured_single_result_renders_classification_details_for_both_receptors(self):
        class _Variable:
            def __init__(self):
                self.value = ""

            def set(self, value):
                self.value = value

        class _Widget:
            def __init__(self):
                self.values = {}

            def configure(self, **kwargs):
                self.values.update(kwargs)

        class _TaggedText:
            def __init__(self):
                self.inserts = []

            def delete(self, _start, _end):
                self.inserts.clear()

            def insert(self, _start, value, tag=None):
                self.inserts.append((value, tag))

        tab = ErbaTab.__new__(ErbaTab)
        tab.single_prediction_summary_var = _Variable()
        tab.single_negative_probability_var = _Variable()
        tab.single_positive_probability_var = _Variable()
        tab.single_pic50_var, tab.single_ic50_var = _Variable(), _Variable()
        tab.single_detail_var = _Variable()
        tab.single_result_text = _TaggedText()
        tab.single_negative_bar, tab.single_positive_bar = _Widget(), _Widget()
        tab._draw_single_structure = lambda smiles: None
        tab._configure_single_result_task = lambda task: None
        tab._render_single_result(ERBAResult(
            row_index=0, task=ERBATask.CLASSIFICATION, subtype=ERBASubtype.ER_ALPHA,
            raw_smiles="CCO", model_smiles="CCO", status_code=ERBAStatusCode.OK,
            status_message="ok", non_binding_probability=0.2, binding_probability=0.8,
            binding_label="binding", model_id="model-v1",
        ))
        self.assertEqual(tab.single_prediction_summary_var.value, "Binding prediction: Binding")
        self.assertEqual(tab.single_negative_probability_var.value, "Non-binding probability: 20.0%")
        self.assertEqual(tab.single_positive_probability_var.value, "Binding probability: 80.0%")
        self.assertEqual(tab.single_positive_bar.values["value"], 80.0)
        self.assertIn(("Binding prediction:", "detail_key"), tab.single_result_text.inserts)
        self.assertTrue(any(tag == "detail_value" for _, tag in tab.single_result_text.inserts))
        tab._render_single_result(ERBAResult(
            row_index=0, task=ERBATask.CLASSIFICATION, subtype=ERBASubtype.ER_BETA,
            raw_smiles="CCO", model_smiles="CCO", status_code=ERBAStatusCode.OK,
            status_message="ok", non_binding_probability=0.7, binding_probability=0.3,
            binding_label="non_binding", model_id="model-v2",
        ))
        self.assertEqual(tab.single_prediction_summary_var.value, "Binding prediction: Non-binding")
        self.assertEqual(tab.single_negative_probability_var.value, "Non-binding probability: 70.0%")
        self.assertEqual(tab.single_positive_probability_var.value, "Binding probability: 30.0%")

    def test_success_then_invalid_prediction_redraws_probability_canvas_immediately(self):
        class _Variable:
            def __init__(self):
                self.value = ""

            def set(self, value):
                self.value = value

        class _Canvas:
            def __init__(self):
                self.text = []

            def delete(self, _target):
                self.text.clear()

            def winfo_width(self):
                return 500

            def create_text(self, *_args, **kwargs):
                self.text.append(str(kwargs.get("text", "")))

            def create_rectangle(self, *_args, **_kwargs):
                return None

        tab = ErbaTab.__new__(ErbaTab)
        tab.single_prediction_summary_var = _Variable()
        tab.single_negative_probability_var = _Variable()
        tab.single_positive_probability_var = _Variable()
        tab.single_pic50_var = _Variable()
        tab.single_ic50_var = _Variable()
        tab.single_detail_var = _Variable()
        tab.single_negative_bar = None
        tab.single_positive_bar = None
        tab.single_prob_canvas = _Canvas()
        tab._configure_single_result_task = lambda _task: None
        tab._draw_single_structure = lambda _smiles: None

        tab._render_single_result(ERBAResult(
            row_index=0,
            task=ERBATask.CLASSIFICATION,
            subtype=ERBASubtype.ER_ALPHA,
            raw_smiles="CCO",
            model_smiles="CCO",
            status_code=ERBAStatusCode.OK,
            status_message="ok",
            non_binding_probability=0.2,
            binding_probability=0.8,
            binding_label="binding",
        ))
        self.assertIn("0.800", tab.single_prob_canvas.text)

        tab._render_single_result(ERBAResult(
            row_index=0,
            task=ERBATask.CLASSIFICATION,
            subtype=ERBASubtype.ER_ALPHA,
            raw_smiles="invalid",
            model_smiles=None,
            status_code=ERBAStatusCode.INVALID_SMILES,
            status_message="Invalid SMILES",
            non_binding_probability=0.1,
            binding_probability=0.9,
            binding_label="binding",
        ))

        self.assertNotIn("0.800", tab.single_prob_canvas.text)
        self.assertNotIn("0.900", tab.single_prob_canvas.text)
        self.assertEqual(tab.single_prob_canvas.text.count("0.000"), 2)
        self.assertIn("Run prediction to show probabilities", tab.single_prob_canvas.text)
        self.assertTrue(
            any(
                "historically exposed" in line
                for line in tab._single_detail_lines
            )
        )

    def test_single_result_reset_clears_reference_image_and_preserves_caveat(self):
        class Variable:
            def __init__(self, value=""):
                self.value = value

            def set(self, value):
                self.value = value

        class Widget:
            def __init__(self):
                self.options = {}

            def configure(self, **kwargs):
                self.options.update(kwargs)

        tab = ErbaTab.__new__(ErbaTab)
        for name in (
            "single_prediction_summary_var",
            "single_negative_probability_var",
            "single_positive_probability_var",
            "single_pic50_var",
            "single_ic50_var",
            "single_detail_var",
            "single_ad_domain_var",
            "single_nearest_reference_var",
        ):
            setattr(tab, name, Variable())
        tab.single_negative_bar = None
        tab.single_positive_bar = None
        tab.draw_probability_graph = lambda: None
        tab._render_detail_lines = lambda: None
        tab.single_structure_label = Widget()
        tab.single_nearest_reference_structure_label = Widget()
        tab.single_ad_graph_label = Widget()
        tab._nearest_reference_structure_image = object()
        tab._single_ad_graph_image = object()

        tab._clear_single_result("Prediction: -")

        self.assertEqual(
            tab.single_nearest_reference_structure_label.options,
            {"text": "No reference", "image": ""},
        )
        self.assertIsNone(tab._nearest_reference_structure_image)
        self.assertIsNone(tab._single_ad_graph_image)
        self.assertTrue(
            any("historically exposed" in line for line in tab._single_detail_lines)
        )

    def test_reference_preview_matches_erta_caption_compound_and_wrapping(self):
        class Widget:
            def __init__(self):
                self.options = {}

            def configure(self, **kwargs):
                self.options.update(kwargs)

        with tempfile.TemporaryDirectory() as directory:
            tab = ErbaTab.__new__(ErbaTab)
            tab.fixed_subtype = ERBASubtype.ER_ALPHA
            tab.output_root = Path(directory)
            tab.nearest_reference_preview_size = (250, 180)
            tab.single_nearest_reference_structure_label = Widget()
            tab._make_preview_photo = lambda _path, _size: "photo"
            with patch("gui.erba_tab.PIL_AVAILABLE", True), patch(
                "gui.erba_tab.save_molecule_image"
            ):
                tab._draw_nearest_reference_structure(
                    {"SMILES": "CCO", "label": "binding", "CID": 702}
                )

        self.assertEqual(
            tab.single_nearest_reference_structure_label.options,
            {
                "image": "photo",
                "text": "label: Binding / CID: 702",
                "compound": "top",
                "wraplength": 180,
            },
        )

    def test_deferred_single_failure_callbacks_capture_error_text(self):
        class _Variable:
            def __init__(self, value=""):
                self.value = value

            def get(self):
                return self.value

            def set(self, value):
                self.value = value

        class _Control:
            def __init__(self):
                self.state = "normal"

            def configure(self, **kwargs):
                self.state = kwargs.get("state", self.state)

        class _ImmediateThread:
            def __init__(self, target, daemon, args=()):
                self.target = target
                self.args = args

            def start(self):
                self.target(*self.args)

        scheduled = []
        tab = ErbaTab.__new__(ErbaTab)
        tab.smiles_var = _Variable()
        tab.cas_var = _Variable("50-00-0")
        tab.single_predict_button = _Control()
        tab.single_status_var = _Variable()
        tab.single_prediction_summary_var = _Variable()
        tab.single_negative_probability_var = _Variable()
        tab.single_positive_probability_var = _Variable()
        tab.single_pic50_var = _Variable()
        tab.single_ic50_var = _Variable()
        tab.single_detail_var = _Variable()
        tab.single_negative_bar = None
        tab.single_positive_bar = None
        tab.predictor = SimpleNamespace(
            predict=lambda _request: self.fail(
                "predictor must not run after CAS lookup failure"
            )
        )
        tab._snapshot = lambda _workflow: (
            1,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
        )
        tab._finish_duration_ms = lambda _request_id: 0
        tab._snapshot_is_current = lambda *_args: True
        tab._route_changed = lambda _workflow: tab.single_predict_button.configure(state="normal")
        tab._emit = lambda *_args, **_kwargs: None
        tab.after = lambda _delay, callback: scheduled.append(callback)

        with patch("gui.erba_tab.threading.Thread", _ImmediateThread), patch(
            "gui.erba_tab.cas_lookup_smiles",
            side_effect=RuntimeError("lookup offline"),
        ):
            tab.single_predict_clicked()

        self.assertEqual(tab.single_predict_button.state, "disabled")
        self.assertEqual(len(scheduled), 1)
        scheduled.pop()()
        self.assertEqual(tab.single_predict_button.state, "normal")
        self.assertIn("lookup offline", tab.single_status_var.value)

        tab._active_ad_request_id = 2
        tab.erba_ad = SimpleNamespace(
            evaluate=lambda *_args: (_ for _ in ()).throw(
                RuntimeError("AD reference unavailable")
            )
        )
        tab._selected_route = lambda _workflow: (
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
        )
        tab.single_ad_domain_var = _Variable("Applicability domain: Evaluating")
        tab._single_detail_lines = []
        tab._render_detail_lines = lambda: None
        tab._evaluate_single_ad(
            2,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "CCO",
        )

        self.assertEqual(len(scheduled), 1)
        scheduled.pop()()
        self.assertIsNone(tab._active_ad_request_id)
        self.assertIn("AD reference unavailable", tab.single_ad_domain_var.value)

    def test_regression_metadata_is_task_specific_and_requires_parity_provenance(self):
        payload = {
            "metadata": {
                "source_manifest_sha256": "a" * 64,
                "preprocessing_parity_sha256": "b" * 64,
                "report_sha256": "c" * 64,
            }
        }
        row = metadata_row(
            ERBATask.IC50_REGRESSION,
            ERBASubtype.ER_ALPHA,
            self.spec,
            payload,
        )
        self.assertEqual(row["Performance_Evidence_Scope"], "internal_regression_model")
        self.assertEqual(row["Preprocessing_Parity_SHA256"], "b" * 64)
        self.assertEqual(row["Evidence_Caveat"], "")
        self.assertEqual(row["Protocol_SHA256"], "")
    def test_empty_and_fatal_batches_publish_nothing(self):
        with tempfile.TemporaryDirectory() as directory:
            empty = self._input(directory, pd.DataFrame(columns=["Row_ID", "SMILES"]))
            with self.assertRaisesRegex(ValueError, "no rows"):
                export_erba_batch(empty, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA,
                                  _Predictor(), _BatchAD(), self.spec, self.provenance)
            fatal = self._input(directory, pd.DataFrame([["x", "CCO"]], columns=["Row_ID", "SMILES"]))
            with self.assertRaisesRegex(RuntimeError, "no output"):
                export_erba_batch(fatal, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA,
                                  _Predictor(ERBAStatusCode.MODEL_INTEGRITY_FAILED),
                                  _BatchAD(), self.spec, self.provenance)
            self.assertEqual(sorted(path.name for path in Path(directory).glob("ERBA_*.xlsx")), [])

    def test_batch_completion_ignores_stale_request(self):
        tab = ErbaTab.__new__(ErbaTab)
        tab._active_batch_request_id = 2
        ErbaTab._batch_complete(tab, 1, ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA,
                                "model-v1", None, "")
        self.assertEqual(tab._active_batch_request_id, 2)

    def test_worker_progress_is_scheduled_through_after(self):
        class _Variable:
            def __init__(self):
                self.value = None

            def set(self, value):
                self.value = value

        tab = ErbaTab.__new__(ErbaTab)
        tab._active_batch_request_id = 7
        tab._model_generation = 0
        tab._request_generations = {7: 0}
        tab.batch_progress_var = _Variable()
        tab.batch_progress_value = _Variable()
        scheduled = []
        tab.after = lambda delay, callback: scheduled.append((delay, callback))
        tab._batch_progress_from_worker(7, "Predicting", 3, 10, 61)
        self.assertEqual(len(scheduled), 1)
        self.assertEqual(tab.batch_progress_var.value, None)
        scheduled[0][1]()
        self.assertEqual(tab.batch_progress_value.value, 61)
        self.assertEqual(
            tab.batch_progress_var.value,
            "61% - 3/10 - Preprocessing and prediction",
        )

        tab._model_generation = 1
        tab._batch_progress_from_worker(7, "Writing workbook", 9, 10, 90)
        scheduled[1][1]()
        self.assertEqual(tab.batch_progress_value.value, 61)
        self.assertEqual(
            tab.batch_progress_var.value,
            "61% - 3/10 - Preprocessing and prediction",
        )

    def test_worker_lower_status_discards_queued_stale_model_update(self):
        class _Variable:
            def __init__(self):
                self.value = None

            def set(self, value):
                self.value = value

        tab = ErbaTab.__new__(ErbaTab)
        tab._active_batch_request_id = 7
        tab._model_generation = 0
        tab._request_generations = {7: 0}
        tab.batch_status_var = _Variable()
        scheduled = []
        tab.after = lambda delay, callback: scheduled.append((delay, callback))

        tab._batch_status_from_worker(
            7,
            "Fetching SMILES from PubChem: 1 / 2 (50-00-0)",
        )
        tab._model_generation = 1
        scheduled[0][1]()

        self.assertIsNone(tab.batch_status_var.value)

    def test_export_avoids_existing_destination_collision(self):
        frame = pd.DataFrame([["x", "CCO"]], columns=["Row_ID", "SMILES"])
        with tempfile.TemporaryDirectory() as directory:
            existing = Path(directory, "ERBA_classification_er_alpha_results.xlsx")
            existing.write_bytes(b"prior")
            with patch("gui.erba_tab.save_erba_batch_graphs", return_value=((), None)):
                export_result = export_erba_batch(
                    self._input(directory, frame), ERBATask.CLASSIFICATION,
                    ERBASubtype.ER_ALPHA, _Predictor(), _BatchAD(),
                    self.spec, self.provenance,
                )
            self.assertEqual(existing.read_bytes(), b"prior")
            self.assertEqual(
                export_result.destination.name,
                "ERBA_classification_er_alpha_results_2.xlsx",
            )

    def test_export_uses_exact_input_parent_and_rejects_protected_parent(self):
        frame = pd.DataFrame([["x", "CCO"]], columns=["Row_ID", "SMILES"])
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_directory = root / "selected-input-folder"
            prior_output_setting = root / "Documents" / "ER_Predictor" / "Exports"
            input_directory.mkdir()
            prior_output_setting.mkdir(parents=True)
            input_path = self._input(input_directory, frame)

            with patch("gui.erba_tab.save_erba_batch_graphs", return_value=((), None)):
                export_result = export_erba_batch(
                    input_path,
                    ERBATask.CLASSIFICATION,
                    ERBASubtype.ER_ALPHA,
                    _Predictor(),
                    _BatchAD(),
                    self.spec,
                    self.provenance,
                )

            self.assertEqual(export_result.destination.parent, input_directory.resolve())
            self.assertEqual(list(prior_output_setting.glob("ERBA_*.xlsx")), [])

            protected_input = self._input(prior_output_setting, frame)
            with self.assertRaisesRegex(RuntimeError, "read-only resources"):
                export_erba_batch(
                    protected_input,
                    ERBATask.CLASSIFICATION,
                    ERBASubtype.ER_ALPHA,
                    _Predictor(),
                    _BatchAD(),
                    self.spec,
                    self.provenance,
                    forbidden_roots=(prior_output_setting,),
                )
            self.assertEqual(
                list(prior_output_setting.glob("ERBA_*_results*.xlsx")),
                [],
            )

            resource_root = root / "resource-root"
            install_root = root / "install-root"
            resource_root.mkdir()
            install_root.mkdir()
            install_input = self._input(install_root, frame)
            install_predictor = _Predictor()
            install_ad = _BatchAD()
            with self.assertRaisesRegex(RuntimeError, "read-only resources"):
                export_erba_batch(
                    install_input,
                    ERBATask.CLASSIFICATION,
                    ERBASubtype.ER_ALPHA,
                    install_predictor,
                    install_ad,
                    self.spec,
                    self.provenance,
                    forbidden_roots=(resource_root, install_root),
                )
            self.assertEqual(install_predictor.calls, [])
            self.assertEqual(install_ad.calls, [])
            self.assertEqual(list(install_root.glob("ERBA_*.xlsx")), [])

    def test_primary_predictions_suffix_duplicate_passthrough_headers_and_trust_results(self):
        input_rows = pd.DataFrame(
            [["first", "second", "spoof"]],
            columns=["Note", "Note", " Probability_Positive_1 "],
        )
        trusted = pd.DataFrame([
            {"Probability_Positive_1": 0.8, "Prediction_label": "Binding"}
        ])
        ad_rows = pd.DataFrame(
            [{
                "AD": "In-domain",
                "AD_MeanDistance": 1.25,
                "AD_DistanceThreshold": 2.5,
                "AD_Distance_InDomain": True,
                "AD_SimilarityMax": 0.75,
                "AD_SimilarityThreshold": 0.6,
                "AD_Similarity_InDomain": True,
                "AD_PC1": 0.5,
                "AD_PC2": -0.25,
            }],
            columns=ERBA_BATCH_AD_COLUMNS,
        )
        pubchem_rows = pd.DataFrame([
            {"PubChem_CID": 702, "PubChem_status": "Found"}
        ])

        primary = build_primary_predictions(
            input_rows,
            trusted,
            ad_rows,
            pubchem_rows,
        )

        self.assertEqual(
            primary.columns.tolist(),
            [
                "Note",
                "Note_input_2",
                "Probability_Positive_1",
                "Prediction_label",
                *ERBA_BATCH_AD_COLUMNS,
                "PubChem_CID",
                "PubChem_status",
            ],
        )
        self.assertEqual(len(primary.columns), len(set(primary.columns)))
        self.assertEqual(primary.iloc[0]["Probability_Positive_1"], 0.8)
        self.assertNotIn("spoof", primary.iloc[0].tolist())

    def test_ad_failure_keeps_binding_prediction_and_marks_only_ad_unavailable(self):
        frame = pd.DataFrame([["row", "CCO"]], columns=["Row_ID", "SMILES"])
        with tempfile.TemporaryDirectory() as directory:
            export_result = export_erba_batch(
                self._input(directory, frame),
                ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA,
                _Predictor(),
                _BatchAD(fail_for={"CCO"}),
                self.spec,
                self.provenance,
            )
            predictions = pd.read_excel(
                export_result.destination,
                sheet_name="Predictions",
                dtype=str,
                keep_default_na=False,
            )
            diagnostics = pd.read_excel(
                export_result.destination,
                sheet_name="Diagnostics",
                dtype=str,
                keep_default_na=False,
            )
            guide = pd.read_excel(
                export_result.destination,
                sheet_name="Guide",
                dtype=str,
                keep_default_na=False,
            )
        self.assertEqual(diagnostics.loc[0, "Result_Status"], "Predicted")
        self.assertEqual(predictions.loc[0, "Probability_Positive_1"], "0.8")
        self.assertEqual(predictions.loc[0, "Prediction_label"], "Binding")
        self.assertEqual(predictions.loc[0, "AD"], "Unavailable")
        self.assertTrue(all(
            predictions.loc[0, column] == ""
            for column in ERBA_BATCH_AD_COLUMNS[1:]
        ))
        self.assertEqual(export_result.ad_unavailable_count, 1)
        self.assertIn("route AD unavailable", export_result.ad_error)
        self.assertIn("AD evaluation warning", guide["Topic"].tolist())

    def test_workbook_write_failure_removes_reserved_workbook_and_new_graph_directory(self):
        frame = pd.DataFrame([["row", "CCO"]], columns=["Row_ID", "SMILES"])

        def create_graph_artifacts(_predictions, destination, _calculator, _fingerprints):
            graph_directory = Path(destination).with_suffix("")
            graph_directory = graph_directory.with_name(
                f"{graph_directory.name}_graphs"
            )
            graph_directory.mkdir()
            graph_path = graph_directory / "binding_class_count.png"
            graph_path.touch()
            return (graph_path,), graph_directory

        with tempfile.TemporaryDirectory() as directory:
            input_path = self._input(directory, frame)
            with patch(
                "gui.erba_tab.save_erba_batch_graphs",
                side_effect=create_graph_artifacts,
            ), patch(
                "gui.erba_tab.pd.ExcelWriter",
                side_effect=RuntimeError("write failed"),
            ):
                with self.assertRaisesRegex(RuntimeError, "write failed"):
                    export_erba_batch(
                        input_path,
                        ERBATask.CLASSIFICATION,
                        ERBASubtype.ER_ALPHA,
                        _Predictor(),
                        _BatchAD(),
                        self.spec,
                        self.provenance,
                    )
            self.assertEqual(list(Path(directory).glob("ERBA_*.xlsx")), [])
            self.assertEqual(list(Path(directory).glob("ERBA_*_graphs*")), [])

    def test_binding_graphs_use_aligned_route_ad_inputs_and_collision_safe_directory(self):
        predictions = pd.DataFrame({
            "Row_ID": ["a", "b"],
            "CAS": ["50-00-0", "64-17-5"],
            "binding_label": ["binding", "non_binding"],
            "binding_probability": [0.9, 0.1],
        })
        fingerprints = np.asarray([[1.0, 0.0], [0.0, 1.0]])
        calculator = SimpleNamespace(fitted=True)
        captured = []

        def save_ad(_calculator, matrix, _predictions, directory):
            captured.append((_calculator, matrix.copy(), Path(directory)))
            path = Path(directory, "ad_pca_plot.png")
            path.touch()
            return str(path)

        def save_decision(_calculator, matrix, _predictions, directory):
            captured.append((_calculator, matrix.copy(), Path(directory)))
            path = Path(directory, "ad_decision_plot.png")
            path.touch()
            return str(path)

        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory, "ERBA_classification_er_alpha_results.xlsx")
            destination.touch()
            prior_graph_directory = Path(directory, f"{destination.stem}_graphs")
            prior_graph_directory.mkdir()
            prior_graph = prior_graph_directory / "binding_class_count.png"
            prior_graph.write_bytes(b"prior")
            with patch("gui.erba_tab.save_ad_plot", side_effect=save_ad), patch(
                "gui.erba_tab.save_ad_decision_plot",
                side_effect=save_decision,
            ):
                graph_paths, graph_directory = save_erba_batch_graphs(
                    predictions,
                    destination,
                    calculator,
                    fingerprints,
                )

            self.assertEqual(
                graph_directory,
                Path(directory, f"{destination.stem}_graphs_2"),
            )
            self.assertEqual(
                [path.name for path in graph_paths],
                [
                    "binding_class_count.png",
                    "binding_probability_histogram.png",
                    "top_binding_chemicals.png",
                    "ad_pca_plot.png",
                    "ad_decision_plot.png",
                ],
            )
            self.assertTrue(all(path.parent == graph_directory for path in graph_paths))
            self.assertTrue(all(path.exists() for path in graph_paths))
            self.assertTrue(all(call[0] is calculator for call in captured))
            self.assertTrue(all(np.array_equal(call[1], fingerprints) for call in captured))
            self.assertTrue(all(call[2] == graph_directory for call in captured))
            self.assertEqual(prior_graph.read_bytes(), b"prior")


class EralphaBatchUsabilityContractTests(unittest.TestCase):
    class _Variable:
        def __init__(self, value=""):
            self.value = value

        def get(self):
            return self.value

        def set(self, value):
            self.value = value

    class _Widget:
        def __init__(self):
            self.state = None

        def configure(self, **kwargs):
            self.state = kwargs.get("state", self.state)

    class _Text(_Widget):
        def __init__(self):
            super().__init__()
            self.value = ""

        def delete(self, _start, _end):
            self.value = ""

        def insert(self, _start, value):
            self.value = value

    def _tab_with_batch_controls(self):
        tab = ErbaTab.__new__(ErbaTab)
        tab.project_root = Path.cwd() / "project-root"
        tab.install_root = Path.cwd() / "install-root"
        tab.batch_input_button = self._Widget()
        tab.download_template_button = self._Widget()
        tab.run_batch_button = self._Widget()
        tab.batch_input_display_var = self._Variable()
        tab._show_batch_dialog = lambda *_args: None
        return tab

    def test_batch_destination_is_read_only_and_tracks_selected_input_parent(self):
        source = Path("gui/erba_tab.py").read_text(encoding="utf-8")
        batch_source = source.split("    def _build_batch(self):", 1)[1].split(
            "    def _selected_route", 1
        )[0]
        self.assertIn('ttk.LabelFrame(self.batch_tab, text="Batch")', batch_source)
        self.assertIn('text="Result folder (same as input)"', batch_source)
        self.assertIn('text="Input xlsx"', batch_source)
        self.assertIn('text="Download template"', batch_source)
        self.assertIn('text="Run batch"', batch_source)
        self.assertIn("batch_progress_var", batch_source)
        self.assertIn("self.batch_result = tk.Text", batch_source)
        self.assertNotIn("textvariable=self.batch_status_var", batch_source)
        self.assertEqual(source.count("self.workflow_status_label.grid("), 1)
        self.assertIn('value="0% - 0/0 - Ready"', source)
        self.assertNotIn("batch_output_button", source)
        self.assertNotIn("batch_output_var", source)
        self.assertNotIn("_choose_output_dir", source)
        self.assertNotIn("askdirectory", source)

        tab = self._tab_with_batch_controls()
        tab.batch_input_var = self._Variable()
        tab.batch_destination_var = self._Variable("C:/Documents/ER_Predictor/Exports")
        tab.batch_status_var = self._Variable()
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory, "selected.xlsx")
            input_path.touch()
            with patch(
                "gui.erba_tab.filedialog.askopenfilename",
                return_value=str(input_path),
            ):
                tab.browse_batch_input()
            self.assertEqual(tab.batch_input_var.value, str(input_path))
            self.assertEqual(
                tab.batch_destination_var.value,
                str(input_path.resolve().parent),
            )
            self.assertEqual(
                erba_batch_destination_display(input_path),
                str(input_path.resolve().parent),
            )

    def test_selected_protected_input_shows_unavailable_destination(self):
        tab = self._tab_with_batch_controls()
        tab.batch_input_var = self._Variable()
        tab.batch_destination_var = self._Variable()
        tab.batch_status_var = self._Variable()
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory, "selected.xlsx")
            input_path.touch()
            with patch(
                "gui.erba_tab.filedialog.askopenfilename",
                return_value=str(input_path),
            ), patch(
                "gui.erba_tab.validate_mutable_directory",
                side_effect=RuntimeError("read-only resources"),
            ):
                tab.browse_batch_input()
            self.assertEqual(tab.batch_input_var.value, str(input_path.resolve()))
            self.assertIn("Unavailable", tab.batch_destination_var.value)
            self.assertIn(str(input_path.resolve().parent), tab.batch_destination_var.value)
            self.assertIn("Copy or download", tab.batch_status_var.value)
            self.assertIn("read-only resources", tab.batch_status_var.value)

    def test_new_eralpha_batch_replaces_prior_completion_before_worker_starts(self):
        tab = self._tab_with_batch_controls()
        tab.batch_input_var = self._Variable()
        tab.batch_destination_var = self._Variable()
        tab.batch_status_var = self._Variable()
        tab.batch_progress_var = self._Variable()
        tab.batch_progress_value = self._Variable()
        tab.batch_result = self._Text()
        tab.batch_result.value = "Prior completed result"
        tab._active_batch_request_id = None
        tab._model_generation = 0
        tab._request_generations = {7: 0}
        tab._request_predictors = {}
        tab._snapshot = lambda _workflow: (
            7,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
        )
        tab.predictor = SimpleNamespace(
            preflight=lambda _task, _subtype: ERBARoutePreflight(
                task=ERBATask.CLASSIFICATION,
                subtype=ERBASubtype.ER_ALPHA,
                status_code=ERBAStatusCode.OK,
                status_message="ready",
            )
        )
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory, "input.xlsx")
            source.touch()
            tab.batch_input_var.set(str(source))
            with patch("gui.erba_tab.threading.Thread") as worker:
                tab.batch_predict_clicked()

        worker.assert_called_once()
        self.assertEqual(tab.batch_result.value, BATCH_RUNNING_RESULT)
        self.assertEqual(
            tab.batch_progress_var.value,
            "0% - 0/0 - Reading input workbook",
        )
        self.assertEqual(tab.batch_status_var.value, "Batch prediction started.")

    def test_distinct_install_root_rejects_erba_input_template_and_execution(self):
        tab = self._tab_with_batch_controls()
        tab.batch_input_var = self._Variable()
        tab.batch_destination_var = self._Variable()
        tab.batch_status_var = self._Variable()
        tab.batch_progress_var = self._Variable()
        tab.batch_progress_value = self._Variable()
        tab.batch_result = self._Text()
        tab._active_batch_request_id = None
        tab._started_at = {1: 0.0}
        tab._snapshot = lambda _workflow: (
            1,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
        )
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            resource_root = root / "resource-root"
            install_root = root / "install-root"
            output_root = root / "single-artifacts"
            resource_root.mkdir()
            install_root.mkdir()
            output_root.mkdir()
            input_path = install_root / "input.xlsx"
            input_path.touch()
            template_path = install_root / "template.xlsx"
            tab.project_root = resource_root
            tab.install_root = install_root
            tab.output_root = output_root

            with patch(
                "gui.erba_tab.filedialog.askopenfilename",
                return_value=str(input_path),
            ):
                tab.browse_batch_input()
            self.assertIn("Unavailable", tab.batch_destination_var.value)

            with patch("gui.erba_tab.threading.Thread") as worker:
                tab.batch_predict_clicked()
            worker.assert_not_called()

            with patch(
                "gui.erba_tab.filedialog.asksaveasfilename",
                return_value=str(template_path),
            ), patch("gui.erba_tab.messagebox.showerror") as show_error:
                tab.download_template_clicked()

            show_error.assert_called_once()
            self.assertIn("read-only resources", str(show_error.call_args))
            self.assertEqual(tab.batch_progress_var.value, "100% - 0/0 - Failed")
            self.assertIn("Batch prediction failed", tab.batch_result.value)
            self.assertFalse(template_path.exists())
            self.assertEqual(list(output_root.iterdir()), [])

    def test_downloaded_template_has_only_the_required_cas_header_and_becomes_input(self):
        tab = self._tab_with_batch_controls()
        tab.batch_input_var = self._Variable()
        tab.batch_destination_var = self._Variable()
        tab.batch_status_var = self._Variable()
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory, "ERalpha_batch_template.xlsx")
            with patch("gui.erba_tab.filedialog.asksaveasfilename", return_value=str(destination)):
                with patch("gui.erba_tab.messagebox.showinfo"):
                    tab.download_template_clicked()
            template = pd.read_excel(destination, dtype=str, keep_default_na=False)
        self.assertEqual(template.columns.tolist(), ["CAS"])
        self.assertTrue(template.empty)
        self.assertEqual(tab.batch_input_var.value, str(destination))
        self.assertEqual(
            tab.batch_destination_var.value,
            str(destination.resolve().parent),
        )

    def test_unwritable_input_parent_fails_before_worker_and_keeps_destination_visible(self):
        tab = self._tab_with_batch_controls()
        tab._snapshot = lambda _workflow: (
            1,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
        )
        tab._started_at = {1: 0.0}
        tab._active_batch_request_id = None
        tab.project_root = Path("C:/application/resources")
        tab.install_root = Path("C:/application")
        tab.batch_input_var = self._Variable()
        tab.batch_destination_var = self._Variable("C:/Documents/ER_Predictor/Exports")
        tab.batch_status_var = self._Variable()
        tab.batch_progress_var = self._Variable()
        tab.batch_progress_value = self._Variable()
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory, "input.xlsx")
            input_path.touch()
            tab.batch_input_var.set(str(input_path))
            with patch(
                "gui.erba_tab.validate_mutable_directory",
                side_effect=RuntimeError("input folder is not writable"),
            ), patch("gui.erba_tab.threading.Thread") as worker:
                tab.batch_predict_clicked()
            worker.assert_not_called()
            self.assertEqual(
                tab.batch_destination_var.value,
                str(input_path.resolve().parent),
            )
            self.assertIn("saved beside the input workbook", tab.batch_status_var.value)
            self.assertIn("not writable", tab.batch_status_var.value)
            self.assertEqual(tab.batch_progress_value.value, 100)
            self.assertEqual(tab.batch_progress_var.value, "100% - 0/0 - Failed")
            self.assertIsNone(tab._active_batch_request_id)
            self.assertEqual(list(input_path.parent.glob("ERBA_*.xlsx")), [])

    def test_batch_dialogs_route_early_failure_and_terminal_outcomes_directly(self):
        tab = self._tab_with_batch_controls()
        tab._show_batch_dialog = ErbaTab._show_batch_dialog
        tab.batch_input_var = self._Variable("")
        tab.batch_destination_var = self._Variable()
        tab.batch_status_var = self._Variable()
        tab.batch_progress_var = self._Variable()
        tab.batch_progress_value = self._Variable()
        tab.batch_result = self._Text()
        tab._started_at = {1: 0.0}
        tab._snapshot = lambda _workflow: (
            1,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
        )
        callbacks = []
        tab.after = lambda _delay, callback: callbacks.append(callback)

        with patch("gui.erba_tab.messagebox.showinfo") as show_info, patch(
            "gui.erba_tab.messagebox.showwarning"
        ) as show_warning, patch(
            "gui.erba_tab.messagebox.showerror"
        ) as show_error:
            tab.batch_predict_clicked()
            show_error.assert_called_once()
            self.assertIn("Choose an input xlsx", str(show_error.call_args))
            self.assertEqual(callbacks, [])
            show_error.reset_mock()

            tab._finish_duration_ms = lambda _request_id: 0
            tab._set_batch_progress = lambda *_args: None
            tab._snapshot_is_current = lambda *_args: True
            tab._route_changed = lambda _workflow: None
            tab._emit = lambda *_args, **_kwargs: None

            destination = Path(
                "C:/published/ERBA_classification_er_alpha_results.xlsx"
            )
            successful = ERBABatchExportResult(
                destination=destination,
                count=2,
                binding_count=1,
                non_binding_count=1,
                not_predicted_count=0,
                ad_in_domain_count=1,
                ad_out_of_domain_count=1,
                ad_unavailable_count=0,
                graph_paths=(destination.parent / "graphs" / "classes.png",),
                graph_directory=destination.parent / "graphs",
            )
            tab._active_batch_request_id = 2
            tab._batch_complete(
                2,
                ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA,
                "model-v1",
                successful,
                "",
            )
            show_info.assert_called_once()
            self.assertIn(str(destination), show_info.call_args.args[1])
            self.assertEqual(callbacks, [])
            show_info.reset_mock()

            mixed = ERBABatchExportResult(
                destination=destination,
                count=2,
                binding_count=1,
                non_binding_count=0,
                not_predicted_count=1,
                ad_in_domain_count=1,
                ad_out_of_domain_count=0,
                ad_unavailable_count=1,
                graph_paths=(),
                graph_directory=None,
                graph_error="No route-specific AD rows were available.",
                ad_error="Row 2: AD unavailable",
            )
            tab._active_batch_request_id = 3
            tab._batch_complete(
                3,
                ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA,
                "model-v1",
                mixed,
                "",
            )
            show_info.assert_called_once()
            self.assertIn("Not predicted: 1", str(show_info.call_args))
            self.assertIn("Graph details:", str(show_info.call_args))
            show_warning.assert_not_called()
            self.assertEqual(callbacks, [])
            show_info.reset_mock()

            tab._active_batch_request_id = 4
            unavailable = ERBABatchExportResult(
                destination=destination,
                count=2,
                binding_count=0,
                non_binding_count=0,
                not_predicted_count=2,
                ad_in_domain_count=0,
                ad_out_of_domain_count=0,
                ad_unavailable_count=0,
                graph_paths=(),
                graph_directory=None,
                graph_error="No route-specific AD rows were available.",
            )
            tab._batch_complete(
                4,
                ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA,
                "model-v1",
                unavailable,
                "",
            )
            show_info.assert_called_once()
            self.assertIn("No rows could be predicted.", str(show_info.call_args))
            show_warning.assert_not_called()
            show_info.reset_mock()

            tab._active_batch_request_id = 5
            tab._batch_complete(
                5,
                ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA,
                "model-v1",
                None,
                "write failed",
            )
            show_error.assert_called_once()
            self.assertIn("write failed", str(show_error.call_args))
            show_info.assert_not_called()
            show_warning.assert_not_called()
            self.assertEqual(callbacks, [])

    def test_batch_worker_ignores_single_output_root_and_exports_beside_input(self):
        tab = ErbaTab.__new__(ErbaTab)
        task, subtype = ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA
        tab.predictor = _Predictor()
        tab.erba_ad = _BatchAD()
        tab.catalog = {(task, subtype): ErbaExcelContractTests.spec}
        tab.catalog_payload = ErbaExcelContractTests.provenance
        tab._batch_progress_from_worker = lambda *_args: None
        completed = []
        tab.after = lambda _delay, callback: callback()
        tab._batch_complete = lambda *args: completed.append(args)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_directory = root / "selected-input-folder"
            prior_output_setting = root / "Documents" / "ER_Predictor" / "Exports"
            input_directory.mkdir()
            prior_output_setting.mkdir(parents=True)
            tab.project_root = root / "application-resources"
            tab.install_root = root / "application"
            tab.output_root = prior_output_setting
            input_path = input_directory / "input.xlsx"
            pd.DataFrame([["x", "CCO"]], columns=["Row_ID", "SMILES"]).to_excel(
                input_path,
                index=False,
            )

            with patch("gui.erba_tab.save_erba_batch_graphs", return_value=((), None)):
                tab._batch_work(1, task, subtype, "model-v1", str(input_path))

            export_result = completed[0][4]
            self.assertEqual(
                export_result.destination.parent,
                input_directory.resolve(),
            )
            self.assertEqual(list(prior_output_setting.glob("*.xlsx")), [])

    def test_completion_uses_exported_parent_for_displayed_directory_and_exact_destination_everywhere(self):
        tab = self._tab_with_batch_controls()
        tab._active_batch_request_id = 1
        tab.batch_destination_var = self._Variable("C:/Documents/ER_Predictor/Exports")
        tab.batch_status_var = self._Variable()
        tab.batch_result = self._Text()
        tab._finish_duration_ms = lambda _request_id: 0
        tab._set_batch_progress = lambda *args: None
        tab._snapshot_is_current = lambda *args: True
        tab._route_changed = lambda _workflow: None
        tab._emit = lambda *args, **kwargs: None
        destination = Path("C:/published/output/ERBA_classification_er_alpha_results.xlsx")
        export_result = ERBABatchExportResult(
            destination=destination,
            count=7,
            binding_count=3,
            non_binding_count=2,
            not_predicted_count=2,
            ad_in_domain_count=2,
            ad_out_of_domain_count=2,
            ad_unavailable_count=1,
            graph_paths=(
                destination.parent / "graphs" / "binding_class_count.png",
                destination.parent / "graphs" / "ad_pca_plot.png",
            ),
            graph_directory=destination.parent / "graphs",
            ad_error="Row 7: AD unavailable",
        )

        tab._batch_complete(
            1,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
            export_result,
            "",
        )

        self.assertEqual(tab.batch_destination_var.value, str(destination.parent))
        self.assertEqual(
            tab.batch_status_var.value,
            f"Batch prediction completed: {destination}",
        )
        self.assertIn(str(destination), tab.batch_result.value)
        for detail in (
            "Total rows: 7",
            "Predicted: 5",
            "Binding: 3",
            "Non-binding: 2",
            "Not predicted: 2",
            "AD In-domain: 2",
            "AD Out-of-domain: 2",
            "AD Unavailable: 1",
            PREDICTED_AD_COUNT_NOTE,
            "Graph files: 2",
            f"Graph directory: {destination.parent / 'graphs'}",
            "Applicability-domain details: Row 7: AD unavailable",
            "historically exposed",
        ):
            self.assertIn(detail, tab.batch_result.value)
        self.assertTrue(all(
            widget.state == "normal"
            for widget in (
                tab.batch_input_button,
                tab.download_template_button,
                tab.run_batch_button,
            )
        ))

    def test_failure_and_stale_completion_restore_or_preserve_control_lock_correctly(self):
        tab = self._tab_with_batch_controls()
        tab._active_batch_request_id = 1
        tab.batch_destination_var = self._Variable()
        tab.batch_status_var = self._Variable()
        tab._finish_duration_ms = lambda _request_id: 0
        tab._set_batch_progress = lambda *args: None
        tab._snapshot_is_current = lambda *args: True
        tab._route_changed = lambda _workflow: None
        tab._emit = lambda *args, **kwargs: None

        tab.batch_result = self._Text()
        tab._batch_complete(
            1,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
            None,
            "write failed",
        )

        self.assertIn("write failed", tab.batch_status_var.value)
        self.assertIn("write failed", tab.batch_result.value)
        self.assertTrue(all(
            widget.state == "normal"
            for widget in (
                tab.batch_input_button,
                tab.download_template_button,
                tab.run_batch_button,
            )
        ))

        tab._set_batch_controls_active(True)
        tab._active_batch_request_id = None
        tab._route_changed = lambda _workflow: tab._set_batch_controls_active(False)
        tab._batch_complete(
            2,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "model-v1",
            None,
            "stale",
        )
        self.assertTrue(all(
            widget.state == "normal"
            for widget in (
                tab.batch_input_button,
                tab.download_template_button,
                tab.run_batch_button,
            )
        ))

    def test_generation_stale_completion_restores_controls_without_terminal_progress(self):
        tab = self._tab_with_batch_controls()
        tab._set_batch_controls_active(True)
        tab._active_batch_request_id = 7
        tab._finish_duration_ms = lambda _request_id: 0
        progress_calls = []
        tab._set_batch_progress = lambda *args: progress_calls.append(args)
        tab._snapshot_is_current = lambda *args: False
        tab._route_changed = lambda _workflow: None
        tab._emit = lambda *args, **kwargs: None

        tab._batch_complete(
            7,
            ERBATask.CLASSIFICATION,
            ERBASubtype.ER_ALPHA,
            "old-model",
            None,
            "obsolete completion",
        )

        self.assertEqual(progress_calls, [])
        self.assertIsNone(tab._active_batch_request_id)
        self.assertTrue(
            all(
                widget.state == "normal"
                for widget in (
                    tab.batch_input_button,
                    tab.download_template_button,
                    tab.run_batch_button,
                )
            )
        )

class _ErtaBatchVariable:
    def __init__(self, value=None):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class _ErtaBatchControl:
    def __init__(self):
        self.state = "normal"

    def configure(self, **kwargs):
        self.state = kwargs.get("state", self.state)


class _ErtaBatchText:
    def __init__(self):
        self.value = ""
        self.state = "disabled"

    def configure(self, **kwargs):
        self.state = kwargs.get("state", self.state)

    def delete(self, _start, _end):
        self.value = ""

    def insert(self, _start, value):
        self.value = value


class _ErtaBatchPredictor:
    model_name = "legacy"

    def is_loaded(self):
        return True

    def predict_dataframe(self, frame, smiles_col):
        result = frame.copy()
        result["Prediction_label"] = "Positive"
        return result, pd.DataFrame([[0.0] for _ in range(len(frame))])


class ErtaBatchUsabilityContractTests(unittest.TestCase):
    def _window(self):
        window = MainWindow.__new__(MainWindow)
        window.project_root = str(Path.cwd() / "project-root")
        window.install_root = str(Path.cwd() / "install-root")
        window.batch_input_button = _ErtaBatchControl()
        window.download_template_button = _ErtaBatchControl()
        window.run_batch_button = _ErtaBatchControl()
        window.model_entry = _ErtaBatchControl()
        window.model_browse_button = _ErtaBatchControl()
        window.model_reload_button = _ErtaBatchControl()
        window.ad_entry = _ErtaBatchControl()
        window.ad_browse_button = _ErtaBatchControl()
        window.ad_reload_button = _ErtaBatchControl()
        window.batch_result = _ErtaBatchText()
        window.batch_destination_var = _ErtaBatchVariable()
        window.batch_progress_var = _ErtaBatchVariable()
        window.batch_progress_value = _ErtaBatchVariable()
        window._batch_active = False
        window._batch_request_id = 0
        window._active_batch_request_id = None
        window._erta_reload_in_flight = False
        window._active_batch_total = 1
        window.statuses = []
        window.set_status = window.statuses.append
        window.main_thread = threading.current_thread()
        window.after = lambda _delay, callback: callback()
        window.ui = lambda callback, *args: callback(*args)
        window.run_threaded = lambda job: job()
        return window

    def test_compact_erta_batch_controls_and_cas_notice_are_present(self):
        source = Path("gui/main_window.py").read_text(encoding="utf-8")
        batch_source = source.split("    def _build_batch_tab(self):", 1)[1].split(
            "    # ---------- UI helpers ----------", 1
        )[0]
        self.assertIn('text="Input xlsx"', source)
        self.assertIn('text="Result folder (same as input)"', source)
        self.assertIn('text="Download template"', source)
        self.assertIn('text="Run batch"', source)
        self.assertIn("required CAS column", source)
        self.assertNotIn("output_dir_button", source)
        self.assertNotIn("output_dir_var", source)
        self.assertNotIn("output_dir_display_var", source)
        self.assertNotIn("browse_output_dir", source)
        self.assertNotIn("_active_batch_output_dir", source)
        self.assertNotIn("askdirectory", batch_source)
        self.assertIn("def validated_single_output_dir(self)", source)
        self.assertIn("self.output_root", source)
        self.assertNotIn('text="Run batch prediction"', source)
        self.assertNotIn('bg="#0969da"', source)
        self.assertIn("self.batch_result = tk.Text(", batch_source)
        self.assertNotIn("self.batch_tree = ttk.Treeview", batch_source)
        self.assertNotIn('text="Applicability domain"', batch_source)
        self.assertNotIn("self.graph_combo", batch_source)

    def test_selected_erta_input_updates_read_only_destination_display(self):
        window = self._window()
        window.batch_input_var = _ErtaBatchVariable()
        window.batch_input_display_var = _ErtaBatchVariable()
        window.batch_destination_var = _ErtaBatchVariable(
            "C:/Documents/ER_Predictor/Exports"
        )
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory, "selected.xlsx")
            input_path.touch()
            with patch(
                "gui.main_window.filedialog.askopenfilename",
                return_value=str(input_path),
            ):
                window.browse_batch_input()
            self.assertEqual(window.batch_input_var.value, str(input_path))
            self.assertEqual(window.batch_input_display_var.value, input_path.name)
            self.assertEqual(
                window.batch_destination_var.value,
                str(input_path.resolve().parent),
            )
            self.assertEqual(
                erba_batch_destination_display(input_path),
                str(input_path.resolve().parent),
            )

    def test_erta_output_path_reservation_is_collision_safe_in_input_parent(self):
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory, "input.xlsx")
            input_path.touch()
            first = allocate_erta_output_path(input_path, "legacy")
            first.write_bytes(b"prior")
            second = allocate_erta_output_path(input_path, "legacy")
            try:
                self.assertEqual(first.parent, input_path.resolve().parent)
                self.assertEqual(second.parent, input_path.resolve().parent)
                self.assertEqual(first.name, "ERTA_input_legacy_prediction.xlsx")
                self.assertEqual(second.name, "ERTA_input_legacy_prediction_2.xlsx")
                self.assertEqual(first.read_bytes(), b"prior")
            finally:
                first.unlink(missing_ok=True)
                second.unlink(missing_ok=True)

    def test_erta_carsrn_alias_resolves_valid_cas_and_preserves_invalid_rows(self):
        window = MainWindow.__new__(MainWindow)
        window.set_status = lambda _message: None
        statuses = []
        progress = []
        frame = pd.DataFrame(
            {"CARSRN": ["50-00-0", 12345, "", None]},
        )
        with patch(
            "gui.main_window.cas_to_smiles",
            return_value={"CanonicalSMILES": "C=O", "PubChem_CID": 712},
        ) as lookup:
            prepared = window.prepare_batch_input(
                frame,
                status_callback=statuses.append,
                progress_callback=lambda stage, current, total, percent: (
                    progress.append((stage, current, total, percent))
                ),
            )

        lookup.assert_called_once_with("50-00-0")
        self.assertEqual(
            statuses,
            ["Fetching SMILES from PubChem: 1 / 4 (50-00-0)"],
        )
        self.assertEqual(
            progress[-1],
            ("Resolving CAS/SMILES", 4, 4, 35),
        )
        self.assertEqual(prepared["CAS"].tolist()[:3], ["50-00-0", 12345, ""])
        self.assertEqual(prepared["SMILES"].tolist(), ["C=O", "", "", ""])
        self.assertEqual(prepared.loc[0, "PubChem_status"], "Found")
        self.assertIn("CAS is invalid", prepared.loc[1, "PubChem_status"])
        self.assertEqual(
            prepared.loc[2, "PubChem_status"],
            "Skipped: CAS is empty",
        )
        self.assertEqual(
            prepared.loc[3, "PubChem_status"],
            "Skipped: CAS is empty",
        )

    def test_erta_rejects_workbooks_without_a_recognized_input_column(self):
        window = MainWindow.__new__(MainWindow)
        with self.assertRaisesRegex(
            ValueError,
            "recognized CAS column.*recognized SMILES column",
        ):
            window.prepare_batch_input(
                pd.DataFrame({"Analyst note": ["not an input"]})
            )

    def test_success_then_protected_erta_batch_replaces_stale_result_with_failure(self):
        window = self._window()
        window.predictor = _ErtaBatchPredictor()
        window.batch_input_var = _ErtaBatchVariable()
        started = []
        window.run_threaded = started.append
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory, "input.xlsx")
            pd.DataFrame({"CAS": ["50-00-0"]}).to_excel(input_path, index=False)
            window.project_root = directory
            window.install_root = str(Path(directory).parent / "install-root")
            window.batch_input_var.set(str(input_path))
            window._set_batch_progress(None, "Completed", 1, 1, 100)
            window._set_erta_batch_result("Prior successful batch")

            with patch("gui.main_window.messagebox.showerror") as show_error:
                window.batch_predict_clicked()

            self.assertFalse(window._batch_active)
            self.assertEqual(
                window.batch_destination_var.value,
                str(input_path.resolve().parent),
            )
            show_error.assert_called_once()
            self.assertIn(
                "writable folder outside the application files",
                show_error.call_args.args[1],
            )
            self.assertIn("read-only resources", show_error.call_args.args[1])
            self.assertEqual(window.batch_progress_var.value, "100% - 0/0 - Failed")
            self.assertIn("Batch prediction failed", window.batch_result.value)
            self.assertNotIn("Prior successful batch", window.batch_result.value)
            self.assertEqual(started, [])
            self.assertEqual(list(input_path.parent.glob("*_prediction*.xlsx")), [])

    def test_distinct_install_root_rejects_erta_batch_and_template_without_fallback(self):
        window = self._window()
        window.predictor = _ErtaBatchPredictor()
        window.batch_input_var = _ErtaBatchVariable()
        window.batch_input_display_var = _ErtaBatchVariable()
        started = []
        window.run_threaded = started.append
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            resource_root = root / "resource-root"
            install_root = root / "install-root"
            output_root = root / "single-artifacts"
            resource_root.mkdir()
            install_root.mkdir()
            output_root.mkdir()
            input_path = install_root / "input.xlsx"
            pd.DataFrame({"CAS": ["50-00-0"]}).to_excel(input_path, index=False)
            template_path = install_root / "Template.xlsx"
            window.project_root = str(resource_root)
            window.install_root = str(install_root)
            window.output_root = str(output_root)
            with patch(
                "gui.main_window.filedialog.askopenfilename",
                return_value=str(input_path),
            ):
                window.browse_batch_input()
            self.assertIn("Unavailable", window.batch_destination_var.value)

            with patch("gui.main_window.messagebox.showerror") as show_error:
                window.batch_predict_clicked()
                with patch(
                    "gui.main_window.filedialog.asksaveasfilename",
                    return_value=str(template_path),
                ):
                    window.download_template_clicked()

            self.assertEqual(show_error.call_count, 2)
            self.assertTrue(
                all(
                    "read-only resources" in call.args[1]
                    for call in show_error.call_args_list
                )
            )
            self.assertEqual(window.batch_progress_var.value, "100% - 0/0 - Failed")
            self.assertIn("Batch prediction failed", window.batch_result.value)
            self.assertEqual(started, [])
            self.assertFalse(template_path.exists())
            self.assertEqual(list(output_root.iterdir()), [])

    def test_portable_single_artifact_root_under_install_remains_writable(self):
        window = self._window()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            resource_root = root / "resource-root"
            install_root = root / "install-root"
            single_output_root = install_root / "ER_Predictor_UserData" / "Exports"
            resource_root.mkdir()
            install_root.mkdir()
            window.project_root = str(resource_root)
            window.install_root = str(install_root)
            window.output_root = str(single_output_root)

            validated = window.validated_single_output_dir()

            self.assertEqual(Path(validated), single_output_root.resolve())
            self.assertTrue(single_output_root.is_dir())

    def test_template_contains_only_the_required_cas_header(self):
        window = self._window()
        window.project_root = str(Path.cwd() / "project-root")
        window.install_root = str(Path.cwd() / "install-root")
        window.batch_input_var = _ErtaBatchVariable()
        window.batch_input_display_var = _ErtaBatchVariable()
        window.batch_destination_var = _ErtaBatchVariable()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "Template.xlsx")
            with patch("gui.main_window.filedialog.asksaveasfilename", return_value=str(path)), \
                 patch("gui.main_window.messagebox.showinfo"):
                window.download_template_clicked()
            self.assertEqual(pd.read_excel(path).columns.tolist(), ["CAS"])
            self.assertEqual(
                window.batch_destination_var.value,
                str(path.resolve().parent),
            )

    def test_new_erta_batch_replaces_prior_completion_before_worker_starts(self):
        window = self._window()
        window.predictor = _ErtaBatchPredictor()
        window.batch_input_var = _ErtaBatchVariable()
        window.batch_result.value = "Prior completed result"
        jobs = []
        window.run_threaded = jobs.append
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory, "input.xlsx")
            source.touch()
            window.batch_input_var.set(str(source))
            window.batch_predict_clicked()

        self.assertEqual(len(jobs), 1)
        self.assertEqual(window.batch_result.value, BATCH_RUNNING_RESULT)
        self.assertEqual(
            window.batch_progress_var.value,
            "0% - 0/0 - Reading input workbook",
        )
        self.assertEqual(window.statuses[-1], "Batch prediction started.")

    def test_erta_batch_and_reload_are_bidirectionally_exclusive(self):
        window = self._window()
        window.model_path_var = _ErtaBatchVariable("C:/models/reloaded.keras")
        window.loaded_model_path_var = _ErtaBatchVariable()
        window.ad_ref_path_var = _ErtaBatchVariable("C:/models/reference.xlsx")

        class Predictor:
            model_path = "C:/models/reloaded.keras"
            model_name = "reloaded"
            input_shape = 2048

            def __init__(self):
                self.loaded = []

            def load_model(self, path):
                self.loaded.append(path)

        window.predictor = Predictor()
        jobs = []
        window.run_threaded = jobs.append
        window.batch_result.value = "Prior completed result"
        window.batch_progress_var.value = "100% - 2/2 - Completed"
        window._set_batch_controls_active(True)
        option_controls = (
            window.model_entry,
            window.model_browse_button,
            window.model_reload_button,
            window.ad_entry,
            window.ad_browse_button,
            window.ad_reload_button,
        )
        self.assertTrue(all(control.state == "disabled" for control in option_controls))

        with patch("gui.main_window.filedialog.askopenfilename") as browse:
            window.browse_model()
            window.browse_ad_reference()
            window.load_model_clicked()
            window.fit_ad_clicked()
        browse.assert_not_called()
        self.assertEqual(jobs, [])

        window._set_batch_controls_active(False)
        window.load_model_clicked()
        self.assertTrue(window._erta_reload_in_flight)
        self.assertEqual(window.run_batch_button.state, "disabled")
        self.assertTrue(all(control.state == "disabled" for control in option_controls))
        self.assertEqual(len(jobs), 1)

        with patch("gui.main_window.messagebox.showerror") as show_error:
            window.batch_predict_clicked()
        show_error.assert_called_once()
        self.assertIn("reload is running", str(show_error.call_args))

        with patch("gui.main_window.messagebox.showinfo") as show_info:
            jobs[0]()
        show_info.assert_called_once()
        self.assertFalse(window._erta_reload_in_flight)
        self.assertEqual(window.run_batch_button.state, "normal")
        self.assertTrue(all(control.state == "normal" for control in option_controls))
        self.assertEqual(
            window.loaded_model_path_var.value,
            "C:/models/reloaded.keras",
        )
        self.assertEqual(
            window.batch_result.value,
            "Run a batch to show the completion summary.",
        )
        self.assertEqual(window.batch_progress_var.value, "0% - 0/0 - Ready")

    def test_erta_failed_ad_reload_unlocks_batch_without_clearing_prior_result(self):
        window = self._window()
        window.batch_result.value = "Prior completed result"
        window._set_erta_reload_active(True)

        with patch("gui.main_window.messagebox.showerror") as show_error:
            window._finish_erta_ad_reload(
                "C:/models/reference.xlsx",
                RuntimeError("reference load failed"),
            )

        show_error.assert_called_once_with(
            "AD fitting failed",
            "reference load failed",
        )
        self.assertFalse(window._erta_reload_in_flight)
        self.assertEqual(window.run_batch_button.state, "normal")
        self.assertEqual(window.batch_result.value, "Prior completed result")

    def test_erta_empty_message_reload_exception_is_not_treated_as_success(self):
        window = self._window()
        window.batch_result.value = "Prior completed result"
        window._set_erta_reload_active(True)

        with patch("gui.main_window.messagebox.showinfo") as show_info, patch(
            "gui.main_window.messagebox.showerror"
        ) as show_error:
            window._finish_erta_model_reload(
                "C:/models/reference.keras",
                RuntimeError(""),
            )

        show_error.assert_called_once_with("Model load failed", "")
        show_info.assert_not_called()
        self.assertEqual(window.batch_result.value, "Prior completed result")

    def test_batch_controls_lock_and_restore_with_terminal_progress(self):
        window = self._window()
        window._active_batch_request_id = 1
        window._set_batch_controls_active(True)
        self.assertTrue(window._batch_active)
        self.assertEqual(
            [control.state for control in (
                window.batch_input_button, window.download_template_button,
                window.run_batch_button,
            )],
            ["disabled"] * 3,
        )
        self.assertTrue(
            all(
                control.state == "disabled"
                for control in (
                    window.model_entry,
                    window.model_browse_button,
                    window.model_reload_button,
                    window.ad_entry,
                    window.ad_browse_button,
                    window.ad_reload_button,
                )
            )
        )
        window._finish_batch(1, False, "Batch prediction failed: write failed")
        self.assertFalse(window._batch_active)
        self.assertEqual(window.batch_progress_var.value, "100% - 0/1 - Failed")
        self.assertEqual(
            [control.state for control in (
                window.batch_input_button, window.download_template_button,
                window.run_batch_button,
            )],
            ["normal"] * 3,
        )
        self.assertTrue(
            all(
                control.state == "normal"
                for control in (
                    window.model_entry,
                    window.model_browse_button,
                    window.model_reload_button,
                    window.ad_entry,
                    window.ad_browse_button,
                    window.ad_reload_button,
                )
            )
        )

    def test_progress_updates_are_scheduled_through_after(self):
        window = self._window()
        window._active_batch_request_id = 7
        scheduled = []
        window.after = lambda delay, callback: scheduled.append((delay, callback))
        window._batch_progress_from_worker(7, "Writing workbook", 3, 4, 85)
        self.assertEqual(window.batch_progress_var.value, None)
        self.assertEqual(len(scheduled), 1)
        scheduled[0][1]()
        self.assertEqual(window.batch_progress_var.value, "85% - 3/4 - Writing workbook")

    def test_erta_queued_progress_and_status_ignore_stale_request(self):
        window = self._window()
        window._active_batch_request_id = 2
        window.batch_progress_var.set("current progress")
        window.batch_progress_value.set(44)
        scheduled = []
        window.after = lambda delay, callback: scheduled.append((delay, callback))

        window._batch_progress_from_worker(1, "Writing workbook", 1, 1, 90)
        window._batch_status_from_worker(
            1,
            "Fetching SMILES from PubChem: 1 / 1 (50-00-0)",
        )
        for _, callback in scheduled:
            callback()

        self.assertEqual(window.batch_progress_var.value, "current progress")
        self.assertEqual(window.batch_progress_value.value, 44)
        self.assertEqual(window.statuses, [])

    def test_erta_published_mixed_and_all_unavailable_batches_use_info_dialog(self):
        for mol_valid, expected_not_predicted, expected_phrase in (
            ([True, False], 1, "1 row(s) could not be predicted."),
            ([False, False], 2, "No rows could be predicted."),
        ):
            window = self._window()
            window._active_batch_request_id = 1
            window._active_batch_total = 2
            result = pd.DataFrame(
                {
                    "Prediction_label": ["Positive", "Negative"],
                    "Mol_valid": mol_valid,
                }
            )
            with tempfile.TemporaryDirectory() as directory:
                destination = Path(directory, "published.xlsx")
                destination.touch()
                with patch("gui.main_window.messagebox.showinfo") as show_info, patch(
                    "gui.main_window.messagebox.showwarning"
                ) as show_warning, patch(
                    "gui.main_window.messagebox.showerror"
                ) as show_error:
                    window._complete_erta_batch(
                        1,
                        result,
                        str(destination),
                        [],
                        "Optional applicability-domain results were not generated.",
                        "No route-specific AD rows were available.",
                    )

            show_info.assert_called_once()
            self.assertIn(
                f"Not predicted: {expected_not_predicted}",
                show_info.call_args.args[1],
            )
            self.assertIn(expected_phrase, show_info.call_args.args[1])
            self.assertIn(
                ERTA_LEGACY_VALID_COUNT_NOTE,
                show_info.call_args.args[1],
            )
            self.assertIn(
                PREDICTED_AD_COUNT_NOTE,
                show_info.call_args.args[1],
            )
            show_warning.assert_not_called()
            show_error.assert_not_called()

    def test_erta_no_output_failure_uses_error_dialog_only(self):
        window = self._window()
        window._active_batch_request_id = 1
        with patch("gui.main_window.messagebox.showinfo") as show_info, patch(
            "gui.main_window.messagebox.showerror"
        ) as show_error:
            window._fail_erta_batch(1, RuntimeError("write failed"))

        show_error.assert_called_once_with("Batch prediction failed", "write failed")
        show_info.assert_not_called()
        self.assertIn("write failed", window.batch_result.value)
        self.assertEqual(
            window.statuses[-1],
            "Batch prediction failed: write failed",
        )

    def test_batch_ignores_prior_output_root_and_publishes_in_input_parent(self):
        window = self._window()
        window.predictor = _ErtaBatchPredictor()
        window.ad_calculator = type("AD", (), {"fitted": False})()
        window.model_path_var = _ErtaBatchVariable("")
        window.ad_ref_path_var = _ErtaBatchVariable("")
        reported = []
        window.update_batch_result_summary = (
            lambda _result, path, _graphs: reported.append(path)
        )
        window.generate_batch_graphs = lambda *_args, **_kwargs: []
        window.show_error = lambda _title, error: self.fail(str(error))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_directory = root / "selected-input-folder"
            prior_output_setting = root / "Documents" / "ER_Predictor" / "Exports"
            input_directory.mkdir()
            prior_output_setting.mkdir(parents=True)
            window.output_root = str(prior_output_setting)
            source = input_directory / "input.xlsx"
            pd.DataFrame({"CAS": ["50-00-0"], "SMILES": ["CCO"]}).to_excel(source, index=False)
            window.batch_input_var = _ErtaBatchVariable(str(source))

            def prepare(frame, **callbacks):
                callbacks["status_callback"](
                    "Fetching SMILES from PubChem: 1 / 1 (50-00-0)"
                )
                return frame

            window.prepare_batch_input = prepare
            with patch("gui.main_window.messagebox.showinfo"):
                window.batch_predict_clicked()
            self.assertEqual(Path(reported[0]).parent, input_directory.resolve())
            self.assertTrue(Path(reported[0]).exists())
            self.assertEqual(
                window.batch_destination_var.value,
                str(input_directory.resolve()),
            )
            self.assertEqual(list(prior_output_setting.glob("*.xlsx")), [])
            lookup_index = window.statuses.index(
                "Fetching SMILES from PubChem: 1 / 1 (50-00-0)"
            )
            self.assertEqual(
                window.statuses[lookup_index + 1],
                "Batch prediction started.",
            )


if __name__ == "__main__":
    unittest.main()
