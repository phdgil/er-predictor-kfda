"""Opt-in native-package acceptance driver used only by release verification."""
from __future__ import annotations

from copy import copy
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import time
import traceback

from openpyxl import Workbook, load_workbook

from core.contracts import (
    CLASSIFICATION_PARENT_POLICY_ID,
    CLASSIFICATION_OUTPUT_COLUMNS,
    ERBA_BINDING_CLASSIFICATION_EXCEL_CONTRACT_ID,
    ERBA_CLASSIFICATION_DIAGNOSTIC_COLUMNS,
    ERBA_CLASSIFICATION_METADATA_COLUMNS,
    ERBA_CLASSIFICATION_PREDICTION_COLUMNS,
    ERBA_CLASSIFICATION_PUBCHEM_COLUMNS,
    ERBASubtype,
    ERBATask,
    HISTORICAL_EXPOSURE_CAVEAT_COMPACT,
)

DISTRIBUTED_TEST_XLSX_SHA256 = (
    "5a1f569f8f6a5cd47bff67a189645c3f9461fcf07bd24f4e5b4f83197f3350aa"
)

BATCH_RECOGNITION_GATE_ID = "batch-recognition-contract-v1"
BATCH_RECOGNITION_CRITERIA = {
    "BRG-01-FEEDBACK-PLACEMENT": (
        "Aggregate progress is bound below Run batch, while per-row PubChem "
        "detail is bound to the endpoint status line below the result pane."
    ),
    "BRG-02-NEW-RUN-RESET": (
        "Starting a new run removes the prior Completed/saved presentation and "
        "locks all batch controls."
    ),
    "BRG-03-RESOLUTION-DETAIL": (
        "CAS resolution exposes aggregate progress and the same per-row "
        "PubChem detail format at both endpoints."
    ),
    "BRG-04-PUBLISHED-SUCCESS": (
        "A published workbook finishes as Completed and uses the blue "
        "showinfo dialog semantic with an explicit Not predicted count."
    ),
    "BRG-05-PARTIAL-UNAVAILABLE": (
        "A partially unavailable batch remains a successful publication and "
        "retains every unavailable row in both the UI and workbook."
    ),
    "BRG-06-ALL-UNAVAILABLE": (
        "An all-unavailable batch remains a successful publication and "
        "reports every unavailable row in both the UI and workbook."
    ),
    "BRG-07-NO-OUTPUT-FAILURE": (
        "An actual run/save failure publishes no workbook and uses Failed plus "
        "the red showerror dialog semantic."
    ),
    "BRG-08-OWN-WINDOW-EVIDENCE": (
        "Running and completion screenshots target only the application Tk "
        "window when native own-window capture is supported."
    ),
    "BRG-09-STUDY-SCOPE": (
        "The receipt identifies this as an automated recognition-contract "
        "check, not a real human usability study."
    ),
}
BATCH_UNAVAILABLE_LABEL = "Not predicted"
ERALPHA_AD_COLUMNS = (
    "AD",
    "AD_MeanDistance",
    "AD_DistanceThreshold",
    "AD_Distance_InDomain",
    "AD_SimilarityMax",
    "AD_SimilarityThreshold",
    "AD_Similarity_InDomain",
    "AD_PC1",
    "AD_PC2",
)
ERALPHA_PRIMARY_TRUSTED_COLUMNS = (
    *ERBA_CLASSIFICATION_PREDICTION_COLUMNS,
    *ERALPHA_AD_COLUMNS,
    *ERBA_CLASSIFICATION_PUBCHEM_COLUMNS,
)
ERALPHA_WORKBOOK_SHEET_ORDER = (
    "Predictions",
    "Guide",
    "Diagnostics",
    "Input",
    "Metadata",
)
ERALPHA_PRIMARY_FORBIDDEN_COLUMNS = tuple(
    dict.fromkeys(
        (
            *CLASSIFICATION_OUTPUT_COLUMNS,
            *ERBA_CLASSIFICATION_METADATA_COLUMNS,
            *(
                column
                for column in ERBA_CLASSIFICATION_DIAGNOSTIC_COLUMNS
                if column not in {"Row_ID", "CAS"}
            ),
        )
    )
)


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normalized_example_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().casefold())


def _read_shared_example_workbook(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        normalized = {
            _normalized_example_header(value): index
            for index, value in enumerate(headers)
        }
        cas_index = next(
            (
                normalized[alias]
                for alias in (
                    "cas",
                    "casno",
                    "casrn",
                    "carsrn",
                    "casnumber",
                )
                if alias in normalized
            ),
            None,
        )
        if cas_index is None:
            raise AssertionError(
                f"shared example workbook lacks a recognized CAS column: {headers}"
            )
        smiles_index = next(
            (
                normalized[alias]
                for alias in (
                    "smiles",
                    "canonicalsmiles",
                    "isomericsmiles",
                )
                if alias in normalized
            ),
            None,
        )
        values = list(sheet.iter_rows(min_row=2, values_only=True))
        cas_values = [
            str(row[cas_index] or "").strip()
            for row in values
        ]
        smiles_values = [
            (
                str(row[smiles_index] or "").strip()
                if smiles_index is not None
                else ""
            )
            for row in values
        ]
        return {
            "sheet": sheet.title,
            "headers": headers,
            "row_count": len(values),
            "cas_header": headers[cas_index],
            "cas_values": cas_values,
            "smiles_header": (
                headers[smiles_index] if smiles_index is not None else None
            ),
            "smiles_values": smiles_values,
        }
    finally:
        workbook.close()


def _openpyxl_cell_kind(cell) -> str:
    if cell.row == 1 and cell.value is not None:
        return "header"
    if cell.value is None:
        return "blank"
    if isinstance(cell.value, bool):
        return "boolean"
    if cell.is_date:
        return "date"
    if isinstance(cell.value, (int, float)):
        return "number"
    return "text"


def _openpyxl_format_signature(cell) -> dict:
    """Return the complete value-independent openpyxl cell-format signature."""
    return {
        "font": repr(copy(cell.font)),
        "fill": repr(copy(cell.fill)),
        "border": repr(copy(cell.border)),
        "alignment": repr(copy(cell.alignment)),
        "number_format": str(cell.number_format),
        "protection": repr(copy(cell.protection)),
        "quote_prefix": bool(cell.quotePrefix),
        "pivot_button": bool(cell.pivotButton),
    }


def _worksheet_format_inventory(sheet) -> dict:
    inventories: dict[str, dict[str, dict]] = {}
    counts: dict[str, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            kind = _openpyxl_cell_kind(cell)
            signature = _openpyxl_format_signature(cell)
            canonical = json.dumps(signature, sort_keys=True)
            inventories.setdefault(kind, {}).setdefault(
                canonical,
                {
                    "signature": signature,
                    "example_coordinate": cell.coordinate,
                },
            )
            counts[kind] = counts.get(kind, 0) + 1
    return {
        kind: {
            "cell_count": counts[kind],
            "signatures": [
                receipt
                for _, receipt in sorted(inventories[kind].items())
            ],
        }
        for kind in sorted(inventories)
    }


def compare_plain_primary_workbook_formatting(
    erta_workbook: str | Path,
    eralpha_workbook: str | Path,
) -> dict:
    """Compare exact cell formatting by kind, never endpoint data semantics."""
    erta = load_workbook(erta_workbook, read_only=False, data_only=False)
    eralpha = load_workbook(eralpha_workbook, read_only=False, data_only=False)
    try:
        erta_sheet = erta.active
        if "Predictions" not in eralpha.sheetnames:
            raise AssertionError("ERalpha workbook lacks the Predictions sheet")
        eralpha_sheet = eralpha["Predictions"]
        inventories = {
            "erta": _worksheet_format_inventory(erta_sheet),
            "eralpha": _worksheet_format_inventory(eralpha_sheet),
        }
        compared_kinds = sorted(
            set(inventories["erta"]) & set(inventories["eralpha"])
        )
        required_kinds = {"header", "text", "number", "boolean"}
        missing = required_kinds - set(compared_kinds)
        if missing:
            raise AssertionError(
                "primary worksheet formatting comparison lacks common cell "
                f"kinds: {sorted(missing)}"
            )
        non_uniform = {
            f"{endpoint}/{kind}": len(
                inventories[endpoint][kind]["signatures"]
            )
            for endpoint in ("erta", "eralpha")
            for kind in required_kinds
            if len(inventories[endpoint][kind]["signatures"]) != 1
        }
        if non_uniform:
            raise AssertionError(
                "plain primary worksheets have non-uniform formatting within "
                f"a cell kind: {non_uniform}"
            )
        mismatches = {}
        for kind in compared_kinds:
            signatures = {
                endpoint: [
                    item["signature"]
                    for item in inventories[endpoint][kind]["signatures"]
                ]
                for endpoint in ("erta", "eralpha")
            }
            if signatures["erta"] != signatures["eralpha"]:
                mismatches[kind] = signatures
        if mismatches:
            raise AssertionError(
                "ERalpha Predictions formatting differs from plain ERTA cells "
                "of the same kind: "
                + json.dumps(mismatches, sort_keys=True)
            )
        return {
            "comparison": (
                "exact openpyxl font/fill/border/alignment/number-format/"
                "protection/quote-prefix/pivot-button signature by cell kind"
            ),
            "data_semantics_compared": False,
            "cell_kind_policy": (
                "header row, then blank/boolean/date/number/text by openpyxl "
                "cell value type; values and column meanings are ignored"
            ),
            "erta_sheet": erta_sheet.title,
            "eralpha_sheet": eralpha_sheet.title,
            "compared_kinds": compared_kinds,
            "required_kinds": sorted(required_kinds),
            "inventories": inventories,
            "uniform_within_required_kinds": True,
            "exact_match": True,
        }
    finally:
        erta.close()
        eralpha.close()


def _widget_geometry_receipt(widget, endpoint_tab) -> dict:
    mapped = bool(widget.winfo_ismapped())
    receipt = {
        "mapped": mapped,
        "class": str(widget.winfo_class()),
        "manager": str(widget.winfo_manager()),
    }
    try:
        receipt["state"] = str(widget.cget("state"))
    except Exception:
        pass
    # Tcl accepts abbreviated option names.  Asking an Entry for "text" can
    # therefore return its endpoint-local -textvariable name (PY_VAR...), which
    # is widget plumbing rather than visible content.  Compare static button
    # captions, while input/status values are checked through their variables
    # by the focused semantic inspections.
    if receipt["class"] in {"Button", "TButton"}:
        receipt["text"] = str(widget.cget("text"))
    if mapped:
        width = int(widget.winfo_width())
        height = int(widget.winfo_height())
        if width <= 1 or height <= 1:
            raise AssertionError(
                f"mapped {widget!s} has an unresolved {width}x{height} bounding box"
            )
        receipt["box"] = {
            "x": int(widget.winfo_rootx() - endpoint_tab.winfo_rootx()),
            "y": int(widget.winfo_rooty() - endpoint_tab.winfo_rooty()),
            "width": width,
            "height": height,
        }
    return receipt


def _batch_recognition_widget_receipt(owner, controls: dict) -> dict:
    """Inspect the visible batch feedback reading order and its real Tk bindings."""
    endpoint_tab = controls["endpoint_tab"]
    progress = controls["batch_progress"]
    progress_frame = progress.master
    progress_variable = str(owner.batch_progress_var)
    progress_labels = []
    for child in progress_frame.winfo_children():
        try:
            if str(child.cget("textvariable")) == progress_variable:
                progress_labels.append(child)
        except Exception:
            continue
    if len(progress_labels) != 1:
        raise AssertionError(
            "batch aggregate progress must have exactly one visible label bound "
            f"to {progress_variable}; found {len(progress_labels)}"
        )
    progress_label = progress_labels[0]
    detail_variable = getattr(owner, "status_var", None)
    if detail_variable is None:
        detail_variable = owner.batch_status_var

    run_receipt = _widget_geometry_receipt(
        controls["batch_predict_button"],
        endpoint_tab,
    )
    progress_receipt = _widget_geometry_receipt(progress, endpoint_tab)
    progress_label_receipt = _widget_geometry_receipt(
        progress_label,
        endpoint_tab,
    )
    result_receipt = _widget_geometry_receipt(
        controls["batch_result"],
        endpoint_tab,
    )
    detail_receipt = _widget_geometry_receipt(
        controls["batch_status"],
        endpoint_tab,
    )
    mapped_receipts = (
        run_receipt,
        progress_receipt,
        progress_label_receipt,
        result_receipt,
        detail_receipt,
    )
    if not all(receipt["mapped"] for receipt in mapped_receipts):
        raise AssertionError(
            "batch recognition widgets must all be mapped while the batch "
            "surface is selected"
        )

    run_command = str(controls["batch_predict_button"].cget("command"))
    actual_progress_value_variable = str(progress.cget("variable"))
    actual_progress_text_variable = str(progress_label.cget("textvariable"))
    actual_detail_variable = str(
        controls["batch_status"].cget("textvariable")
    )
    if (
        "batch_predict_clicked" not in run_command
        or actual_progress_value_variable != str(owner.batch_progress_value)
        or actual_progress_text_variable != progress_variable
        or actual_detail_variable != str(detail_variable)
    ):
        raise AssertionError(
            "batch recognition widget binding drift: "
            + json.dumps(
                {
                    "run_command": run_command,
                    "progress_value_variable": {
                        "actual": actual_progress_value_variable,
                        "expected": str(owner.batch_progress_value),
                    },
                    "progress_text_variable": {
                        "actual": actual_progress_text_variable,
                        "expected": progress_variable,
                    },
                    "detail_text_variable": {
                        "actual": actual_detail_variable,
                        "expected": str(detail_variable),
                    },
                },
                sort_keys=True,
            )
        )

    run_box = run_receipt["box"]
    progress_box = progress_receipt["box"]
    progress_label_box = progress_label_receipt["box"]
    result_box = result_receipt["box"]
    detail_box = detail_receipt["box"]
    run_bottom = run_box["y"] + run_box["height"]
    progress_bottom = progress_box["y"] + progress_box["height"]
    result_bottom = result_box["y"] + result_box["height"]
    label_bottom = progress_label_box["y"] + progress_label_box["height"]
    progress_vertical_overlap = (
        progress_label_box["y"] < progress_bottom
        and label_bottom > progress_box["y"]
    )
    if not (
        run_bottom <= progress_box["y"]
        and progress_vertical_overlap
        and progress_label_box["x"]
        >= progress_box["x"] + progress_box["width"]
        and progress_bottom <= result_box["y"]
        and result_bottom <= detail_box["y"]
    ):
        raise AssertionError(
            "batch feedback is not in Run, aggregate progress, result, "
            "per-row detail reading order: "
            + json.dumps(
                {
                    "run": run_box,
                    "aggregate_progress": progress_box,
                    "aggregate_progress_label": progress_label_box,
                    "result": result_box,
                    "per_row_detail": detail_box,
                },
                sort_keys=True,
            )
        )

    return {
        "criterion_id": "BRG-01-FEEDBACK-PLACEMENT",
        "passed": True,
        "coordinate_space": "relative to the selected endpoint_tab",
        "reading_order": [
            "run_batch",
            "aggregate_progress",
            "batch_result",
            "per_row_detail",
        ],
        "bindings": {
            "run_command": run_command,
            "run_callback": "batch_predict_clicked",
            "progress_value_variable": actual_progress_value_variable,
            "progress_text_variable": actual_progress_text_variable,
            "per_row_detail_variable": actual_detail_variable,
        },
        "boxes": {
            "run_batch": run_box,
            "aggregate_progress": progress_box,
            "aggregate_progress_label": progress_label_box,
            "batch_result": result_box,
            "per_row_detail": detail_box,
        },
    }


def _capture_own_tk_window(window, destination: Path) -> dict:
    """Capture only this Tk HWND via PrintWindow; never fall back to the desktop."""
    if os.name != "nt":
        return {
            "supported": False,
            "method": "PrintWindow(PW_RENDERFULLCONTENT)",
            "reason": "own-window native capture is only available on Windows",
        }
    try:
        import ctypes

        from PIL import Image, ImageStat
        import win32gui
        import win32ui

        handle = int(window.winfo_id())
        handle = int(win32gui.GetAncestor(handle, 2) or handle)
        left, top, right, bottom = win32gui.GetWindowRect(handle)
        width, height = right - left, bottom - top
        if width <= 1 or height <= 1:
            raise RuntimeError(f"invalid own-window bounds {width}x{height}")

        window_dc = win32gui.GetWindowDC(handle)
        source_dc = win32ui.CreateDCFromHandle(window_dc)
        memory_dc = source_dc.CreateCompatibleDC()
        bitmap = win32ui.CreateBitmap()
        bitmap.CreateCompatibleBitmap(source_dc, width, height)
        memory_dc.SelectObject(bitmap)
        try:
            if not ctypes.windll.user32.PrintWindow(
                handle,
                memory_dc.GetSafeHdc(),
                2,
            ):
                raise RuntimeError("PrintWindow failed")
            image = Image.frombuffer(
                "RGB",
                (width, height),
                bitmap.GetBitmapBits(True),
                "raw",
                "BGRX",
                0,
                1,
            ).copy()
        finally:
            win32gui.DeleteObject(bitmap.GetHandle())
            memory_dc.DeleteDC()
            source_dc.DeleteDC()
            win32gui.ReleaseDC(handle, window_dc)

        destination.parent.mkdir(parents=True, exist_ok=True)
        image.save(destination)
        extrema = ImageStat.Stat(image).extrema
        if not any(low != high for low, high in extrema):
            raise RuntimeError("PrintWindow returned a uniform image")
        return {
            "supported": True,
            "method": "PrintWindow(PW_RENDERFULLCONTENT)",
            "path": str(destination),
            "sha256": _sha256_path(destination),
            "size": [width, height],
            "rgb_extrema": [list(channel) for channel in extrema],
            "target_hwnd": handle,
        }
    except Exception as error:
        try:
            destination.unlink(missing_ok=True)
        except OSError:
            pass
        return {
            "supported": False,
            "method": "PrintWindow(PW_RENDERFULLCONTENT)",
            "reason": f"{type(error).__name__}: {error}",
        }


def _alternate_window_extent(
    current: int,
    minimum: int,
    maximum: int,
    delta: int,
) -> int:
    smaller = current - delta
    if smaller >= minimum:
        return smaller
    larger = current + delta
    if maximum <= 0 or larger <= maximum:
        return larger
    if minimum != current:
        return minimum
    if maximum != current:
        return maximum
    raise AssertionError(f"Tk window extent {current} cannot be resized")


def _layout_box(widget, root) -> dict:
    mapped = bool(widget.winfo_ismapped())
    result = {"mapped": mapped}
    if mapped:
        result["box"] = {
            "x": int(widget.winfo_rootx() - root.winfo_rootx()),
            "y": int(widget.winfo_rooty() - root.winfo_rooty()),
            "width": int(widget.winfo_width()),
            "height": int(widget.winfo_height()),
        }
    return result


def _settle_tk_layout(
    app,
    watched_widgets: dict,
    *,
    expected_root_size: tuple[int, int] | None,
    context: str,
    max_cycles: int = 12,
) -> dict:
    """Drain real Tk events until root and watched boxes repeat unchanged."""
    previous = None
    stable_samples = 0
    observations = []
    for cycle in range(1, max_cycles + 1):
        idle_reached = []
        app.after_idle(lambda: idle_reached.append(True))
        app.update()
        app.update_idletasks()
        root_size = (int(app.winfo_width()), int(app.winfo_height()))
        boxes = {
            name: _layout_box(widget, app)
            for name, widget in sorted(watched_widgets.items())
        }
        snapshot = {
            "root_size": list(root_size),
            "idle_reached": bool(idle_reached),
            "boxes": boxes,
        }
        observations.append(snapshot)
        signature = json.dumps(snapshot, sort_keys=True)
        root_matches = (
            expected_root_size is None or root_size == expected_root_size
        )
        if root_matches and idle_reached and signature == previous:
            stable_samples += 1
        elif root_matches and idle_reached:
            stable_samples = 1
        else:
            stable_samples = 0
        if stable_samples >= 2:
            return {
                "context": context,
                "cycles": cycle,
                "stable_samples": stable_samples,
                **snapshot,
            }
        previous = signature

    requested = (
        list(expected_root_size) if expected_root_size is not None else None
    )
    raise AssertionError(
        f"Tk layout did not settle for {context} after {max_cycles} cycles: "
        + json.dumps(
            {
                "requested_root_size": requested,
                "last_observations": observations[-3:],
            },
            sort_keys=True,
        )
    )


def inspect_endpoint_geometry_parity(
    app,
    screenshot_directory: str | Path | None = None,
) -> dict:
    """Inspect the real mapped ERTA/ERalpha controls at identical Tk geometries."""
    endpoints = {
        "erta": app,
        "eralpha": app.eralpha_tab,
    }
    maps = {
        name: dict(getattr(owner, "parity_widgets", {}))
        for name, owner in endpoints.items()
    }
    if not maps["erta"] or not maps["eralpha"]:
        raise AssertionError("both endpoints must expose a non-empty parity_widgets map")
    if maps["erta"].keys() != maps["eralpha"].keys():
        raise AssertionError(
            "endpoint parity widget keys differ: "
            f"ERTA={sorted(maps['erta'])}, ERalpha={sorted(maps['eralpha'])}"
        )
    required_keys = {
        "endpoint_tab",
        "options_button",
        "options_frame",
        "model_entry",
        "model_browse_button",
        "model_reload_button",
        "ad_entry",
        "ad_browse_button",
        "ad_reload_button",
        "mode_notebook",
        "single_tab",
        "batch_tab",
        "cas_label",
        "cas_entry",
        "pubchem_button",
        "smiles_label",
        "smiles_entry",
        "example_button",
        "single_predict_button",
        "draw_structure_button",
        "batch_input_entry",
        "batch_browse_button",
        "batch_destination_entry",
        "batch_example_button",
        "batch_predict_button",
        "batch_progress",
        "batch_status",
        "batch_result",
    }
    missing = required_keys - maps["erta"].keys()
    if missing:
        raise AssertionError(
            f"endpoint parity map is missing required controls: {sorted(missing)}"
        )
    containers = {
        name: {
            "header_box": controls["options_button"].master,
            "single_input_box": controls["cas_entry"].master,
            "batch_input_box": controls["batch_input_entry"].master,
            "batch_progress_box": controls["batch_progress"].master,
            "batch_result_box": controls["batch_result"].master,
        }
        for name, controls in maps.items()
    }

    app.update()
    app.update_idletasks()
    original_geometry = app.geometry()
    original_state = str(app.state())
    original_root_size = (int(app.winfo_width()), int(app.winfo_height()))
    original_endpoint = app.notebook.select()
    original_modes = {
        name: controls["mode_notebook"].select()
        for name, controls in maps.items()
    }
    original_options = {
        name: bool(owner.options_visible)
        for name, owner in endpoints.items()
    }
    screenshot_root = (
        Path(screenshot_directory)
        if screenshot_directory is not None
        else None
    )
    screenshot_available = screenshot_root is not None
    screenshots = []
    states = []
    recognition_widgets = {}
    toggle_invocations = {"erta": 0, "eralpha": 0}
    window_size_settling = {}

    def select_surface(
        name: str,
        mode: str,
        options_open: bool,
        expected_root_size: tuple[int, int],
    ) -> dict:
        nonlocal toggle_invocations
        owner = endpoints[name]
        controls = maps[name]
        watched = {
            **{
                f"control:{key}": widget
                for key, widget in controls.items()
            },
            **{
                f"container:{key}": widget
                for key, widget in containers[name].items()
            },
        }
        app.notebook.select(controls["endpoint_tab"])
        controls["mode_notebook"].select(controls[f"{mode}_tab"])
        _settle_tk_layout(
            app,
            watched,
            expected_root_size=expected_root_size,
            context=f"{name}/{mode}/before Options transition",
        )
        actual_open = bool(controls["options_frame"].winfo_ismapped())
        if bool(owner.options_visible) != actual_open:
            raise AssertionError(
                f"{name} Options state flag and mapped state disagree"
            )
        if actual_open != options_open:
            controls["options_button"].invoke()
            toggle_invocations[name] += 1
        settled = _settle_tk_layout(
            app,
            watched,
            expected_root_size=expected_root_size,
            context=(
                f"{name}/{mode}/Options "
                f"{'open' if options_open else 'closed'}"
            ),
        )
        if (
            bool(owner.options_visible) != options_open
            or bool(controls["options_frame"].winfo_ismapped()) != options_open
            or bool(controls["options_frame"].grid_info()) != options_open
        ):
            raise AssertionError(
                f"{name} Options button did not produce the requested "
                f"{'open' if options_open else 'closed'} state"
            )
        selected = controls["mode_notebook"].select()
        if selected != str(controls[f"{mode}_tab"]):
            raise AssertionError(f"{name} did not select its {mode} tab")
        return settled

    try:
        if original_state != "normal":
            app.state("normal")
        initial_settling = _settle_tk_layout(
            app,
            {"top_notebook": app.notebook},
            expected_root_size=None,
            context="initial root geometry",
        )
        initial_size = (int(app.winfo_width()), int(app.winfo_height()))
        minimum = tuple(int(value) for value in app.minsize())
        maximum = tuple(int(value) for value in app.maxsize())
        resized_size = (
            _alternate_window_extent(
                initial_size[0],
                minimum[0],
                maximum[0],
                73,
            ),
            _alternate_window_extent(
                initial_size[1],
                minimum[1],
                maximum[1],
                47,
            ),
        )
        if resized_size == initial_size:
            raise AssertionError("the alternate Tk geometry did not resize the window")

        for size_name, requested_size in (
            ("initial", initial_size),
            ("resized", resized_size),
        ):
            app.geometry(f"{requested_size[0]}x{requested_size[1]}")
            size_settling = _settle_tk_layout(
                app,
                {"top_notebook": app.notebook},
                expected_root_size=requested_size,
                context=f"{size_name} requested root geometry",
            )
            window_size_settling[size_name] = size_settling
            actual_size = tuple(size_settling["root_size"])
            if size_name == "resized" and actual_size == initial_size:
                raise AssertionError("Tk ignored the requested resize")

            for mode in ("single", "batch"):
                for options_open in (False, True):
                    endpoint_receipts = {}
                    container_receipts = {}
                    endpoint_sizes = {}
                    endpoint_settling = {}
                    for endpoint_name in endpoints:
                        endpoint_settling[endpoint_name] = select_surface(
                            endpoint_name,
                            mode,
                            options_open,
                            requested_size,
                        )
                        endpoint_sizes[endpoint_name] = [
                            int(app.winfo_width()),
                            int(app.winfo_height()),
                        ]
                        endpoint_tab = maps[endpoint_name]["endpoint_tab"]
                        endpoint_receipts[endpoint_name] = {
                            key: _widget_geometry_receipt(widget, endpoint_tab)
                            for key, widget in maps[endpoint_name].items()
                        }
                        container_receipts[endpoint_name] = {
                            key: _widget_geometry_receipt(widget, endpoint_tab)
                            for key, widget in containers[endpoint_name].items()
                        }
                        if (
                            size_name == "initial"
                            and mode == "batch"
                            and not options_open
                        ):
                            recognition_widgets[endpoint_name] = (
                                _batch_recognition_widget_receipt(
                                    endpoints[endpoint_name],
                                    maps[endpoint_name],
                                )
                            )
                        if screenshot_available and screenshot_root is not None:
                            option_name = "options-open" if options_open else "options-closed"
                            screenshot = _capture_own_tk_window(
                                app,
                                screenshot_root
                                / (
                                    f"{size_name}-{mode}-{option_name}-"
                                    f"{endpoint_name}.png"
                                ),
                            )
                            screenshot.update(
                                {
                                    "window_geometry": size_name,
                                    "mode": mode,
                                    "options_open": options_open,
                                    "endpoint": endpoint_name,
                                }
                            )
                            screenshots.append(screenshot)
                            if not screenshot["supported"]:
                                screenshot_available = False

                    if endpoint_sizes["erta"] != endpoint_sizes["eralpha"]:
                        raise AssertionError(
                            f"root geometry changed between endpoints: {endpoint_sizes}"
                        )
                    if endpoint_receipts["erta"] != endpoint_receipts["eralpha"]:
                        different = [
                            key
                            for key in maps["erta"]
                            if endpoint_receipts["erta"][key]
                            != endpoint_receipts["eralpha"][key]
                        ]
                        raise AssertionError(
                            "actual ERTA/ERalpha control geometry differs for "
                            f"{size_name}/{mode}/Options "
                            f"{'open' if options_open else 'closed'}: "
                            + json.dumps(
                                {
                                    key: {
                                        endpoint: endpoint_receipts[endpoint][key]
                                        for endpoint in endpoints
                                    }
                                    for key in different
                                },
                                sort_keys=True,
                            )
                        )
                    if container_receipts["erta"] != container_receipts["eralpha"]:
                        different = [
                            key
                            for key in containers["erta"]
                            if container_receipts["erta"][key]
                            != container_receipts["eralpha"][key]
                        ]
                        raise AssertionError(
                            "actual ERTA/ERalpha container boxes differ for "
                            f"{size_name}/{mode}/Options "
                            f"{'open' if options_open else 'closed'}: "
                            + json.dumps(
                                {
                                    key: {
                                        endpoint: container_receipts[endpoint][key]
                                        for endpoint in endpoints
                                    }
                                    for key in different
                                },
                                sort_keys=True,
                            )
                        )
                    states.append(
                        {
                            "window_geometry": size_name,
                            "requested_size": list(requested_size),
                            "actual_size": endpoint_sizes["erta"],
                            "mode": mode,
                            "options_open": options_open,
                            "exact_match": True,
                            "controls": endpoint_receipts,
                            "containers": container_receipts,
                            "settling": endpoint_settling,
                        }
                    )
    finally:
        for name, owner in endpoints.items():
            controls = maps[name]
            app.notebook.select(controls["endpoint_tab"])
            controls["mode_notebook"].select(original_modes[name])
            app.update()
            app.update_idletasks()
            if bool(owner.options_visible) != original_options[name]:
                controls["options_button"].invoke()
                app.update()
                app.update_idletasks()
        app.notebook.select(original_endpoint)
        app.geometry(original_geometry)
        if original_state != "normal":
            app.state(original_state)
        if original_state == "normal":
            _settle_tk_layout(
                app,
                {"top_notebook": app.notebook},
                expected_root_size=original_root_size,
                context="restored root geometry",
            )
        else:
            app.update_idletasks()

    expected_state_count = 2 * 2 * 2
    if len(states) != expected_state_count:
        raise AssertionError(
            f"expected {expected_state_count} geometry comparisons, found {len(states)}"
        )
    if set(recognition_widgets) != {"erta", "eralpha"}:
        raise AssertionError(
            "batch recognition widget receipts are incomplete: "
            f"{sorted(recognition_widgets)}"
        )
    comparable_recognition_widgets = {
        endpoint: {
            "reading_order": receipt["reading_order"],
            "boxes": receipt["boxes"],
            "run_callback": receipt["bindings"]["run_callback"],
        }
        for endpoint, receipt in recognition_widgets.items()
    }
    if (
        comparable_recognition_widgets["erta"]
        != comparable_recognition_widgets["eralpha"]
    ):
        raise AssertionError(
            "ERTA/ERalpha batch feedback placement differs: "
            + json.dumps(comparable_recognition_widgets, sort_keys=True)
        )
    batch_recognition_widget_contract = {
        "criterion_id": "BRG-01-FEEDBACK-PLACEMENT",
        "passed": True,
        "exact_layout_match": True,
        "endpoints": recognition_widgets,
    }
    return {
        "coordinate_space": "relative to each selected endpoint_tab",
        "compared_widget_keys": sorted(maps["erta"]),
        "compared_container_keys": sorted(containers["erta"]),
        "state_count": len(states),
        "toggle_invocations": toggle_invocations,
        "initial_settling": initial_settling,
        "window_size_settling": window_size_settling,
        "states": states,
        "batch_recognition_widgets": batch_recognition_widget_contract,
        "screenshots": screenshots,
        "screenshot_policy": (
            "Own Tk HWND only via PrintWindow; no desktop-capture fallback."
        ),
    }


def inspect_shared_example_parity(
    app,
    isolated_profile_root: str | Path | None = None,
) -> dict:
    """Verify both endpoints share the writable, byte-faithful test.xlsx."""
    from core.paths import (
        resolve_shared_example_input,
        validate_mutable_directory,
    )

    endpoints = {
        "erta": app,
        "eralpha": app.eralpha_tab,
    }
    expected = {
        name: {
            "cas": str(owner.example_cas),
            "smiles": str(owner.example_smiles),
            "workbook": Path(owner.example_workbook).expanduser().resolve(
                strict=True
            ),
            "batch_workbook": Path(owner.batch_input_var.get()).expanduser().resolve(
                strict=True
            ),
            "batch_display": owner.batch_input_display_var.get(),
            "batch_destination": owner.batch_destination_var.get(),
        }
        for name, owner in endpoints.items()
    }
    for name, owner in endpoints.items():
        if not expected[name]["cas"]:
            raise AssertionError(f"{name} preset CAS is blank")
        if (
            owner.cas_var.get() != expected[name]["cas"]
            or owner.smiles_var.get() != expected[name]["smiles"]
        ):
            raise AssertionError(f"{name} did not start with its preset example")
        if expected[name]["workbook"] != expected[name]["batch_workbook"]:
            raise AssertionError(
                f"{name} batch example path differs from its preset workbook"
            )
        if expected[name]["batch_display"] != expected[name]["workbook"].name:
            raise AssertionError(
                f"{name} batch example display does not name the preset workbook"
            )
        if expected[name]["batch_destination"] != str(
            expected[name]["workbook"].parent
        ):
            raise AssertionError(
                f"{name} batch destination does not name the writable preset parent"
            )

    if (
        expected["erta"]["cas"],
        expected["erta"]["smiles"],
    ) != (
        expected["eralpha"]["cas"],
        expected["eralpha"]["smiles"],
    ):
        raise AssertionError("ERTA and ERalpha preset CAS/SMILES values differ")
    if expected["erta"]["workbook"] != expected["eralpha"]["workbook"]:
        raise AssertionError("ERTA and ERalpha batch example paths differ")
    if expected["erta"]["batch_display"] != expected["eralpha"]["batch_display"]:
        raise AssertionError("ERTA and ERalpha batch example displays differ")
    if (
        expected["erta"]["batch_destination"]
        != expected["eralpha"]["batch_destination"]
    ):
        raise AssertionError("ERTA and ERalpha batch destinations differ")
    if expected["erta"]["workbook"].name.casefold() != "test.xlsx":
        raise AssertionError("shared batch default is not named test.xlsx")

    erta_bytes = expected["erta"]["workbook"].read_bytes()
    eralpha_bytes = expected["eralpha"]["workbook"].read_bytes()
    if erta_bytes != eralpha_bytes:
        raise AssertionError("ERTA and ERalpha batch example workbook bytes differ")

    default_metadata = _read_shared_example_workbook(
        expected["erta"]["workbook"]
    )
    if not default_metadata["cas_values"]:
        raise AssertionError("shared example workbook has no data rows")
    source_cas = default_metadata["cas_values"][0]
    source_smiles = default_metadata["smiles_values"][0]
    if (source_cas, source_smiles) != (
        expected["erta"]["cas"],
        expected["erta"]["smiles"],
    ):
        raise AssertionError(
            "preset CAS/SMILES does not equal the first shared workbook example"
        )

    default_hash_before = hashlib.sha256(erta_bytes).hexdigest()
    validate_mutable_directory(
        expected["erta"]["workbook"].parent,
        forbidden_roots=(app.project_root, app.install_root),
    )
    if not os.access(expected["erta"]["workbook"], os.W_OK):
        raise AssertionError("resolved shared test.xlsx is not writable")
    if hashlib.sha256(
        expected["erta"]["workbook"].read_bytes()
    ).hexdigest() != default_hash_before:
        raise AssertionError(
            "writability inspection changed the shared default workbook"
        )

    bundled_workbook = (
        Path(app.project_root).expanduser().resolve(strict=True)
        / "templates"
        / "test.xlsx"
    )
    if not bundled_workbook.is_file():
        raise AssertionError(
            f"bundled shared example is missing: {bundled_workbook}"
        )
    bundled_bytes = bundled_workbook.read_bytes()
    bundled_sha256 = hashlib.sha256(bundled_bytes).hexdigest()
    if bundled_sha256 != DISTRIBUTED_TEST_XLSX_SHA256:
        raise AssertionError(
            "bundled test.xlsx bytes differ from the supplied distribution source"
        )
    bundled_metadata = _read_shared_example_workbook(bundled_workbook)
    if (
        bundled_metadata["headers"] != ["CARSRN"]
        or bundled_metadata["cas_header"] != "CARSRN"
        or bundled_metadata["smiles_header"] is not None
        or bundled_metadata["row_count"] != 25
        or not all(bundled_metadata["cas_values"])
    ):
        raise AssertionError(
            "bundled test.xlsx does not preserve the supplied CARSRN/25-row contract"
        )

    fresh_copy = None
    if isolated_profile_root is not None:
        qa_profile = Path(isolated_profile_root).expanduser().resolve(
            strict=False
        )
        requested_fresh_copy = (
            qa_profile
            / "Documents"
            / "ER_Predictor"
            / "Examples"
            / "test.xlsx"
        )
        if qa_profile.exists():
            raise AssertionError(
                f"fresh shared-example QA profile already exists: {qa_profile}"
            )
        original_environment = {
            name: os.environ.get(name)
            for name in (
                "USERPROFILE",
                "LOCALAPPDATA",
                "ER_PREDICTOR_PORTABLE",
            )
        }
        had_frozen = hasattr(os.sys, "frozen")
        original_frozen = getattr(os.sys, "frozen", None)
        try:
            os.environ["USERPROFILE"] = str(qa_profile)
            os.environ["LOCALAPPDATA"] = str(qa_profile / "LocalAppData")
            os.environ["ER_PREDICTOR_PORTABLE"] = "0"
            setattr(os.sys, "frozen", False)
            fresh_copy = resolve_shared_example_input()
        finally:
            if had_frozen:
                setattr(os.sys, "frozen", original_frozen)
            else:
                delattr(os.sys, "frozen")
            for name, value in original_environment.items():
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value
        if fresh_copy != requested_fresh_copy or not fresh_copy.is_file():
            raise AssertionError(
                "shared-example resolver did not create the isolated first-run copy"
            )
        if fresh_copy.read_bytes() != bundled_bytes:
            raise AssertionError(
                "newly created shared default does not byte-match bundled test.xlsx"
            )
        validate_mutable_directory(
            fresh_copy.parent,
            forbidden_roots=(app.project_root, app.install_root),
        )
        if not os.access(fresh_copy, os.W_OK):
            raise AssertionError("newly created shared-example QA copy is not writable")

    original_endpoint = app.notebook.select()
    original_modes = {
        name: owner.parity_widgets["mode_notebook"].select()
        for name, owner in endpoints.items()
    }
    callback_receipts = []
    try:
        for name, owner in endpoints.items():
            other_name = "eralpha" if name == "erta" else "erta"
            other = endpoints[other_name]
            other_values = (
                other.cas_var.get(),
                other.smiles_var.get(),
                other.batch_input_var.get(),
                other.batch_input_display_var.get(),
                other.batch_destination_var.get(),
            )
            owner.cas_var.set(f"{name}-changed-cas")
            owner.smiles_var.set(f"{name}-changed-smiles")
            owner.batch_input_var.set(
                str(expected[name]["workbook"].parent / f"{name}-changed.xlsx")
            )
            owner.batch_input_display_var.set(f"{name}-changed.xlsx")
            owner.batch_destination_var.set(f"{name}-changed-destination")
            app.notebook.select(owner.parity_widgets["endpoint_tab"])
            owner.parity_widgets["mode_notebook"].select(
                owner.parity_widgets["single_tab"]
            )
            app.update_idletasks()
            owner.parity_widgets["example_button"].invoke()
            app.update_idletasks()
            restored = (owner.cas_var.get(), owner.smiles_var.get())
            expected_values = (expected[name]["cas"], expected[name]["smiles"])
            restored_batch = (
                Path(owner.batch_input_var.get()).expanduser().resolve(
                    strict=True
                ),
                owner.batch_input_display_var.get(),
                owner.batch_destination_var.get(),
            )
            expected_batch = (
                expected[name]["workbook"],
                expected[name]["workbook"].name,
                str(expected[name]["workbook"].parent),
            )
            if restored != expected_values:
                raise AssertionError(
                    f"{name} Example input button did not restore the preset"
                )
            if restored_batch != expected_batch:
                raise AssertionError(
                    f"{name} Example input button did not restore the batch default"
                )
            if (
                other.cas_var.get(),
                other.smiles_var.get(),
                other.batch_input_var.get(),
                other.batch_input_display_var.get(),
                other.batch_destination_var.get(),
            ) != other_values:
                raise AssertionError(
                    f"{name} Example input callback leaked into {other_name}"
                )
            callback_receipts.append(
                {
                    "endpoint": name,
                    "widget": "example_button",
                    "invoked": True,
                    "restored": list(restored),
                    "restored_batch": [
                        str(restored_batch[0]),
                        restored_batch[1],
                        restored_batch[2],
                    ],
                    "other_endpoint_unchanged": True,
                }
            )
    finally:
        for name, owner in endpoints.items():
            owner.parity_widgets["mode_notebook"].select(original_modes[name])
        app.notebook.select(original_endpoint)
        app.update_idletasks()

    return {
        "cas": expected["erta"]["cas"],
        "smiles": expected["erta"]["smiles"],
        "workbook": str(expected["erta"]["workbook"]),
        "workbook_size_bytes": len(erta_bytes),
        "workbook_sha256": default_hash_before,
        "workbook_sheet": default_metadata["sheet"],
        "workbook_headers": default_metadata["headers"],
        "workbook_row_count": default_metadata["row_count"],
        "cas_header": default_metadata["cas_header"],
        "smiles_header": default_metadata["smiles_header"],
        "batch_display": expected["erta"]["batch_display"],
        "batch_destination": expected["erta"]["batch_destination"],
        "path_equal": True,
        "content_equal": True,
        "default_writable": True,
        "default_matches_bundled": erta_bytes == bundled_bytes,
        "bundled_workbook": str(bundled_workbook),
        "bundled_size_bytes": len(bundled_bytes),
        "bundled_sha256": bundled_sha256,
        "bundled_headers": bundled_metadata["headers"],
        "bundled_row_count": bundled_metadata["row_count"],
        "bundled_cas_values": bundled_metadata["cas_values"],
        "fresh_copy": str(fresh_copy) if fresh_copy is not None else None,
        "fresh_copy_writable": fresh_copy is not None,
        "fresh_copy_matches_bundled": fresh_copy is not None,
        "fresh_copy_sha256": (
            _sha256_path(fresh_copy) if fresh_copy is not None else None
        ),
        "fresh_copy_size_bytes": (
            fresh_copy.stat().st_size if fresh_copy is not None else None
        ),
        "fresh_copy_resolution": (
            "resolve_shared_example_input() under a QA-only USERPROFILE and "
            "LOCALAPPDATA with frozen layout temporarily disabled"
            if fresh_copy is not None
            else None
        ),
        "callbacks": callback_receipts,
    }


def _require_recognition(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(f"{BATCH_RECOGNITION_GATE_ID}: {message}")


def _dialog_recognition_semantic(dialog: dict, context: str) -> str:
    expected = {
        "info": ("showinfo", "blue-information"),
        "warning": ("showwarning", "yellow-warning"),
        "error": ("showerror", "red-error"),
    }
    kind = str(dialog.get("kind", ""))
    _require_recognition(
        kind in expected,
        f"{context} has an unknown dialog kind: {dialog}",
    )
    function, icon = expected[kind]
    _require_recognition(
        dialog.get("function") == function
        and dialog.get("icon_semantic") == icon,
        f"{context} dialog function/icon semantic drift: {dialog}",
    )
    return icon


def _recognition_records_by_phase(records: list[dict]) -> dict:
    grouped = {
        endpoint: {}
        for endpoint in ("erta", "eralpha")
    }
    for record in records:
        endpoint = record.get("endpoint")
        phase = record.get("phase")
        _require_recognition(
            endpoint in grouped and isinstance(phase, str) and phase,
            f"malformed batch recognition record: {record}",
        )
        grouped[endpoint].setdefault(phase, []).append(record)
    return grouped


def _count_normalized_progress(text: str) -> str:
    match = re.fullmatch(r"(\d+)% - (\d+)/(\d+) - (.+)", text)
    _require_recognition(
        match is not None,
        f"batch progress has an unrecognized visible format: {text!r}",
    )
    return "<percent>% - <current>/<total> - " + match.group(4)


def _resolution_transition_receipt(record: dict) -> dict:
    events = list(record.get("transitions") or ())
    progress_events = [
        event
        for event in events
        if event.get("channel") == "progress"
        and str(event.get("text", "")).endswith(
            " - Resolving CAS/SMILES"
        )
    ]
    progress_values = [
        str(event.get("text", ""))
        for event in progress_events
    ]
    detail_pattern = re.compile(
        r"^Fetching SMILES from PubChem: "
        r"(?P<current>\d+) / (?P<total>\d+) "
        r"\((?P<cas>[^()\r\n]+)\)$"
    )
    detail_events = [
        event
        for event in events
        if event.get("channel") == "detail"
        and detail_pattern.fullmatch(str(event.get("text", "")))
    ]
    detail_values = [
        str(event.get("text", ""))
        for event in detail_events
    ]
    _require_recognition(
        progress_values,
        f"{record['endpoint']} did not display aggregate CAS-resolution progress",
    )
    _require_recognition(
        detail_values,
        f"{record['endpoint']} did not display per-row PubChem detail at the bottom",
    )
    for value in progress_values:
        _count_normalized_progress(value)
    parsed_details = [detail_pattern.fullmatch(value) for value in detail_values]
    _require_recognition(
        int(progress_events[0]["sequence"])
        < int(detail_events[0]["sequence"]),
        f"{record['endpoint']} per-row PubChem detail appeared before "
        "aggregate resolution progress",
    )
    _require_recognition(
        all(
            1 <= int(match.group("current"))
            <= int(match.group("total"))
            for match in parsed_details
            if match is not None
        ),
        f"{record['endpoint']} displayed an invalid PubChem row counter",
    )
    return {
        "actual_progress_strings": progress_values,
        "actual_detail_strings": detail_values,
        "progress_format": (
            "<percent>% - <current>/<total> - Resolving CAS/SMILES"
        ),
        "detail_format": (
            "Fetching SMILES from PubChem: "
            "<current> / <total> (<CAS>)"
        ),
        "cas_values": [
            match.group("cas")
            for match in parsed_details
            if match is not None
        ],
    }


def _assert_batch_controls(
    record: dict,
    expected_state: str,
    context: str,
) -> None:
    expected = {
        "input": expected_state,
        "template": expected_state,
        "run": expected_state,
    }
    _require_recognition(
        record.get("controls") == expected,
        f"{context} control-state drift: {record.get('controls')}",
    )


def _assert_unavailable_count_line(
    text: str,
    unavailable_count: int,
    context: str,
) -> None:
    expected = f"{BATCH_UNAVAILABLE_LABEL}: {unavailable_count}"
    _require_recognition(
        expected in text.splitlines(),
        f"{context} does not visibly contain the exact line {expected!r}",
    )


def _published_recognition_receipt(
    record: dict,
    *,
    availability: str,
) -> dict:
    endpoint = record["endpoint"]
    context = f"{endpoint}/{availability}"
    total = int(record.get("total_count", -1))
    unavailable = int(record.get("unavailable_count", -1))
    workbook_unavailable = int(
        record.get("workbook_unavailable_count", -1)
    )
    available = int(record.get("available_count", -1))
    _require_recognition(total > 0, f"{context} has no result rows")
    _require_recognition(
        available + unavailable == total,
        f"{context} available/unavailable counts do not total {total}",
    )
    _require_recognition(
        workbook_unavailable == unavailable,
        f"{context} UI count {unavailable} differs from workbook count "
        f"{workbook_unavailable}",
    )
    if availability == "success":
        _require_recognition(
            unavailable == 0 and available == total,
            f"{context} is not an all-available success",
        )
    elif availability == "partial_unavailable":
        _require_recognition(
            0 < unavailable < total and available > 0,
            f"{context} is not partially unavailable",
        )
    elif availability == "all_unavailable":
        _require_recognition(
            unavailable == total and available == 0,
            f"{context} is not all unavailable",
        )
    else:
        raise AssertionError(f"unknown publication availability: {availability}")

    output = str(record.get("output", ""))
    result = str(record.get("result", ""))
    status = str(record.get("status", ""))
    dialog = record.get("dialog")
    _require_recognition(
        bool(output) and record.get("output_exists") is True,
        f"{context} did not publish its recorded workbook",
    )
    _require_recognition(
        str(record.get("progress", ""))
        == f"100% - {total}/{total} - Completed",
        f"{context} does not display terminal Completed progress: "
        f"{record.get('progress')!r}",
    )
    _require_recognition(
        result.startswith("Batch job completed and workbook saved.\n")
        and output in result
        and status == f"Batch prediction completed: {output}",
        f"{context} does not visibly identify the completed workbook",
    )
    _require_recognition(
        isinstance(dialog, dict)
        and dialog.get("kind") == "info"
        and dialog.get("title") == "Batch prediction done",
        f"{context} did not use the successful publication dialog: {dialog}",
    )
    _require_recognition(
        output in str(dialog.get("message", "")),
        f"{context} success dialog omits the published workbook",
    )
    _require_recognition(
        (dialog.get("batch_controls_at_dialog") or {}).get(endpoint)
        == {
            "input": "normal",
            "template": "normal",
            "run": "normal",
        },
        f"{context} dialog appeared before batch controls were restored",
    )
    icon = _dialog_recognition_semantic(dialog, context)
    _require_recognition(
        icon == "blue-information",
        f"{context} successful publication is not blue/info",
    )
    _assert_unavailable_count_line(result, unavailable, f"{context} result")
    _assert_unavailable_count_line(
        str(dialog["message"]),
        unavailable,
        f"{context} dialog",
    )
    if availability == "success":
        availability_sentence = "All rows were predicted."
    elif availability == "partial_unavailable":
        availability_sentence = (
            f"{unavailable} row(s) could not be predicted. "
            "Row-level reasons are saved in the workbook."
        )
    else:
        availability_sentence = (
            "No rows could be predicted. "
            "Row-level reasons are saved in the workbook."
        )
    _require_recognition(
        availability_sentence in result
        and availability_sentence in str(dialog["message"]),
        f"{context} omits the exact row-availability explanation",
    )
    _assert_batch_controls(record, "normal", context)
    return {
        "state_meaning": "published",
        "terminal_stage": "Completed",
        "dialog_title": dialog["title"],
        "dialog_function": dialog["function"],
        "icon_semantic": icon,
        "availability": availability,
        "count_fields_normalized": True,
        "endpoint_label_normalized": True,
        "raw_progress": record["progress"],
        "raw_status": status,
        "raw_result": result,
        "raw_dialog_message": dialog["message"],
        "total_count": total,
        "available_count": available,
        "unavailable_count": unavailable,
        "workbook_unavailable_count": workbook_unavailable,
        "availability_sentence": availability_sentence,
    }


def evaluate_batch_recognition_contract(
    records: list[dict],
    widget_contract: dict,
) -> dict:
    """Fail closed on the automated cross-endpoint batch recognition contract."""
    _require_recognition(
        widget_contract.get("passed") is True
        and widget_contract.get("exact_layout_match") is True
        and set(widget_contract.get("endpoints", {}))
        == {"erta", "eralpha"},
        "batch feedback widget placement/binding evidence is incomplete",
    )
    for endpoint, receipt in widget_contract["endpoints"].items():
        bindings = receipt.get("bindings") or {}
        _require_recognition(
            receipt.get("passed") is True
            and receipt.get("criterion_id")
            == "BRG-01-FEEDBACK-PLACEMENT"
            and receipt.get("reading_order")
            == [
                "run_batch",
                "aggregate_progress",
                "batch_result",
                "per_row_detail",
            ],
            f"{endpoint} feedback placement/binding contract failed",
        )
        _require_recognition(
            bindings.get("run_callback") == "batch_predict_clicked"
            and "batch_predict_clicked"
            in str(bindings.get("run_command", ""))
            and bool(bindings.get("progress_value_variable"))
            and bool(bindings.get("progress_text_variable"))
            and bool(bindings.get("per_row_detail_variable")),
            f"{endpoint} feedback widget bindings are incomplete: {bindings}",
        )

    grouped = _recognition_records_by_phase(records)
    required_phases = {
        "new_run",
        "resolving",
        "success",
        "partial_unavailable",
        "all_unavailable",
        "no_output_failure",
    }
    for endpoint in ("erta", "eralpha"):
        missing = required_phases - set(grouped[endpoint])
        _require_recognition(
            not missing,
            f"{endpoint} is missing recognition phases: {sorted(missing)}",
        )

    new_run_receipts = {}
    resolution_receipts = {}
    publication_receipts = {
        phase: {}
        for phase in (
            "success",
            "partial_unavailable",
            "all_unavailable",
        )
    }
    failure_receipts = {}
    screenshot_receipts = []
    for endpoint in ("erta", "eralpha"):
        prior_terminal_cleared = False
        for record in grouped[endpoint]["new_run"]:
            context = f"{endpoint}/new_run/{record.get('scenario', '')}"
            progress = str(record.get("progress", ""))
            status = str(record.get("status", ""))
            result = str(record.get("result", ""))
            prior_progress = str(record.get("prior_progress", ""))
            prior_result = str(record.get("prior_result", ""))
            _assert_batch_controls(record, "disabled", context)
            _require_recognition(
                record.get("dialog") is None
                and progress
                == "0% - 0/0 - Reading input workbook"
                and status == "Batch prediction started."
                and result
                == (
                    "Batch prediction is running.\n\n"
                    "Completion details will appear after the workbook is "
                    "saved."
                )
                and "Completed" not in progress
                and "Completed" not in status
                and "completed" not in result.casefold()
                and "Output workbook:" not in result,
                f"{context} retained terminal success presentation",
            )
            if (
                "Completed" in prior_progress
                and "Output workbook:" in prior_result
            ):
                prior_terminal_cleared = True
            screenshot = record.get("screenshot")
            if isinstance(screenshot, dict):
                screenshot_receipts.append(
                    {
                        "endpoint": endpoint,
                        "state": "running",
                        **screenshot,
                    }
                )
        _require_recognition(
            prior_terminal_cleared,
            f"{endpoint} did not prove a prior Completed result was cleared "
            "by a subsequent real callback",
        )
        representative_new_run = grouped[endpoint]["new_run"][-1]
        new_run_receipts[endpoint] = {
            "state_meaning": "running",
            "progress_format": _count_normalized_progress(
                str(representative_new_run["progress"])
            ),
            "stale_completed_visible": False,
            "controls": "disabled",
        }

        resolving_candidates = [
            _resolution_transition_receipt(record)
            for record in grouped[endpoint]["resolving"]
        ]
        resolution_receipts[endpoint] = resolving_candidates[0]

        for phase in publication_receipts:
            record = grouped[endpoint][phase][-1]
            publication_receipts[phase][endpoint] = (
                _published_recognition_receipt(
                    record,
                    availability=phase,
                )
            )
            screenshot = record.get("screenshot")
            if phase == "success" and isinstance(screenshot, dict):
                screenshot_receipts.append(
                    {
                        "endpoint": endpoint,
                        "state": "completion",
                        **screenshot,
                    }
                )

        failure = grouped[endpoint]["no_output_failure"][-1]
        context = f"{endpoint}/no_output_failure"
        dialog = failure.get("dialog")
        result = str(failure.get("result", ""))
        status = str(failure.get("status", ""))
        _require_recognition(
            failure.get("output_exists") is False
            and not failure.get("output")
            and failure.get("fresh_output_count") == 0,
            f"{context} created or claimed an output",
        )
        _require_recognition(
            failure.get("progress") == "100% - 0/0 - Failed"
            and result.startswith("Batch prediction failed.\n")
            and "Output workbook:" not in result
            and "completed" not in result.casefold()
            and status.startswith("Batch prediction failed:")
            and "saved" not in status.casefold(),
            f"{context} does not visibly distinguish failure from publication",
        )
        _require_recognition(
            isinstance(dialog, dict)
            and dialog.get("kind") == "error"
            and dialog.get("title") == "Batch prediction failed",
            f"{context} did not use the failure dialog: {dialog}",
        )
        _require_recognition(
            (dialog.get("batch_controls_at_dialog") or {}).get(endpoint)
            == {
                "input": "normal",
                "template": "normal",
                "run": "normal",
            },
            f"{context} dialog appeared before controls were restored",
        )
        icon = _dialog_recognition_semantic(dialog, context)
        _require_recognition(
            icon == "red-error",
            f"{context} failure is not red/error",
        )
        _assert_batch_controls(failure, "normal", context)
        failure_receipts[endpoint] = {
            "state_meaning": "no-output failure",
            "terminal_stage": "Failed",
            "dialog_title": dialog["title"],
            "dialog_function": dialog["function"],
            "icon_semantic": icon,
            "fresh_output_count": 0,
            "raw_progress": failure["progress"],
            "raw_status": status,
            "raw_result": result,
            "raw_dialog_message": dialog["message"],
        }

    _require_recognition(
        {
            receipt["progress_format"]
            for receipt in new_run_receipts.values()
        }
        == {
            "<percent>% - <current>/<total> - Reading input workbook"
        },
        f"new-run meaning differs by endpoint: {new_run_receipts}",
    )
    _require_recognition(
        {
            receipt["progress_format"]
            for receipt in resolution_receipts.values()
        }
        == {
            "<percent>% - <current>/<total> - Resolving CAS/SMILES"
        }
        and {
            receipt["detail_format"]
            for receipt in resolution_receipts.values()
        }
        == {
            "Fetching SMILES from PubChem: "
            "<current> / <total> (<CAS>)"
        }
        and resolution_receipts["erta"]["cas_values"]
        == resolution_receipts["eralpha"]["cas_values"],
        f"CAS-resolution feedback format differs by endpoint: "
        f"{resolution_receipts}",
    )
    for phase, endpoint_receipts in publication_receipts.items():
        meanings = {
            (
                receipt["state_meaning"],
                receipt["terminal_stage"],
                receipt["dialog_title"],
                receipt["dialog_function"],
                receipt["icon_semantic"],
                receipt["availability"],
            )
            for receipt in endpoint_receipts.values()
        }
        _require_recognition(
            len(meanings) == 1,
            f"{phase} state meaning differs by endpoint: {endpoint_receipts}",
        )
    _require_recognition(
        len(
            {
                (
                    receipt["state_meaning"],
                    receipt["terminal_stage"],
                    receipt["dialog_title"],
                    receipt["dialog_function"],
                    receipt["icon_semantic"],
                )
                for receipt in failure_receipts.values()
            }
        )
        == 1,
        f"no-output failure meaning differs by endpoint: {failure_receipts}",
    )

    expected_screenshots = {
        (endpoint, state)
        for endpoint in ("erta", "eralpha")
        for state in ("running", "completion")
    }
    observed_screenshots = {
        (receipt["endpoint"], receipt["state"])
        for receipt in screenshot_receipts
    }
    _require_recognition(
        expected_screenshots <= observed_screenshots,
        "running/completion own-window screenshot receipts are incomplete: "
        f"{sorted(expected_screenshots - observed_screenshots)}",
    )
    for screenshot in screenshot_receipts:
        _require_recognition(
            screenshot.get("method")
            == "PrintWindow(PW_RENDERFULLCONTENT)"
            and screenshot.get("capture_scope") == "own Tk HWND only"
            and (
                (
                    screenshot.get("supported") is True
                    and bool(screenshot.get("path"))
                    and bool(screenshot.get("target_hwnd"))
                    and Path(str(screenshot["path"])).is_file()
                    and re.fullmatch(
                        r"[0-9a-f]{64}",
                        str(screenshot.get("sha256", "")),
                    )
                    is not None
                )
                or (
                    screenshot.get("supported") is False
                    and bool(screenshot.get("reason"))
                )
            ),
            "screenshot evidence used a non-own-window method or is "
            f"incomplete: {screenshot}",
        )

    passed_ids = list(BATCH_RECOGNITION_CRITERIA)
    return {
        "gate_id": BATCH_RECOGNITION_GATE_ID,
        "passed": True,
        "criteria": [
            {
                "id": criterion_id,
                "description": BATCH_RECOGNITION_CRITERIA[criterion_id],
                "passed": True,
            }
            for criterion_id in passed_ids
        ],
        "criteria_ids": passed_ids,
        "widget_contract": widget_contract,
        "comparison": {
            "new_run": new_run_receipts,
            "resolving": resolution_receipts,
            **publication_receipts,
            "no_output_failure": failure_receipts,
        },
        "records": records,
        "screenshots": screenshot_receipts,
        "normalization": {
            "normalized_fields": [
                "endpoint display label",
                "endpoint-specific prediction-label counts",
            ],
            "whole_messages_normalized": False,
            "raw_displayed_strings_retained": True,
        },
        "validation_scope": {
            "automated_recognition_contract_checks": True,
            "real_human_usability_study_performed": False,
            "statement": (
                "This gate checks machine-observable recognition cues, real "
                "callback transitions, displayed strings, widget placement, "
                "and messagebox API semantics. It is not evidence from a real "
                "human usability study; no such study was performed."
            ),
        },
    }


class NativePackageQa:
    """Drive real Tk callbacks and record fail-closed packaged-app evidence."""

    BASELINE_CHECKS = {
        "erta_eralpha_exact_ui_geometry_and_examples",
        "eralpha_real_resource_buttons_release_enforcement_and_isolation",
        "erta_startup_resources_and_controls",
        "erta_options_and_resource_controls",
        "erta_exact_valid_single",
        "erta_legacy_invalid_direct_smiles",
        "erta_mocked_pubchem_callback",
        "erta_mocked_pubchem_prediction",
        "erta_model_ad_browse_callbacks",
        "erta_template_callback_and_dialog",
        "erta_batch_workbook_graphs_and_order",
        "erba_evidence_caveat_and_dialog",
        "erba_route",
        "erba_invalid_wildcard",
        "erba_invalid_pipe",
        "erba_invalid_metal",
        "erba_invalid_blank",
        "erba_invalid_invalid_cas",
        "erba_eralpha_batch_workbook_ad_graphs_and_summary",
        "fixed_eralpha_classification_tab",
        "erta_erba_erta_state_isolation",
        "read_only_install_and_writable_state",
    }
    BATCH_FEEDBACK_CHECKS = {
        "erta_batch_failure_dialog_and_control_recovery",
        "eralpha_batch_failure_dialog_and_control_recovery",
    }
    SHARED_EXAMPLE_BATCH_CHECKS = {
        "erta_shared_test_batch_success_dialog_and_predictions",
        "eralpha_shared_test_batch_success_dialog_and_predictions",
    }
    VALID_SMILES = "C[C@]12CC[C@H]3[C@@H]([C@@H]1CC[C@@H]2O)CCC4=CC(=CC=C34)O"
    LEGACY_INVALID_SMILES = "not a SMILES"
    MOCK_CAS = "50-00-0"
    MOCK_SMILES = "C=O"
    MOCK_UNAVAILABLE_CAS_VALUES = ("7732-18-5", "58-08-2")
    ERBA_BATCH_AD_COLUMNS = ERALPHA_AD_COLUMNS
    ERBA_BATCH_GRAPH_FILES = {
        "binding_class_count.png",
        "binding_probability_histogram.png",
        "top_binding_chemicals.png",
        "ad_pca_plot.png",
        "ad_decision_plot.png",
    }
    ERBA_PROVENANCE_COLUMNS = (
        "Protocol_SHA256",
        "Source_Manifest_SHA256",
        "Historical_Exposure_Manifest_SHA256",
        "Split_Manifest_SHA256",
        "Nested_CV_SHA256",
        "Internal_Resplit_SHA256",
        "Report_SHA256",
        "Caveat_SHA256",
    )

    def __init__(self, app, destination: str | Path) -> None:
        self.app = app
        self.destination = Path(destination)
        self.destination.parent.mkdir(parents=True, exist_ok=True)
        self.root = Path(app.output_root) / "native-package-qa"
        self.root.mkdir(parents=True, exist_ok=True)
        self.run_root = self.root / f"run-{os.getpid()}-{time.time_ns()}"
        self.run_root.mkdir(parents=True, exist_ok=False)
        self.stage = -4
        self._tick_active = False
        self.route_index = 0
        self.invalid_index = 0
        self.started_at = time.monotonic()
        self.deadline = self.started_at + 900
        self.erta_summary = ""
        self.erta_probability = ""
        self.erta_batch_input_path: Path | None = None
        self.erta_batch_outputs_before: set[Path] = set()
        self.erba_batch_input_path: Path | None = None
        self.erba_batch_outputs_before: set[Path] = set()
        self.erba_batch_graph_dirs_before: set[Path] = set()
        self.dialogs: list[dict] = []
        self.batch_recognition_records: list[dict] = []
        self.batch_recognition_widget_contract: dict | None = None
        self._active_batch_transition_observer: dict | None = None
        self.patches: list[tuple[object, str, object]] = []
        self.mock_cas_enabled = False
        self.mock_cas_values: set[str] = {self.MOCK_CAS}
        self.shared_default_path: Path | None = None
        self.shared_default_sha256 = ""
        self.shared_bundle_path: Path | None = None
        self.shared_bundle_sha256 = ""
        self.qa_shared_example_path: Path | None = None
        self.erta_batch_output_path: Path | None = None
        gate_value = os.environ.get("ER_PREDICTOR_EXTERNAL_DRIVER_GATE", "").strip()
        self.external_driver_gate = Path(gate_value) if gate_value else None
        self.run_nonce = os.environ.get("ER_PREDICTOR_AUTOMATION_NONCE", "").strip()
        self.target_pid = os.getpid()
        self.routes = [
            (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA),
        ]
        self.invalid_cases = [
            ("wildcard", "[*]CC", "", "wildcard_smiles"),
            ("pipe", "CC|CC", "", "cxsmiles_not_allowed"),
            ("metal", "[Na+]", "", "no_carbon"),
            ("blank", "", "", "blank_smiles"),
            ("invalid_cas", "", "123-45-6", "CAS is invalid"),
        ]
        self.transcript = {
            "schema_version": 2,
            "kind": "native-app-automation-transcript",
            "surface": "native-desktop",
            "executable": os.fspath(Path(os.sys.executable).resolve()),
            "run_nonce": self.run_nonce,
            "target_pid": self.target_pid,
            "checks": [],
            "automation_scopes": [],
        }

    def start(self) -> None:
        self._install_dialog_observer()
        callback = self._wait_for_external_driver if self.external_driver_gate else self.tick
        self.app.after(100, callback)

    def _wait_for_external_driver(self) -> None:
        try:
            if time.monotonic() >= self.deadline:
                raise TimeoutError("external pywinauto driver did not release the native QA gate")
            if not self.external_driver_gate.is_file():
                self.app.after(100, self._wait_for_external_driver)
                return
            payload = json.loads(self.external_driver_gate.read_text(encoding="utf-8"))
            self._require(payload.get("passed") is True, "external pywinauto driver gate is not passing")
            self._require(bool(self.run_nonce), "automation run nonce is missing")
            self._require(payload.get("run_nonce") == self.run_nonce, "external driver gate nonce mismatch")
            self._require(payload.get("target_pid") == self.target_pid, "external driver gate target PID mismatch")
            self._record(
                "external_pywinauto_driver_gate",
                driver_kind=payload.get("driver_kind"),
                run_nonce=self.run_nonce,
                target_pid=self.target_pid,
                real_input_actions=payload.get("real_input_actions"),
                evidence_path=str(self.external_driver_gate),
            )
            self.app.after(100, self.tick)
        except Exception:
            self._finish(False, traceback.format_exc())

    def _body(self, tab=None) -> str:
        tab = tab or self.app.erba_tab
        fields = (
            tab.single_prediction_summary_var,
            tab.single_negative_probability_var,
            tab.single_positive_probability_var,
            tab.single_pic50_var,
            tab.single_ic50_var,
            tab.single_ad_domain_var,
            tab.single_detail_var,
            tab.single_status_var,
        )
        return "\n".join(value.get().strip() for value in fields if value.get().strip())

    def _record(self, check: str, **evidence) -> None:
        self.transcript["checks"].append({"check": check, **evidence, "passed": True})

    def _require(self, condition: bool, message: str) -> None:
        if not condition:
            raise AssertionError(message)

    def _require_eralpha_primary_headers(
        self,
        headers: list,
        passthrough_headers: tuple[str, ...],
        context: str,
    ) -> None:
        expected = [
            *passthrough_headers,
            *ERALPHA_PRIMARY_TRUSTED_COLUMNS,
        ]
        self._require(
            headers == expected,
            f"{context} primary header order drift: {headers}",
        )
        self._require(
            len(headers) == len(set(headers))
            and set(headers).isdisjoint(
                ERALPHA_PRIMARY_FORBIDDEN_COLUMNS
            ),
            f"{context} repeats technical diagnostics, model provenance, "
            "or evidence caveats in Predictions",
        )

    def _require_eralpha_diagnostics(
        self,
        headers: list,
        diagnostic_rows: list[dict],
        prediction_rows: list[dict],
        context: str,
    ) -> None:
        self._require(
            headers == list(ERBA_CLASSIFICATION_DIAGNOSTIC_COLUMNS),
            f"{context} Diagnostics header order drift: {headers}",
        )
        self._require(
            len(diagnostic_rows) == len(prediction_rows)
            and [
                row["row_index"] for row in diagnostic_rows
            ]
            == list(range(len(prediction_rows)))
            and [
                str(row["CAS"] or "") for row in diagnostic_rows
            ]
            == [
                str(row["CAS"] or "") for row in prediction_rows
            ]
            and all(
                str(row[column] or "").strip()
                for row in diagnostic_rows
                for column in (
                    "Row_ID",
                    "Status_Code",
                    "SMILES_Provenance",
                    "Result_Status",
                    "Reason_Category",
                    "Reason_Description",
                    "Recommended_Action",
                )
            ),
            f"{context} Diagnostics rows are not one-to-one, ordered, "
            "reasoned mappings of Predictions",
        )

    @staticmethod
    def _eralpha_prediction_is_consistent(row: dict) -> bool:
        negative = row["Probability_Negative_0"]
        positive = row["Probability_Positive_1"]
        prediction = row["Prediction"]
        if (
            isinstance(negative, bool)
            or not isinstance(negative, (int, float))
            or isinstance(positive, bool)
            or not isinstance(positive, (int, float))
            or isinstance(prediction, bool)
            or prediction not in {0, 1}
        ):
            return False
        expected_prediction = 1 if positive >= 0.5 else 0
        return (
            0.0 <= negative <= 1.0
            and 0.0 <= positive <= 1.0
            and abs((negative + positive) - 1.0) <= 1e-6
            and prediction == expected_prediction
            and row["Prediction_label"]
            == ("Binding" if prediction == 1 else "Non-binding")
        )

    @staticmethod
    def _eralpha_prediction_payload_is_blank(row: dict) -> bool:
        return all(
            row[column] is None
            for column in (
                "Probability_Negative_0",
                "Probability_Positive_1",
                "Prediction",
                "Prediction_label",
                *ERALPHA_AD_COLUMNS,
            )
        )

    def _patch(self, owner: object, name: str, replacement: object) -> None:
        self.patches.append((owner, name, getattr(owner, name)))
        setattr(owner, name, replacement)

    def _restore_patches(self) -> None:
        while self.patches:
            owner, name, original = self.patches.pop()
            setattr(owner, name, original)

    def _install_dialog_observer(self) -> None:
        # Native message boxes block unattended Tk automation; record their exact
        # visible title/body while preserving the callback path that invokes them.
        from gui import erba_tab, main_window

        def observe(kind, function, icon_semantic):
            def handler(title, message, **_kwargs):
                controls_at_dialog = {}
                for endpoint, owner in (
                    ("erta", self.app),
                    ("eralpha", self.app.eralpha_tab),
                ):
                    try:
                        controls_at_dialog[endpoint] = (
                            self._batch_control_states(owner)
                        )
                    except Exception as error:
                        controls_at_dialog[endpoint] = {
                            "unavailable": f"{type(error).__name__}: {error}"
                        }
                self.dialogs.append(
                    {
                        "kind": kind,
                        "function": function,
                        "icon_semantic": icon_semantic,
                        "title": str(title),
                        "message": str(message),
                        "batch_controls_at_dialog": controls_at_dialog,
                    }
                )
                return "ok"
            return handler

        messagebox_modules = {
            id(main_window.messagebox): main_window.messagebox,
            id(erba_tab.messagebox): erba_tab.messagebox,
        }
        for messagebox_module in messagebox_modules.values():
            self._patch(
                messagebox_module,
                "showerror",
                observe("error", "showerror", "red-error"),
            )
            self._patch(
                messagebox_module,
                "showinfo",
                observe("info", "showinfo", "blue-information"),
            )
            self._patch(
                messagebox_module,
                "showwarning",
                observe("warning", "showwarning", "yellow-warning"),
            )
        self.transcript["automation_scopes"].append({
            "scope": "messagebox_observer",
            "functions": ["showinfo", "showwarning", "showerror"],
            "reason": "record dialog-visible semantics without blocking opt-in automation",
        })

    def _install_mock_cas(self) -> None:
        if self.mock_cas_enabled:
            return
        from gui import erba_tab, main_window

        def mocked_cas_to_smiles(cas: str) -> dict:
            normalized = str(cas).strip()
            if normalized in self.MOCK_UNAVAILABLE_CAS_VALUES:
                raise RuntimeError(
                    "deterministic native-QA PubChem unavailability"
                )
            if normalized not in self.mock_cas_values:
                raise RuntimeError(f"unexpected deterministic QA CAS: {cas}")
            return {
                "CanonicalSMILES": self.MOCK_SMILES,
                "PubChem_CID": "712",
                "PubChem_status": "Found",
            }

        self._patch(main_window, "cas_to_smiles", mocked_cas_to_smiles)
        self._patch(erba_tab, "cas_to_smiles", mocked_cas_to_smiles)
        self.mock_cas_enabled = True
        self.transcript["automation_scopes"].append({
            "scope": "mocked_pubchem",
            "cas_values": sorted(self.mock_cas_values),
            "unavailable_cas_values": list(
                self.MOCK_UNAVAILABLE_CAS_VALUES
            ),
            "canonical_smiles": self.MOCK_SMILES,
            "cid": "712",
            "reason": (
                "deterministic opt-in UI callback and bundled 25-CAS batch "
                "coverage; recognized success CAS values are deliberately "
                "substituted with the same valid SMILES, while two explicit "
                "CAS values raise a deterministic unavailable response; no "
                "network request is made"
            ),
            "production_equivalent": False,
            "online_verification": (
                "Run the packaged app without native-QA substitution to resolve "
                "the actual 25 CAS values through PubChem."
            ),
        })

    def _install_picker(self, save_path: Path) -> None:
        from gui import main_window

        def choose_open_file(**kwargs):
            title = str(kwargs.get("title", ""))
            return self.app.ad_ref_path_var.get() if "AD reference" in title else self.app.model_path_var.get()

        self._patch(main_window.filedialog, "askopenfilename", choose_open_file)
        self._patch(main_window.filedialog, "asksaveasfilename", lambda **_kwargs: str(save_path))
        self.transcript["automation_scopes"].append({
            "scope": "filedialog_selection",
            "path": str(save_path),
            "reason": "exercise callback state transitions without coordinate-driven native dialogs",
        })

    def _finish(self, passed: bool, error: str = "") -> None:
        self._stop_batch_transition_observer(record_resolution=False)
        self._restore_patches()
        self.transcript["passed"] = passed
        self.transcript["duration_seconds"] = time.monotonic() - self.started_at
        if error:
            self.transcript["error"] = error
        self.destination.write_text(
            json.dumps(self.transcript, indent=2, ensure_ascii=False, default=str) + "\n",
            encoding="utf-8",
        )
        hold_seconds = max(1, int(os.environ.get("ER_PREDICTOR_AUTOMATION_HOLD_SECONDS", "30")))
        self.app.after(hold_seconds * 1000, self.app.destroy)

    def _reschedule(self, delay: int = 100) -> None:
        if time.monotonic() >= self.deadline:
            raise TimeoutError(f"native package QA timed out at stage {self.stage}")
        self.app.after(delay, self.tick)

    def _wait_erta_result(self, previous: str = "") -> bool:
        return (
            self.app.status_var.get().startswith("Single prediction done:")
            and self.app.prediction_summary_var.get() != "Prediction: -"
            and (not previous or self.app.result_text.get("1.0", "end").strip() != previous)
        )

    def _wait_dialog(self, count: int) -> bool:
        return len(self.dialogs) > count

    @staticmethod
    def _batch_control_states(owner) -> dict:
        return {
            "input": str(owner.batch_input_button.cget("state")),
            "template": str(owner.download_template_button.cget("state")),
            "run": str(owner.run_batch_button.cget("state")),
        }

    def _batch_owner(self, endpoint: str):
        if endpoint == "erta":
            return self.app
        if endpoint == "eralpha":
            return self.app.eralpha_tab
        raise AssertionError(f"unknown batch endpoint: {endpoint}")

    @staticmethod
    def _batch_status_value(owner) -> str:
        status_var = getattr(owner, "status_var", None)
        if status_var is None:
            status_var = owner.batch_status_var
        return str(status_var.get())

    @staticmethod
    def _batch_result_value(owner) -> str:
        return str(owner.batch_result.get("1.0", "end")).strip()

    def _select_batch_surface(self, endpoint: str) -> None:
        owner = self._batch_owner(endpoint)
        controls = owner.parity_widgets
        self.app.notebook.select(controls["endpoint_tab"])
        controls["mode_notebook"].select(controls["batch_tab"])
        self.app.update()
        self.app.update_idletasks()

    def _begin_batch_transition_observer(
        self,
        endpoint: str,
        scenario: str,
    ) -> None:
        self._require(
            self._active_batch_transition_observer is None,
            "a batch transition observer is already active",
        )
        owner = self._batch_owner(endpoint)
        self._select_batch_surface(endpoint)
        detail_var = getattr(owner, "status_var", None)
        if detail_var is None:
            detail_var = owner.batch_status_var
        events: list[dict] = []
        traces = []

        def add_trace(channel, variable):
            def observed(*_args):
                events.append(
                    {
                        "sequence": len(events) + 1,
                        "channel": channel,
                        "text": str(variable.get()),
                    }
                )

            token = variable.trace_add("write", observed)
            traces.append((variable, token))

        add_trace("progress", owner.batch_progress_var)
        add_trace("detail", detail_var)
        self._active_batch_transition_observer = {
            "endpoint": endpoint,
            "scenario": scenario,
            "events": events,
            "traces": traces,
            "dialog_count": len(self.dialogs),
            "prior_progress": str(owner.batch_progress_var.get()),
            "prior_status": self._batch_status_value(owner),
            "prior_result": self._batch_result_value(owner),
        }

    def _stop_batch_transition_observer(
        self,
        *,
        record_resolution: bool = True,
    ) -> dict | None:
        observer = self._active_batch_transition_observer
        if observer is None:
            return None
        self._active_batch_transition_observer = None
        for variable, token in observer["traces"]:
            try:
                variable.trace_remove("write", token)
            except Exception:
                pass
        has_resolution_progress = any(
            event["channel"] == "progress"
            and event["text"].endswith(
                " - Resolving CAS/SMILES"
            )
            for event in observer["events"]
        )
        has_pubchem_detail = any(
            event["channel"] == "detail"
            and event["text"].startswith(
                "Fetching SMILES from PubChem:"
            )
            for event in observer["events"]
        )
        if (
            record_resolution
            and has_resolution_progress
            and has_pubchem_detail
        ):
            self.batch_recognition_records.append(
                {
                    "endpoint": observer["endpoint"],
                    "phase": "resolving",
                    "scenario": observer["scenario"],
                    "transitions": list(observer["events"]),
                }
            )
        return observer

    def _capture_batch_screenshot(
        self,
        endpoint: str,
        state: str,
    ) -> dict:
        owner = self._batch_owner(endpoint)
        self._require(
            self.app.notebook.select()
            == str(owner.parity_widgets["endpoint_tab"])
            and owner.parity_widgets["mode_notebook"].select()
            == str(owner.parity_widgets["batch_tab"]),
            f"{endpoint} {state} screenshot is not on its visible batch surface",
        )
        self.app.update_idletasks()
        screenshot = _capture_own_tk_window(
            self.app,
            self.run_root
            / "batch-recognition-screenshots"
            / f"{endpoint}-{state}.png",
        )
        screenshot["capture_scope"] = "own Tk HWND only"
        return screenshot

    def _record_batch_new_run(
        self,
        endpoint: str,
        *,
        screenshot: bool = False,
    ) -> dict:
        observer = self._active_batch_transition_observer
        self._require(
            observer is not None and observer["endpoint"] == endpoint,
            f"{endpoint} new-run snapshot has no matching observer",
        )
        owner = self._batch_owner(endpoint)
        record = {
            "endpoint": endpoint,
            "phase": "new_run",
            "scenario": observer["scenario"],
            "progress": str(owner.batch_progress_var.get()),
            "status": self._batch_status_value(owner),
            "result": self._batch_result_value(owner),
            "controls": self._batch_control_states(owner),
            "dialog": (
                None
                if len(self.dialogs) == observer["dialog_count"]
                else self.dialogs[observer["dialog_count"]]
            ),
            "prior_progress": observer["prior_progress"],
            "prior_status": observer["prior_status"],
            "prior_result": observer["prior_result"],
        }
        if screenshot:
            record["screenshot"] = self._capture_batch_screenshot(
                endpoint,
                "running",
            )
        self.batch_recognition_records.append(record)
        return record

    def _record_batch_terminal(
        self,
        endpoint: str,
        phase: str,
        *,
        dialog: dict,
        output: str | Path | None,
        output_exists: bool,
        total_count: int = 0,
        available_count: int = 0,
        unavailable_count: int = 0,
        workbook_unavailable_count: int = 0,
        fresh_output_count: int = 0,
        screenshot: bool = False,
    ) -> dict:
        owner = self._batch_owner(endpoint)
        observer = self._stop_batch_transition_observer()
        self._require(
            observer is not None
            and observer["endpoint"] == endpoint,
            f"{endpoint} terminal snapshot has no matching observer",
        )
        record = {
            "endpoint": endpoint,
            "phase": phase,
            "scenario": observer["scenario"],
            "progress": str(owner.batch_progress_var.get()),
            "status": self._batch_status_value(owner),
            "result": self._batch_result_value(owner),
            "controls": self._batch_control_states(owner),
            "dialog": dialog,
            "output": str(output) if output is not None else "",
            "output_exists": bool(output_exists),
            "fresh_output_count": int(fresh_output_count),
            "total_count": int(total_count),
            "available_count": int(available_count),
            "unavailable_count": int(unavailable_count),
            "workbook_unavailable_count": int(
                workbook_unavailable_count
            ),
        }
        if screenshot:
            record["screenshot"] = self._capture_batch_screenshot(
                endpoint,
                "completion",
            )
        self.batch_recognition_records.append(record)
        return record

    def _terminal_dialog(
        self,
        initial_count: int,
        *,
        kind: str,
        title: str,
        context: str,
    ) -> dict:
        observed = self.dialogs[initial_count:]
        self._require(
            len(observed) == 1,
            f"{context} emitted {len(observed)} dialogs instead of exactly one: "
            f"{observed}",
        )
        dialog = observed[0]
        self._require(
            dialog["kind"] == kind and dialog["title"] == title,
            f"{context} dialog contract drift: {dialog}",
        )
        return dialog

    def _copy_shared_example_for_batch(self, endpoint: str) -> Path:
        self._require(
            self.qa_shared_example_path is not None
            and self.qa_shared_example_path.is_file(),
            "fresh bundled shared-example QA source is unavailable",
        )
        source = self.qa_shared_example_path
        source_hash = _sha256_path(source)
        destination = (
            self.run_root / "shared-example-success" / endpoint / "test.xlsx"
        )
        destination.parent.mkdir(parents=True, exist_ok=False)
        self._require(
            destination.resolve(strict=False) != source.resolve(strict=True),
            f"{endpoint} QA batch input unexpectedly targets the shared source",
        )
        shutil.copyfile(source, destination)
        self._require(
            destination.read_bytes() == source.read_bytes()
            and _sha256_path(source) == source_hash,
            f"{endpoint} QA copy changed the shared example bytes",
        )
        return destination.resolve(strict=True)

    def _make_corrupt_batch_input(self, endpoint: str) -> Path:
        destination = (
            self.run_root / "deterministic-batch-failure" / endpoint / "test.xlsx"
        )
        destination.parent.mkdir(parents=True, exist_ok=False)
        destination.write_bytes(
            b"ER_Predictor native QA deterministic invalid xlsx fixture\n"
        )
        return destination.resolve(strict=True)

    @staticmethod
    def _same_path(first: str | Path, second: str | Path) -> bool:
        try:
            return (
                Path(first).expanduser().resolve(strict=True)
                == Path(second).expanduser().resolve(strict=True)
            )
        except (OSError, RuntimeError):
            return False

    def _parity_startup_ready(self) -> bool:
        app = self.app
        tab = app.eralpha_tab
        statuses = (app.status_var.get(), tab.single_status_var.get())
        for status in statuses:
            if status.startswith("Startup failed:"):
                raise RuntimeError(status)
        return bool(
            app.predictor.is_loaded()
            and app.ad_calculator.fitted
            and app.loaded_model_path_var.get()
            and app.loaded_ad_ref_path_var.get()
            and tab.loaded_model_path_var.get()
            and tab.loaded_ad_ref_path_var.get()
            and tab.predictor is not None
        )

    def _invoke_open_file_button(
        self,
        button,
        selected_path: str | Path,
        *,
        purpose: str,
    ) -> dict:
        """Invoke a real widget command while replacing only the picker result."""
        from gui import erba_tab

        calls = []
        original = erba_tab.filedialog.askopenfilename

        def selected(**kwargs):
            calls.append(dict(kwargs))
            return os.fspath(selected_path)

        self._require(
            str(button.cget("state")) != "disabled",
            f"{purpose} widget is disabled",
        )
        erba_tab.filedialog.askopenfilename = selected
        try:
            button.invoke()
        finally:
            erba_tab.filedialog.askopenfilename = original
        self._require(
            len(calls) == 1,
            f"{purpose} did not invoke exactly one real filedialog callback",
        )
        receipt = {
            "purpose": purpose,
            "widget_text": str(button.cget("text")),
            "selected_path": os.fspath(selected_path),
            "dialog_title": str(calls[0].get("title", "")),
            "callback_invoked": True,
        }
        self.transcript["automation_scopes"].append(
            {
                "scope": "filedialog_selection_return_only",
                **receipt,
                "reason": (
                    "invoke the actual mapped Tk button/callback without "
                    "coordinate-driving a native chooser"
                ),
            }
        )
        return receipt

    @staticmethod
    def _erta_resource_state(app) -> dict:
        model_path = Path(app.model_path_var.get()).expanduser().resolve(
            strict=True
        )
        ad_path = Path(app.ad_ref_path_var.get()).expanduser().resolve(
            strict=True
        )
        return {
            "model_selection": str(model_path),
            "model_loaded": app.loaded_model_path_var.get(),
            "model_sha256": _sha256_path(model_path),
            "ad_selection": str(ad_path),
            "ad_loaded": app.loaded_ad_ref_path_var.get(),
            "ad_sha256": _sha256_path(ad_path),
            "predictor_identity": id(app.predictor),
            "predictor_model_path": str(
                getattr(app.predictor, "model_path", "")
            ),
            "ad_calculator_identity": id(app.ad_calculator),
            "status": app.status_var.get(),
        }

    @staticmethod
    def _erba_model_state(tab) -> dict:
        loaded = getattr(tab.predictor, "_loaded", {})
        return {
            "model_selection": tab.model_path_var.get(),
            "model_loaded": tab.loaded_model_path_var.get(),
            "predictor_identity": id(tab.predictor),
            "model_generation": int(tab._model_generation),
            "selected_model_id": (
                tab.selected_model_spec.model_id
                if tab.selected_model_spec is not None
                else ""
            ),
            "loaded_route_keys": sorted(
                f"{task.value}/{subtype.value}" for task, subtype in loaded
            ),
            "loaded_artifact_identities": sorted(
                id(value[0]) for value in loaded.values()
            ),
        }

    def _start_exact_ui_parity_stage(self) -> None:
        app = self.app
        tab = app.eralpha_tab
        example_receipt = inspect_shared_example_parity(
            app,
            self.run_root / "fresh-default-profile",
        )
        self.shared_default_path = Path(example_receipt["workbook"])
        self.shared_default_sha256 = example_receipt["workbook_sha256"]
        self.shared_bundle_path = Path(example_receipt["bundled_workbook"])
        self.shared_bundle_sha256 = example_receipt["bundled_sha256"]
        self.qa_shared_example_path = Path(example_receipt["fresh_copy"])
        self.mock_cas_values.update(example_receipt["bundled_cas_values"])
        self.transcript["automation_scopes"].append(
            {
                "scope": "isolated_first_run_shared_example_resolution",
                "profile_root": str(
                    self.run_root / "fresh-default-profile"
                ),
                "temporary_environment": {
                    "USERPROFILE": "QA-only profile_root",
                    "LOCALAPPDATA": "QA-only profile_root/LocalAppData",
                    "ER_PREDICTOR_PORTABLE": "0",
                    "sys.frozen": False,
                },
                "reason": (
                    "exercise resolve_shared_example_input() creation without "
                    "reading, replacing, or publishing beside the user's "
                    "distributed test.xlsx"
                ),
            }
        )
        geometry_receipt = inspect_endpoint_geometry_parity(
            app,
            self.run_root / "ui-parity-screenshots",
        )
        self.batch_recognition_widget_contract = geometry_receipt[
            "batch_recognition_widgets"
        ]
        self._record(
            "erta_eralpha_exact_ui_geometry_and_examples",
            examples=example_receipt,
            geometry=geometry_receipt,
        )

        route = tab._selected_route("single")
        spec = tab.selected_model_spec
        self._require(
            route == (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
            and spec is not None,
            "ERalpha does not expose its released classification model",
        )
        released_model = (
            tab.catalog_path.parent / spec.relative_path
        ).resolve(strict=True)
        released_ad = Path(tab.ad_ref_path_var.get()).resolve(strict=True)
        self._require(
            self._same_path(tab.model_path_var.get(), released_model)
            and self._same_path(tab.loaded_model_path_var.get(), released_model),
            "ERalpha startup did not select and load the released current model",
        )
        self._require(
            self._same_path(tab.ad_ref_path_var.get(), released_ad)
            and self._same_path(tab.loaded_ad_ref_path_var.get(), released_ad),
            "ERalpha startup did not select and load its approved AD reference",
        )
        self._require(
            _sha256_path(released_model) == spec.sha256.lower(),
            "released ERalpha model bytes do not match the catalog before QA",
        )

        self.parity_released_model = released_model
        self.parity_released_ad = released_ad
        self.parity_model_sha256 = _sha256_path(released_model)
        self.parity_ad_sha256 = _sha256_path(released_ad)
        self.parity_erta_before = self._erta_resource_state(app)
        self.parity_model_before = self._erba_model_state(tab)
        self.parity_options_before = bool(tab.options_visible)

        app.notebook.select(tab)
        if not tab.options_visible:
            tab.parity_widgets["options_button"].invoke()
        app.update_idletasks()
        self._require(
            tab.options_visible and tab.options_frame.winfo_ismapped(),
            "ERalpha Options did not open before resource interaction",
        )
        self.parity_model_browse = self._invoke_open_file_button(
            tab.parity_widgets["model_browse_button"],
            released_model,
            purpose="ERalpha released model Browse",
        )
        self._require(
            self._same_path(tab.model_path_var.get(), released_model)
            and self._erba_model_state(tab) == self.parity_model_before,
            "browsing the current released model activated or changed it before Reload",
        )
        self.parity_model_dialog_count = len(self.dialogs)
        tab.parity_widgets["model_reload_button"].invoke()
        self._require(
            tab.single_status_var.get().startswith("Loading model"),
            "ERalpha Reload model button did not start the actual callback",
        )

    def _model_reload_finished(self) -> bool:
        tab = self.app.eralpha_tab
        status = tab.single_status_var.get()
        if status.startswith("Error:"):
            raise RuntimeError(f"approved ERalpha model reload failed: {status}")
        return bool(
            tab._model_generation > self.parity_model_before["model_generation"]
            and self._same_path(
                tab.loaded_model_path_var.get(),
                self.parity_released_model,
            )
            and status
            == f"Model loaded: {tab.selected_model_spec.model_id}"
            and self._wait_dialog(self.parity_model_dialog_count)
        )

    def _finish_model_reload_and_start_ad(self) -> None:
        app = self.app
        tab = app.eralpha_tab
        model_after = self._erba_model_state(tab)
        self._require(
            model_after["predictor_identity"]
            != self.parity_model_before["predictor_identity"]
            and model_after["model_generation"]
            > self.parity_model_before["model_generation"]
            and model_after["loaded_route_keys"]
            == ["classification/er_alpha"],
            "ERalpha Reload model did not activate a fresh preflighted predictor",
        )
        self._require(
            _sha256_path(self.parity_released_model)
            == self.parity_model_sha256,
            "ERalpha released model bytes changed during Reload",
        )
        self.parity_model_after = model_after
        self.parity_model_dialog = self.dialogs[-1]

        self.parity_ad_browse = self._invoke_open_file_button(
            tab.parity_widgets["ad_browse_button"],
            self.parity_released_ad,
            purpose="ERalpha approved AD Browse",
        )
        self._require(
            self._same_path(
                tab.ad_ref_path_var.get(),
                self.parity_released_ad,
            ),
            "ERalpha AD Browse did not retain the approved route reference",
        )
        self.parity_ad_manager_before = id(tab.erba_ad)
        self.parity_ad_dialog_count = len(self.dialogs)
        tab.parity_widgets["ad_reload_button"].invoke()
        self._require(
            tab.single_status_var.get().startswith(
                "Rebuilding AD reference cache"
            ),
            "ERalpha Reload AD button did not start the actual callback",
        )

    def _ad_reload_finished(self) -> bool:
        tab = self.app.eralpha_tab
        status = tab.single_status_var.get()
        if status.startswith("Error:"):
            raise RuntimeError(f"approved ERalpha AD reload failed: {status}")
        return bool(
            self._same_path(
                tab.loaded_ad_ref_path_var.get(),
                self.parity_released_ad,
            )
            and status.startswith("AD fitted.")
            and self._wait_dialog(self.parity_ad_dialog_count)
        )

    def _finish_resource_parity_stage(self) -> None:
        app = self.app
        tab = app.eralpha_tab
        self._require(
            _sha256_path(self.parity_released_ad) == self.parity_ad_sha256,
            "ERalpha approved AD reference bytes changed during Reload",
        )
        self._require(
            id(tab.erba_ad) != self.parity_ad_manager_before,
            "ERalpha Reload AD did not activate a fresh endpoint-local AD manager",
        )
        ad_dialog = self.dialogs[-1]
        model_candidate = self.run_root / "unapproved-erba-model.joblib"
        ad_candidate = self.run_root / "unapproved-erba-ad.xlsx"
        model_candidate.write_bytes(b"native QA unapproved model candidate\n")
        ad_candidate.write_bytes(b"native QA unapproved AD candidate\n")
        try:
            approved_state = self._erba_model_state(tab)
            model_dialog_count = len(self.dialogs)
            rejected_model_browse = self._invoke_open_file_button(
                tab.parity_widgets["model_browse_button"],
                model_candidate,
                purpose="ERalpha unapproved model Browse",
            )
            rejected_model_status = tab.single_status_var.get()
            self._require(
                self._erba_model_state(tab) == approved_state,
                "unapproved ERalpha candidate changed or loaded model state",
            )
            self._require(
                self._wait_dialog(model_dialog_count)
                and rejected_model_status.startswith("Error:")
                and (
                    "Arbitrary joblib files are not allowed"
                    in rejected_model_status
                    or "integrity" in rejected_model_status.casefold()
                ),
                "unapproved ERalpha candidate was not rejected by Browse",
            )
            rejected_model_dialog = self.dialogs[-1]

            ad_selection_before = tab.ad_ref_path_var.get()
            ad_loaded_before = tab.loaded_ad_ref_path_var.get()
            ad_calculator = tab.erba_ad.calculator_for(
                ERBATask.CLASSIFICATION,
                ERBASubtype.ER_ALPHA,
            )
            ad_calculator_identity = id(ad_calculator)
            ad_dialog_count = len(self.dialogs)
            rejected_ad_browse = self._invoke_open_file_button(
                tab.parity_widgets["ad_browse_button"],
                ad_candidate,
                purpose="ERalpha unapproved AD Browse",
            )
            rejected_ad_status = tab.single_status_var.get()
            self._require(
                tab.ad_ref_path_var.get() == ad_selection_before
                and tab.loaded_ad_ref_path_var.get() == ad_loaded_before
                and id(
                    tab.erba_ad.calculator_for(
                        ERBATask.CLASSIFICATION,
                        ERBASubtype.ER_ALPHA,
                    )
                )
                == ad_calculator_identity,
                "unapproved ERalpha AD candidate changed or fitted AD state",
            )
            self._require(
                self._wait_dialog(ad_dialog_count)
                and rejected_ad_status.startswith("Error:")
                and "approved bundled AD reference" in rejected_ad_status,
                "unapproved ERalpha AD candidate was not rejected by Browse",
            )
            rejected_ad_dialog = self.dialogs[-1]
        finally:
            model_candidate.unlink(missing_ok=True)
            ad_candidate.unlink(missing_ok=True)

        erta_after = self._erta_resource_state(app)
        self._require(
            erta_after == self.parity_erta_before,
            "ERalpha Browse/Reload interactions leaked into ERTA resource state",
        )
        self._require(
            _sha256_path(self.parity_released_model)
            == self.parity_model_sha256,
            "released ERalpha model bytes changed during candidate rejection",
        )
        self._require(
            _sha256_path(self.parity_released_ad) == self.parity_ad_sha256,
            "approved ERalpha AD bytes changed during candidate rejection",
        )
        self._record(
            "eralpha_real_resource_buttons_release_enforcement_and_isolation",
            approved_model={
                "path": str(self.parity_released_model),
                "sha256": self.parity_model_sha256,
                "browse": self.parity_model_browse,
                "before": self.parity_model_before,
                "after": self.parity_model_after,
                "dialog": self.parity_model_dialog,
            },
            approved_ad={
                "path": str(self.parity_released_ad),
                "sha256": self.parity_ad_sha256,
                "browse": self.parity_ad_browse,
                "dialog": ad_dialog,
                "manager_before": self.parity_ad_manager_before,
                "manager_after": id(tab.erba_ad),
            },
            rejected_model={
                "candidate": str(model_candidate),
                "candidate_size_bytes": len(
                    b"native QA unapproved model candidate\n"
                ),
                "browse": rejected_model_browse,
                "status": rejected_model_status,
                "dialog": rejected_model_dialog,
                "model_state_unchanged": True,
            },
            rejected_ad={
                "candidate": str(ad_candidate),
                "candidate_size_bytes": len(
                    b"native QA unapproved AD candidate\n"
                ),
                "browse": rejected_ad_browse,
                "status": rejected_ad_status,
                "dialog": rejected_ad_dialog,
                "ad_state_unchanged": True,
            },
            erta_before=self.parity_erta_before,
            erta_after=erta_after,
            endpoint_isolation=True,
            approved_resource_bytes_unchanged=True,
        )
        if bool(tab.options_visible) != self.parity_options_before:
            tab.parity_widgets["options_button"].invoke()
        app.notebook.select(app.erta_tab)
        app.parity_widgets["mode_notebook"].select(
            app.parity_widgets["single_tab"]
        )
        app.update_idletasks()

    def tick(self) -> None:
        if self._tick_active:
            return
        self._tick_active = True
        try:
            app = self.app
            tab = app.eralpha_tab
            if self.stage == -4:
                if not self._parity_startup_ready():
                    self._reschedule()
                    return
                self._start_exact_ui_parity_stage()
                self.stage = -3
                self._reschedule()
                return
            if self.stage == -3:
                if not self._model_reload_finished():
                    self._reschedule()
                    return
                self._finish_model_reload_and_start_ad()
                self.stage = -2
                self._reschedule()
                return
            if self.stage == -2:
                if not self._ad_reload_finished():
                    self._reschedule()
                    return
                self._finish_resource_parity_stage()
                self.stage = 0
                self._reschedule()
                return
            if self.stage == 0:
                self._require(app.notebook.tab(app.notebook.select(), "text") == "ERTA", "ERTA is not startup tab")
                self._require(Path(app.model_path_var.get()).is_file(), "default ERTA model resource missing")
                self._require(Path(app.ad_ref_path_var.get()).is_file(), "default ERTA AD resource missing")
                self._require(app.options_frame.winfo_ismapped() == 0, "options must start hidden")
                self._require(str(app.run_batch_button.cget("state")) == "normal", "ERTA batch control unavailable")
                self._record("erta_startup_resources_and_controls", model=app.model_path_var.get(), ad_reference=app.ad_ref_path_var.get(), options_visible=False, batch_control_state=app.run_batch_button.cget("state"))
                app.toggle_options()
                app.update_idletasks()
                self._require(app.options_visible and bool(app.options_frame.grid_info()), "options did not open")
                self._require(app.model_path_var.get() and app.ad_ref_path_var.get(), "option resources not visible")
                app.toggle_options()
                app.update_idletasks()
                self._record("erta_options_and_resource_controls", model_entry=app.model_path_var.get(), ad_entry=app.ad_ref_path_var.get(), options_hidden_after_toggle=not app.options_visible and not bool(app.options_frame.grid_info()))
                self.stage = 1
            if self.stage == 1:
                if not app.predictor.is_loaded() or not app.ad_calculator.fitted:
                    self._reschedule()
                    return
                app.cas_var.set("")
                app.smiles_var.set(self.VALID_SMILES)
                app.single_predict_clicked()
                self.stage = 2
                self._reschedule()
                return
            if self.stage == 2:
                if not self._wait_erta_result():
                    self._reschedule()
                    return
                probability = float(app.active_probability_var.get().split(":", 1)[1].strip())
                self._require(abs(probability - 0.9801110625267029) < 1e-5, "ERTA exact valid oracle probability drift")
                self._require(app.ad_domain_var.get() in {"Applicability domain: In-domain", "Applicability domain: Out-of-domain"}, "ERTA AD was not rendered")
                self._require(bool(app.structure_label.cget("image")), "input molecule structure was not rendered")
                self._require(bool(app.nearest_reference_structure_label.cget("image")), "nearest reference structure was not rendered")
                self._require(bool(app.single_ad_graph_label.cget("image")), "single AD graph was not rendered")
                self._require(len(app.prob_canvas.find_all()) > 0, "probability graph canvas is empty")
                self.erta_summary = app.prediction_summary_var.get()
                self.erta_probability = app.active_probability_var.get()
                self._record("erta_exact_valid_single", probability_positive=probability, expected_probability_positive=0.9801110625267029, summary=self.erta_summary, ad=app.ad_domain_var.get(), nearest_reference=app.nearest_reference_var.get(), probability_canvas_items=len(app.prob_canvas.find_all()), input_structure=True, reference_structure=True, ad_graph=True)
                self.erta_previous_text = app.result_text.get("1.0", "end").strip()
                app.smiles_var.set(self.LEGACY_INVALID_SMILES)
                app.single_predict_clicked()
                self.stage = 3
                self._reschedule()
                return
            if self.stage == 3:
                if not self._wait_erta_result(self.erta_previous_text):
                    self._reschedule()
                    return
                detail = app.result_text.get("1.0", "end").strip()
                self._require("Mol valid: False" in detail, "legacy invalid SMILES did not reach ERTA zero-vector path")
                self._require(app.prediction_summary_var.get() in {"Prediction: Positive", "Prediction: Negative"}, "legacy invalid prediction label missing")
                self._record("erta_legacy_invalid_direct_smiles", smiles=self.LEGACY_INVALID_SMILES, rendered_summary=app.prediction_summary_var.get(), positive_probability=app.active_probability_var.get(), details=detail)
                self._install_mock_cas()
                app.cas_var.set(self.MOCK_CAS)
                app.smiles_var.set("")
                app.pubchem_clicked()
                self.stage = 4
                self._reschedule()
                return
            if self.stage == 4:
                if app.smiles_var.get() != self.MOCK_SMILES:
                    self._reschedule()
                    return
                self._require("PubChem found CID 712" in app.status_var.get(), "mocked PubChem callback did not update status")
                self._record("erta_mocked_pubchem_callback", cas=self.MOCK_CAS, smiles=app.smiles_var.get(), status=app.status_var.get())
                self.erta_previous_text = app.result_text.get("1.0", "end").strip()
                app.single_predict_clicked()
                self.stage = 5
                self._reschedule()
                return
            if self.stage == 5:
                if not self._wait_erta_result(self.erta_previous_text):
                    self._reschedule()
                    return
                self._require(f"CAS: {self.MOCK_CAS}" in app.result_text.get("1.0", "end"), "CAS-resolved prediction details missing CAS")
                self._record("erta_mocked_pubchem_prediction", summary=app.prediction_summary_var.get(), probability=app.active_probability_var.get(), canonical_details=app.result_text.get("1.0", "end").strip())
                self.stage = 6
            if self.stage == 6:
                template_path = self.run_root / "downloaded-template.xlsx"
                self._install_picker(template_path)
                model_before = app.model_path_var.get()
                ad_before = app.ad_ref_path_var.get()
                app.browse_model()
                app.browse_ad_reference()
                self._require(app.model_path_var.get() == model_before and app.ad_ref_path_var.get() == ad_before, "model or AD browse callback changed bundled resource selection")
                self._record("erta_model_ad_browse_callbacks", model=app.model_path_var.get(), ad_reference=app.ad_ref_path_var.get())
                dialogs_before = len(self.dialogs)
                app.download_template_clicked()
                self._require(template_path.is_file(), "template callback did not create workbook")
                self._require(self._wait_dialog(dialogs_before), "template callback did not expose completion dialog")
                template = load_workbook(template_path, read_only=True, data_only=True)
                try:
                    template_headers = [
                        cell.value for cell in template.active[1]
                    ]
                    self._require(
                        template.active.max_row == 1
                        and template_headers == ["CAS"],
                        "template schema drift",
                    )
                finally:
                    template.close()
                self._record(
                    "erta_template_callback_and_dialog",
                    output=str(template_path),
                    headers=template_headers,
                    dialog=self._terminal_dialog(
                        dialogs_before,
                        kind="info",
                        title="Template downloaded",
                        context="ERTA template download",
                    ),
                )
                self.erta_batch_input_path = (
                    self._copy_shared_example_for_batch("erta")
                )
                self.erta_batch_outputs_before = {
                    path.resolve()
                    for path in self.erta_batch_input_path.parent.glob(
                        "ERTA_*.xlsx"
                    )
                }
                app.batch_input_var.set(str(self.erta_batch_input_path))
                app.batch_input_display_var.set(
                    self.erta_batch_input_path.name
                )
                app.batch_destination_var.set(str(self.erta_batch_input_path.parent))
                self.batch_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "erta",
                    "shared_success",
                )
                app.batch_predict_clicked()
                self.erta_success_locked_controls = (
                    self._batch_control_states(app)
                )
                self._require(
                    app._batch_active
                    and self.erta_success_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERTA shared-example batch did not lock its controls",
                )
                self._record_batch_new_run("erta", screenshot=True)
                self.stage = 7
                self._reschedule()
                return
            if self.stage == 7:
                if (
                    app._batch_active
                    or not self._wait_dialog(self.batch_dialog_count)
                ):
                    self._reschedule()
                    return
                self._require(
                    self.erta_batch_input_path is not None,
                    "ERTA batch input path was not recorded",
                )
                input_parent = self.erta_batch_input_path.parent
                outputs = {
                    path.resolve()
                    for path in input_parent.glob(
                        "ERTA_*.xlsx"
                    )
                } - self.erta_batch_outputs_before
                self._require(
                    len(outputs) == 1,
                    f"expected one fresh ERTA workbook, found {len(outputs)}",
                )
                output = outputs.pop()
                self._require(
                    output.parent == input_parent
                    and output.name.startswith("ERTA_test_")
                    and output.name.endswith("_prediction.xlsx"),
                    "ERTA workbook path or ERTA_ filename contract drift",
                )
                self.erta_batch_output_path = output
                workbook = load_workbook(output, read_only=True, data_only=True)
                try:
                    sheet = workbook.active
                    headers = [cell.value for cell in sheet[1]]
                    rows = [
                        dict(zip(headers, row))
                        for row in sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                    self._require(
                        sheet.title == "Sheet1",
                        "legacy ERTA workbook sheet name drift",
                    )
                finally:
                    workbook.close()
                required_headers = {
                    "CAS",
                    "SMILES",
                    "Canonical_SMILES",
                    "Mol_valid",
                    "Probability_Negative_0",
                    "Probability_Positive_1",
                    "Prediction",
                    "Prediction_label",
                    "Decision_rule",
                    *self.ERBA_BATCH_AD_COLUMNS,
                    "PubChem_CID",
                    "PubChem_status",
                }
                self._require(
                    required_headers <= set(headers),
                    "ERTA shared-example workbook schema is incomplete",
                )
                expected_cas = _read_shared_example_workbook(
                    self.erta_batch_input_path
                )["cas_values"]
                self._require(
                    len(rows) == 25
                    and [str(row["CAS"]) for row in rows] == expected_cas,
                    "ERTA CARSRN alias or shared-example input order drift",
                )
                self._require(
                    all(
                        row["SMILES"] == self.MOCK_SMILES
                        and row["PubChem_status"] == "Found"
                        and row["Mol_valid"] is True
                        and row["Probability_Negative_0"] is not None
                        and row["Probability_Positive_1"] is not None
                        and row["Prediction_label"] in {"Positive", "Negative"}
                        for row in rows
                    ),
                    "ERTA shared CARSRN example did not yield 25 nonblank predictions",
                )
                graph_paths = sorted(
                    path.name
                    for path in (input_parent / "graphs").glob("*.png")
                )
                self._require(graph_paths, "ERTA batch graph files were not generated")
                summary = app.batch_result.get("1.0", "end").strip()
                self._require(
                    f"Output workbook: {output}" in summary
                    and "AD In-domain:" in summary
                    and "Graph directory:" in summary,
                    "ERTA batch completion summary is incomplete",
                )
                self._require(
                    not hasattr(app, "batch_tree")
                    and not hasattr(app, "graph_combo")
                    and not hasattr(app, "graph_label"),
                    "obsolete ERTA batch table/graph preview remains in the UI",
                )
                completion_dialog = self._terminal_dialog(
                    self.batch_dialog_count,
                    kind="info",
                    title="Batch prediction done",
                    context="ERTA shared-example batch success",
                )
                self._require(
                    str(output) in completion_dialog["message"]
                    and completion_dialog[
                        "batch_controls_at_dialog"
                    ]["erta"]
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    },
                    "ERTA success dialog does not report the published workbook",
                )
                self._require(
                    self._batch_control_states(app)
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    }
                    and app.batch_progress_var.get()
                    == "100% - 25/25 - Completed"
                    and int(app.batch_progress_value.get()) == 100,
                    "ERTA batch controls were not restored after success",
                )
                self.erta_success_result_identity = id(app.last_batch_result)
                recognition = self._record_batch_terminal(
                    "erta",
                    "success",
                    dialog=completion_dialog,
                    output=output,
                    output_exists=output.is_file(),
                    total_count=len(rows),
                    available_count=sum(
                        row["Mol_valid"] is True
                        for row in rows
                    ),
                    unavailable_count=sum(
                        row["Mol_valid"] is not True
                        for row in rows
                    ),
                    workbook_unavailable_count=sum(
                        row["Mol_valid"] is not True
                        for row in rows
                    ),
                    fresh_output_count=1,
                    screenshot=True,
                )
                self._record(
                    "erta_shared_test_batch_success_dialog_and_predictions",
                    input=str(self.erta_batch_input_path),
                    output=str(output),
                    sheet=sheet.title,
                    headers=headers,
                    input_order=expected_cas,
                    input_header="CARSRN",
                    row_count=len(rows),
                    predicted_row_count=sum(
                        row["Probability_Positive_1"] is not None
                        for row in rows
                    ),
                    mocked_pubchem_statuses=sorted(
                        {row["PubChem_status"] for row in rows}
                    ),
                    deterministic_smiles_substitution=self.MOCK_SMILES,
                    source_bytes_preserved=(
                        _sha256_path(self.qa_shared_example_path)
                        == self.shared_bundle_sha256
                    ),
                    graph_files=graph_paths,
                    batch_summary=summary,
                    progress=app.batch_progress_var.get(),
                    controls_during_run=self.erta_success_locked_controls,
                    controls=self._batch_control_states(app),
                    completion_dialog=completion_dialog,
                    recognition=recognition,
                )
                self.stage = 72
            if self.stage == 72:
                mixed_directory = self.run_root / "mixed-batch" / "erta"
                mixed_directory.mkdir(parents=True, exist_ok=False)
                mixed_input = mixed_directory / "erta-mixed-input.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["No.", "CAS", "Chemical Name", "SMILES"])
                sheet.append([2, "", "valid", self.VALID_SMILES])
                sheet.append([1, self.MOCK_CAS, "mocked-cas", ""])
                sheet.append(
                    [3, "", "legacy-invalid", self.LEGACY_INVALID_SMILES]
                )
                sheet.append(
                    [
                        4,
                        self.MOCK_UNAVAILABLE_CAS_VALUES[0],
                        "pubchem-unavailable",
                        "",
                    ]
                )
                workbook.save(mixed_input)
                workbook.close()
                self.erta_mixed_input_path = mixed_input.resolve(
                    strict=True
                )
                self.erta_mixed_outputs_before = {
                    path.resolve()
                    for path in mixed_directory.glob("ERTA_*.xlsx")
                }
                app.batch_input_var.set(str(self.erta_mixed_input_path))
                app.batch_input_display_var.set(
                    self.erta_mixed_input_path.name
                )
                app.batch_destination_var.set(str(mixed_directory))
                self.erta_mixed_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "erta",
                    "partial_unavailable",
                )
                app.batch_predict_clicked()
                self.erta_mixed_locked_controls = (
                    self._batch_control_states(app)
                )
                self._require(
                    app._batch_active
                    and self.erta_mixed_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERTA mixed batch did not start with locked controls",
                )
                self._record_batch_new_run("erta")
                self.stage = 73
                self._reschedule()
                return
            if self.stage == 73:
                if (
                    app._batch_active
                    or not self._wait_dialog(
                        self.erta_mixed_dialog_count
                    )
                ):
                    self._reschedule()
                    return
                mixed_parent = self.erta_mixed_input_path.parent
                mixed_outputs = {
                    path.resolve()
                    for path in mixed_parent.glob("ERTA_*.xlsx")
                } - self.erta_mixed_outputs_before
                self._require(
                    len(mixed_outputs) == 1,
                    f"expected one fresh mixed ERTA workbook, found {len(mixed_outputs)}",
                )
                mixed_output = mixed_outputs.pop()
                self._require(
                    mixed_output.parent == mixed_parent
                    and mixed_output.name.startswith(
                        "ERTA_erta-mixed-input_"
                    ),
                    "mixed ERTA workbook path or filename prefix drift",
                )
                workbook = load_workbook(
                    mixed_output,
                    read_only=True,
                    data_only=True,
                )
                try:
                    sheet = workbook.active
                    mixed_headers = [
                        cell.value for cell in sheet[1]
                    ]
                    mixed_rows = list(
                        sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    )
                    mixed_dict_rows = [
                        dict(zip(mixed_headers, row))
                        for row in mixed_rows
                    ]
                    self._require(
                        sheet.title == "Sheet1",
                        "legacy ERTA workbook sheet name drift",
                    )
                finally:
                    workbook.close()
                expected_headers = [
                    "No.",
                    "CAS",
                    "Chemical Name",
                    "SMILES",
                    "Canonical_SMILES",
                    "Mol_valid",
                    "Probability_Negative_0",
                    "Probability_Positive_1",
                    "Prediction",
                    "Prediction_label",
                    "Decision_rule",
                    *self.ERBA_BATCH_AD_COLUMNS,
                    "PubChem_CID",
                    "PubChem_status",
                ]
                self._require(
                    mixed_headers == expected_headers,
                    "legacy ERTA mixed workbook schema/order drift",
                )
                self._require(
                    [row[0] for row in mixed_rows] == [2, 1, 3, 4],
                    "ERTA mixed batch input order drift",
                )
                self._require(
                    mixed_rows[1][3] == self.MOCK_SMILES
                    and mixed_rows[1][-1] == "Found",
                    "mocked CAS mixed-batch row was not resolved",
                )
                self._require(
                    mixed_rows[2][5] is False,
                    "legacy invalid mixed-batch row lost its invalid marker",
                )
                self._require(
                    mixed_dict_rows[3]["CAS"]
                    == self.MOCK_UNAVAILABLE_CAS_VALUES[0]
                    and str(
                        mixed_dict_rows[3].get("PubChem_status") or ""
                    ).startswith("Not found:")
                    and mixed_dict_rows[3]["Mol_valid"] is False
                    and mixed_dict_rows[3][
                        "Probability_Negative_0"
                    ]
                    is not None
                    and mixed_dict_rows[3][
                        "Probability_Positive_1"
                    ]
                    is not None
                    and mixed_dict_rows[3]["Prediction_label"]
                    in {"Positive", "Negative"},
                    "ERTA mixed batch did not retain its unavailable row",
                )
                mixed_graphs = sorted(
                    path.name
                    for path in (mixed_parent / "graphs").glob("*.png")
                )
                mixed_summary = app.batch_result.get(
                    "1.0", "end"
                ).strip()
                self._require(
                    mixed_graphs
                    and f"Output workbook: {mixed_output}" in mixed_summary
                    and f"{BATCH_UNAVAILABLE_LABEL}: 2"
                    in mixed_summary
                    and "AD In-domain:" in mixed_summary
                    and "Graph directory:" in mixed_summary,
                    "ERTA mixed batch artifacts or completion summary are incomplete",
                )
                mixed_dialog = self._terminal_dialog(
                    self.erta_mixed_dialog_count,
                    kind="info",
                    title="Batch prediction done",
                    context="ERTA mixed batch success",
                )
                restored_controls = self._batch_control_states(app)
                self._require(
                    restored_controls
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    }
                    and app.batch_progress_var.get()
                    == "100% - 4/4 - Completed"
                    and int(app.batch_progress_value.get()) == 100,
                    "ERTA mixed batch did not restore controls and progress",
                )
                self._require(
                    mixed_dialog["batch_controls_at_dialog"]["erta"]
                    == restored_controls,
                    "ERTA mixed-batch controls were not restored before its dialog",
                )
                self.erta_success_result_identity = id(
                    app.last_batch_result
                )
                mixed_recognition = self._record_batch_terminal(
                    "erta",
                    "partial_unavailable",
                    dialog=mixed_dialog,
                    output=mixed_output,
                    output_exists=mixed_output.is_file(),
                    total_count=len(mixed_rows),
                    available_count=2,
                    unavailable_count=2,
                    workbook_unavailable_count=sum(
                        row["Mol_valid"] is not True
                        for row in mixed_dict_rows
                    ),
                    fresh_output_count=1,
                )
                self._record(
                    "erta_batch_workbook_graphs_and_order",
                    input=str(self.erta_mixed_input_path),
                    output=str(mixed_output),
                    sheet=sheet.title,
                    headers=mixed_headers,
                    input_order=[row[0] for row in mixed_rows],
                    mocked_cas_status=mixed_rows[1][-1],
                    invalid_row_mol_valid=mixed_rows[2][5],
                    unavailable_row={
                        key: mixed_dict_rows[3].get(key)
                        for key in (
                            "CAS",
                            "PubChem_status",
                            "Mol_valid",
                            "Probability_Negative_0",
                            "Probability_Positive_1",
                            "Prediction_label",
                        )
                    },
                    graph_files=mixed_graphs,
                    batch_summary=mixed_summary,
                    progress=app.batch_progress_var.get(),
                    controls_during_run=self.erta_mixed_locked_controls,
                    controls=restored_controls,
                    completion_dialog=mixed_dialog,
                    recognition=mixed_recognition,
                )
                self.erta_mixed_check = self.transcript["checks"][-1]
                self.stage = 74
            if self.stage == 74:
                unavailable_directory = (
                    self.run_root / "all-unavailable-batch" / "erta"
                )
                unavailable_directory.mkdir(parents=True, exist_ok=False)
                unavailable_input = (
                    unavailable_directory / "all-unavailable-input.xlsx"
                )
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["CAS"])
                for cas in self.MOCK_UNAVAILABLE_CAS_VALUES:
                    sheet.append([cas])
                workbook.save(unavailable_input)
                workbook.close()
                self.erta_unavailable_input_path = (
                    unavailable_input.resolve(strict=True)
                )
                self.erta_unavailable_outputs_before = {
                    path.resolve()
                    for path in unavailable_directory.glob("ERTA_*.xlsx")
                }
                app.batch_input_var.set(
                    str(self.erta_unavailable_input_path)
                )
                app.batch_input_display_var.set(
                    self.erta_unavailable_input_path.name
                )
                app.batch_destination_var.set(
                    str(unavailable_directory)
                )
                self.erta_unavailable_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "erta",
                    "all_unavailable",
                )
                app.batch_predict_clicked()
                self.erta_unavailable_locked_controls = (
                    self._batch_control_states(app)
                )
                self._require(
                    app._batch_active
                    and self.erta_unavailable_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERTA all-unavailable batch did not start with locked "
                    "controls",
                )
                self._record_batch_new_run("erta")
                self.stage = 75
                self._reschedule()
                return
            if self.stage == 75:
                if (
                    app._batch_active
                    or not self._wait_dialog(
                        self.erta_unavailable_dialog_count
                    )
                ):
                    self._reschedule()
                    return
                unavailable_parent = (
                    self.erta_unavailable_input_path.parent
                )
                unavailable_outputs = {
                    path.resolve()
                    for path in unavailable_parent.glob("ERTA_*.xlsx")
                } - self.erta_unavailable_outputs_before
                self._require(
                    len(unavailable_outputs) == 1,
                    "expected one fresh all-unavailable ERTA workbook, "
                    f"found {len(unavailable_outputs)}",
                )
                unavailable_output = unavailable_outputs.pop()
                workbook = load_workbook(
                    unavailable_output,
                    read_only=True,
                    data_only=True,
                )
                try:
                    unavailable_sheet = workbook.active
                    unavailable_headers = [
                        cell.value for cell in unavailable_sheet[1]
                    ]
                    unavailable_rows = [
                        dict(zip(unavailable_headers, row))
                        for row in unavailable_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                finally:
                    workbook.close()
                self._require(
                    len(unavailable_rows)
                    == len(self.MOCK_UNAVAILABLE_CAS_VALUES)
                    and [
                        str(row["CAS"])
                        for row in unavailable_rows
                    ]
                    == list(self.MOCK_UNAVAILABLE_CAS_VALUES)
                    and all(
                        str(row.get("PubChem_status") or "").startswith(
                            "Not found:"
                        )
                        and row["Mol_valid"] is False
                        and row["Probability_Negative_0"] is not None
                        and row["Probability_Positive_1"] is not None
                        and row["Prediction_label"]
                        in {"Positive", "Negative"}
                        for row in unavailable_rows
                    ),
                    "ERTA all-unavailable result did not retain explicit "
                    "Mol_valid/PubChem unavailable row markers",
                )
                unavailable_summary = self._batch_result_value(app)
                unavailable_dialog = self._terminal_dialog(
                    self.erta_unavailable_dialog_count,
                    kind="info",
                    title="Batch prediction done",
                    context="ERTA all-unavailable batch publication",
                )
                restored_controls = self._batch_control_states(app)
                unavailable_count = len(unavailable_rows)
                self._require(
                    f"{BATCH_UNAVAILABLE_LABEL}: {unavailable_count}"
                    in unavailable_summary
                    and f"{BATCH_UNAVAILABLE_LABEL}: "
                    f"{unavailable_count}"
                    in unavailable_dialog["message"]
                    and app.batch_progress_var.get()
                    == (
                        f"100% - {unavailable_count}/"
                        f"{unavailable_count} - Completed"
                    )
                    and restored_controls
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    },
                    "ERTA all-unavailable publication feedback drift",
                )
                unavailable_recognition = self._record_batch_terminal(
                    "erta",
                    "all_unavailable",
                    dialog=unavailable_dialog,
                    output=unavailable_output,
                    output_exists=unavailable_output.is_file(),
                    total_count=unavailable_count,
                    available_count=0,
                    unavailable_count=unavailable_count,
                    workbook_unavailable_count=sum(
                        row["Mol_valid"] is not True
                        for row in unavailable_rows
                    ),
                    fresh_output_count=1,
                )
                self.erta_mixed_check["all_unavailable"] = {
                    "input": str(self.erta_unavailable_input_path),
                    "output": str(unavailable_output),
                    "row_count": unavailable_count,
                    "unavailable_count": unavailable_count,
                    "rows": unavailable_rows,
                    "batch_summary": unavailable_summary,
                    "completion_dialog": unavailable_dialog,
                    "progress": app.batch_progress_var.get(),
                    "controls_during_run": (
                        self.erta_unavailable_locked_controls
                    ),
                    "controls": restored_controls,
                    "recognition": unavailable_recognition,
                }
                self.erta_success_result_identity = id(
                    app.last_batch_result
                )
                self.stage = 70
            if self.stage == 70:
                failure_input = self._make_corrupt_batch_input("erta")
                self.erta_failure_input_path = failure_input
                self.erta_failure_outputs_before = {
                    path.resolve()
                    for path in failure_input.parent.glob("ERTA_*.xlsx")
                }
                self.erta_failure_graphs_before = {
                    path.resolve()
                    for path in failure_input.parent.glob("graphs")
                }
                app.batch_input_var.set(str(failure_input))
                app.batch_input_display_var.set(failure_input.name)
                app.batch_destination_var.set(str(failure_input.parent))
                self.erta_failure_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "erta",
                    "no_output_failure",
                )
                app.batch_predict_clicked()
                self.erta_failure_locked_controls = (
                    self._batch_control_states(app)
                )
                self._require(
                    app._batch_active
                    and self.erta_failure_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERTA deterministic failure did not start with locked controls",
                )
                self._record_batch_new_run("erta")
                self.stage = 71
                self._reschedule()
                return
            if self.stage == 71:
                if (
                    app._batch_active
                    or not self._wait_dialog(
                        self.erta_failure_dialog_count
                    )
                ):
                    self._reschedule()
                    return
                failure_dialog = self._terminal_dialog(
                    self.erta_failure_dialog_count,
                    kind="error",
                    title="Batch prediction failed",
                    context="ERTA deterministic batch failure",
                )
                failure_summary = app.batch_result.get(
                    "1.0", "end"
                ).strip()
                new_outputs = {
                    path.resolve()
                    for path in self.erta_failure_input_path.parent.glob(
                        "ERTA_*.xlsx"
                    )
                } - self.erta_failure_outputs_before
                restored_controls = self._batch_control_states(app)
                self._require(
                    restored_controls
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    }
                    and app.batch_progress_var.get()
                    == "100% - 0/0 - Failed"
                    and int(app.batch_progress_value.get()) == 100
                    and app.status_var.get().startswith(
                        "Batch prediction failed:"
                    )
                    and failure_summary.startswith(
                        "Batch prediction failed."
                    )
                    and "Batch job completed and workbook saved."
                    not in failure_summary
                    and "Output workbook:" not in failure_summary
                    and not new_outputs
                    and {
                        path.resolve()
                        for path in self.erta_failure_input_path.parent.glob(
                            "graphs"
                        )
                    }
                    == self.erta_failure_graphs_before
                    and id(app.last_batch_result)
                    == self.erta_success_result_identity
                    and failure_dialog[
                        "batch_controls_at_dialog"
                    ]["erta"]
                    == restored_controls,
                    "ERTA failure exposed false success or did not restore controls",
                )
                failure_recognition = self._record_batch_terminal(
                    "erta",
                    "no_output_failure",
                    dialog=failure_dialog,
                    output=None,
                    output_exists=False,
                    fresh_output_count=len(new_outputs),
                )
                self._record(
                    "erta_batch_failure_dialog_and_control_recovery",
                    input=str(self.erta_failure_input_path),
                    fixture_kind="intentionally corrupt xlsx bytes",
                    dialog=failure_dialog,
                    controls_during_run=self.erta_failure_locked_controls,
                    controls=restored_controls,
                    progress=app.batch_progress_var.get(),
                    status=app.status_var.get(),
                    result=failure_summary,
                    fresh_outputs=[],
                    fresh_graph_directories=[],
                    success_dialog_emitted=False,
                    prior_success_result_retained=True,
                    recognition=failure_recognition,
                )
                app.notebook.select(tab)
                self.stage = 8
            if self.stage == 8:
                task, subtype = self.routes[self.route_index]
                app.notebook.select(tab)
                if self.route_index == 0:
                    tab.mode_notebook.select(tab.single_tab)
                    app.update_idletasks()
                    evidence_details = tab.single_result_text.get(
                        "1.0",
                        "end",
                    ).strip()
                    expected_caveat = (
                        "Evidence caveat: "
                        f"{HISTORICAL_EXPOSURE_CAVEAT_COMPACT}"
                    )
                    self._require(
                        tab.single_result_text.winfo_ismapped()
                        and expected_caveat in evidence_details,
                        "ERalpha details pane does not visibly present the "
                        "compact evidence caveat",
                    )
                    self._record(
                        "erba_evidence_caveat_and_dialog",
                        presentation="single_prediction_details_pane",
                        compact_caveat=HISTORICAL_EXPOSURE_CAVEAT_COMPACT,
                        rendered_details=evidence_details,
                        details_box=_widget_geometry_receipt(
                            tab.single_result_text,
                            tab,
                        ),
                    )
                tab._route_changed("single")
                tab.smiles_var.set("Oc1ccccc1")
                tab.cas_var.set("")
                tab.single_predict_clicked()
                self.stage = 9
                self._reschedule()
                return
            if self.stage == 9:
                if (
                    tab._active_single_request_id is not None
                    or tab._active_ad_request_id is not None
                    or "Evaluating" in tab.single_ad_domain_var.get()
                ):
                    self._reschedule()
                    return
                body = self._body(tab)
                task, subtype = self.routes[self.route_index]
                self._require(tab.single_status_var.get() == "ok" and body, "ERBA route did not render success")
                self._require(
                    tab.single_ad_domain_var.get()
                    in (
                        "Applicability domain: In-domain",
                        "Applicability domain: Out-of-domain",
                    ),
                    "ERBA route did not complete applicability-domain analysis",
                )
                self._require(
                    bool(tab._single_ad_graph_image)
                    and "Morgan similarity=" in tab.single_nearest_reference_var.get(),
                    "ERBA route did not render AD graph and nearest reference",
                )
                self._record("erba_route", task=task.value, subtype=subtype.value, status=tab.single_status_var.get(), rendered_result=body)
                self.route_index += 1
                if self.route_index < len(self.routes):
                    self.stage = 8
                    self._reschedule()
                    return
                self.stage = 10
            if self.stage == 10:
                tab = app.eralpha_tab
                app.notebook.select(tab)
                name, smiles, cas, _expected = self.invalid_cases[self.invalid_index]
                tab.single_task_var.set(ERBATask.CLASSIFICATION.value)
                tab.single_subtype_var.set(ERBASubtype.ER_ALPHA.value)
                tab.smiles_var.set(smiles)
                tab.cas_var.set(cas)
                tab.single_predict_clicked()
                self.stage = 11
                self._reschedule()
                return
            if self.stage == 11:
                if tab._active_single_request_id is not None:
                    self._reschedule()
                    return
                body = self._body(tab)
                self._require(tab.single_status_var.get() != "ok", f"ERBA invalid {self.invalid_cases[self.invalid_index][0]} succeeded")
                self._record(f"erba_invalid_{self.invalid_cases[self.invalid_index][0]}", rendered_result=body, rendered_status=tab.single_status_var.get())
                self.invalid_index += 1
                if self.invalid_index < len(self.invalid_cases):
                    self.stage = 10
                    self._reschedule()
                    return
                self.stage = 12
            if self.stage == 12:
                app.notebook.select(tab)
                tab.notebook.select(tab.batch_tab)
                tab.batch_task_var.set(ERBATask.CLASSIFICATION.value)
                tab.batch_subtype_var.set(ERBASubtype.ER_ALPHA.value)
                self.erba_batch_input_path = (
                    self._copy_shared_example_for_batch("eralpha")
                )
                input_parent = self.erba_batch_input_path.parent
                self.erba_batch_outputs_before = {
                    path.resolve()
                    for path in input_parent.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                }
                self.erba_batch_graph_dirs_before = {
                    path.resolve()
                    for path in input_parent.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                }
                tab.batch_input_var.set(str(self.erba_batch_input_path))
                tab.batch_input_display_var.set(
                    self.erba_batch_input_path.name
                )
                tab.batch_destination_var.set(str(input_parent))
                self.erba_batch_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "eralpha",
                    "shared_success",
                )
                tab.batch_predict_clicked()
                self.erba_success_locked_controls = (
                    self._batch_control_states(tab)
                )
                self._require(
                    tab._active_batch_request_id is not None
                    and self.erba_success_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERalpha shared-example batch did not start with locked "
                    f"controls: {tab.batch_status_var.get()}",
                )
                self._record_batch_new_run(
                    "eralpha",
                    screenshot=True,
                )
                self.stage = 13
                self._reschedule()
                return
            if self.stage == 13:
                if (
                    tab._active_batch_request_id is not None
                    or not self._wait_dialog(
                        self.erba_batch_dialog_count
                    )
                ):
                    self._reschedule()
                    return
                self._require(
                    self.erba_batch_input_path is not None,
                    "ERalpha batch input path was not recorded",
                )
                input_parent = self.erba_batch_input_path.parent
                outputs = {
                    path.resolve()
                    for path in input_parent.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                } - self.erba_batch_outputs_before
                self._require(
                    len(outputs) == 1,
                    f"expected one new ERalpha workbook, found {len(outputs)}: "
                    f"{tab.batch_status_var.get()}",
                )
                output = outputs.pop()
                self._require(
                    output.parent == input_parent
                    and tab.batch_destination_var.get() == str(input_parent),
                    "ERalpha workbook destination is not the input workbook parent",
                )
                graph_dirs = {
                    path.resolve()
                    for path in input_parent.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                } - self.erba_batch_graph_dirs_before
                self._require(
                    len(graph_dirs) == 1,
                    f"expected one new input-adjacent ERalpha graph directory, found {len(graph_dirs)}",
                )
                graph_directory = graph_dirs.pop()
                graph_files = {
                    path.name for path in graph_directory.iterdir() if path.is_file()
                }
                self._require(
                    graph_directory.parent == input_parent
                    and graph_files == self.ERBA_BATCH_GRAPH_FILES,
                    "ERalpha batch graphs are incomplete or outside the input parent",
                )

                workbook = load_workbook(
                    output,
                    read_only=True,
                    data_only=True,
                )
                try:
                    sheet_names = workbook.sheetnames
                    self._require(
                        sheet_names == list(ERALPHA_WORKBOOK_SHEET_ORDER)
                        and workbook.active.title == "Predictions",
                        "ERalpha batch workbook is not Predictions-first and active",
                    )
                    prediction_sheet = workbook["Predictions"]
                    prediction_headers = [
                        cell.value for cell in prediction_sheet[1]
                    ]
                    self._require_eralpha_primary_headers(
                        prediction_headers,
                        (),
                        "ERalpha shared-example workbook",
                    )
                    prediction_rows = [
                        dict(zip(prediction_headers, row))
                        for row in prediction_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                    diagnostics_sheet = workbook["Diagnostics"]
                    diagnostics_headers = [
                        cell.value for cell in diagnostics_sheet[1]
                    ]
                    diagnostics_rows = [
                        dict(zip(diagnostics_headers, row))
                        for row in diagnostics_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                    input_sheet = workbook["Input"]
                    input_headers = [
                        cell.value for cell in input_sheet[1]
                    ]
                    input_rows = list(
                        input_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    )
                    guide_sheet = workbook["Guide"]
                    guide_rows = list(
                        guide_sheet.iter_rows(values_only=True)
                    )
                    metadata_sheet = workbook["Metadata"]
                    metadata_headers = [
                        cell.value for cell in metadata_sheet[1]
                    ]
                    metadata_rows = list(
                        metadata_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    )
                    self._require(
                        len(metadata_rows) == 1,
                        "ERalpha shared-example Metadata must contain "
                        "one execution-level row",
                    )
                    metadata = dict(
                        zip(metadata_headers, metadata_rows[0])
                    )
                finally:
                    workbook.close()

                expected_cas = _read_shared_example_workbook(
                    self.erba_batch_input_path
                )["cas_values"]
                self._require(
                    input_headers == ["CARSRN"]
                    and input_rows
                    == [(cas,) for cas in expected_cas]
                    and len(input_rows) == 25,
                    "ERalpha Input sheet did not preserve the supplied CARSRN workbook",
                )
                self._require(
                    len(prediction_rows) == 25
                    and [str(row["CAS"]) for row in prediction_rows]
                    == expected_cas,
                    "ERalpha primary predictions lost CARSRN-to-CAS row order",
                )
                self._require_eralpha_diagnostics(
                    diagnostics_headers,
                    diagnostics_rows,
                    prediction_rows,
                    "ERalpha shared-example workbook",
                )
                self._require(
                    all(
                        row["SMILES"] == self.MOCK_SMILES
                        and row["Canonical_SMILES"] == self.MOCK_SMILES
                        and row["Mol_valid"] is True
                        and self._eralpha_prediction_is_consistent(row)
                        and row["AD"] in {"In-domain", "Out-of-domain"}
                        and str(row["PubChem_CID"]) == "712"
                        and row["PubChem_status"] == "Found"
                        for row in prediction_rows
                    ),
                    "ERalpha shared CARSRN example did not preserve its "
                    "numeric binding predictions or captured PubChem fields",
                )
                self._require(
                    [
                        row["Row_ID"] for row in diagnostics_rows
                    ]
                    == [str(index) for index in range(25)]
                    and [
                        str(row["CAS"]) for row in diagnostics_rows
                    ]
                    == expected_cas
                    and all(
                        row["Status_Code"] == "ok"
                        and not row["Status_Message"]
                        and row["SMILES_Provenance"]
                        == "pubchem_lookup"
                        and row["Result_Status"] == "Predicted"
                        for row in diagnostics_rows
                    ),
                    "ERalpha shared-example Diagnostics lost successful "
                    "row status or PubChem provenance",
                )
                self._require(
                    guide_rows
                    and guide_rows[0] == ("Topic", "Details")
                    and any(
                        row[0] == "ERalpha batch prediction guide"
                        for row in guide_rows[1:]
                    ),
                    "ERalpha Guide sheet content is missing",
                )

                expected_spec = tab.catalog[
                    (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
                ]
                hex_digits = set("0123456789abcdef")
                self._require(
                    metadata_headers
                    == list(ERBA_CLASSIFICATION_METADATA_COLUMNS)
                    and metadata.get("Excel_Contract_ID")
                    == ERBA_BINDING_CLASSIFICATION_EXCEL_CONTRACT_ID
                    and metadata.get("Workflow") == "erba"
                    and metadata.get("Task") == "classification"
                    and metadata.get("Subtype") == "er_alpha"
                    and metadata.get("Decision_rule")
                    == "binding_probability>=0.5"
                    and metadata.get("Model_ID") == expected_spec.model_id
                    and metadata.get("Model_SHA256") == expected_spec.sha256
                    and metadata.get("Preprocessing_Policy_ID")
                    == CLASSIFICATION_PARENT_POLICY_ID
                    and metadata.get("Performance_Evidence_Scope")
                    == "internal_historically_exposed"
                    and bool(metadata.get("Evidence_Caveat"))
                    and all(
                        isinstance(metadata.get(column), str)
                        and len(metadata[column]) == 64
                        and set(metadata[column]) <= hex_digits
                        for column in self.ERBA_PROVENANCE_COLUMNS
                    ),
                    "ERalpha batch metadata did not preserve catalog model provenance",
                )

                binding_count = sum(
                    row["Prediction_label"] == "Binding"
                    for row in prediction_rows
                )
                non_binding_count = sum(
                    row["Prediction_label"] == "Non-binding"
                    for row in prediction_rows
                )
                not_predicted_count = sum(
                    row["Result_Status"] != "Predicted"
                    for row in diagnostics_rows
                )
                self._require(
                    binding_count + non_binding_count == 25
                    and not_predicted_count == 0,
                    "ERalpha shared example contains blank or not-predicted results",
                )
                summary = tab.batch_result.get("1.0", "end").strip()
                self._require(
                    tab.batch_status_var.get()
                    == f"Batch prediction completed: {output}"
                    and f"Total rows: {len(prediction_rows)}" in summary
                    and f"Binding: {binding_count}" in summary
                    and f"Non-binding: {non_binding_count}" in summary
                    and f"Not predicted: {not_predicted_count}" in summary
                    and f"Output workbook: {output}" in summary
                    and f"Graph files: {len(graph_files)}" in summary
                    and f"Graph directory: {graph_directory}" in summary,
                    "ERalpha batch completion did not expose binding and artifact summaries",
                )
                completion_dialog = self._terminal_dialog(
                    self.erba_batch_dialog_count,
                    kind="info",
                    title="Batch prediction done",
                    context="ERalpha shared-example batch success",
                )
                self._require(
                    str(output) in completion_dialog["message"]
                    and "Not predicted: 0" in completion_dialog["message"]
                    and completion_dialog[
                        "batch_controls_at_dialog"
                    ]["eralpha"]
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    },
                    "ERalpha success dialog does not report the complete publication",
                )
                restored_controls = self._batch_control_states(tab)
                self._require(
                    restored_controls
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    }
                    and tab.batch_progress_var.get()
                    == "100% - 25/25 - Completed"
                    and int(tab.batch_progress_value.get()) == 100,
                    "ERalpha batch controls were not restored after success",
                )
                self._require(
                    self.erta_batch_output_path is not None,
                    "ERTA primary workbook is unavailable for style comparison",
                )
                formatting_parity = compare_plain_primary_workbook_formatting(
                    self.erta_batch_output_path,
                    output,
                )
                recognition = self._record_batch_terminal(
                    "eralpha",
                    "success",
                    dialog=completion_dialog,
                    output=output,
                    output_exists=output.is_file(),
                    total_count=len(prediction_rows),
                    available_count=sum(
                        row["Result_Status"] == "Predicted"
                        for row in diagnostics_rows
                    ),
                    unavailable_count=not_predicted_count,
                    workbook_unavailable_count=sum(
                        row["Result_Status"] != "Predicted"
                        for row in diagnostics_rows
                    ),
                    fresh_output_count=1,
                    screenshot=True,
                )
                self._record(
                    "eralpha_shared_test_batch_success_dialog_and_predictions",
                    input=str(self.erba_batch_input_path),
                    output=str(output),
                    sheets=sheet_names,
                    prediction_headers=prediction_headers,
                    diagnostics_headers=diagnostics_headers,
                    diagnostics_status_codes=[
                        row["Status_Code"] for row in diagnostics_rows
                    ],
                    smiles_provenance=[
                        row["SMILES_Provenance"]
                        for row in diagnostics_rows
                    ],
                    preserved_input_rows=[list(row) for row in input_rows],
                    input_header="CARSRN",
                    metadata=metadata,
                    graph_directory=str(graph_directory),
                    graph_files=sorted(graph_files),
                    binding_count=binding_count,
                    non_binding_count=non_binding_count,
                    not_predicted_count=not_predicted_count,
                    batch_summary=summary,
                    progress=tab.batch_progress_var.get(),
                    completion_dialog=completion_dialog,
                    controls_during_run=self.erba_success_locked_controls,
                    controls=restored_controls,
                    plain_primary_formatting=formatting_parity,
                    deterministic_smiles_substitution=self.MOCK_SMILES,
                    source_bytes_preserved=(
                        _sha256_path(self.qa_shared_example_path)
                        == self.shared_bundle_sha256
                    ),
                    recognition=recognition,
                )
                self.stage = 132
            if self.stage == 132:
                mixed_directory = (
                    self.run_root / "mixed-batch" / "eralpha"
                )
                mixed_directory.mkdir(parents=True, exist_ok=False)
                mixed_input = (
                    mixed_directory / "erba-alpha-batch-input.xlsx"
                )
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Original_Input"
                sheet.append(["Row_ID", "SMILES", "Analyst_Note"])
                sheet.append(
                    ["native-valid", "Oc1ccccc1", "preserve-valid"]
                )
                sheet.append(
                    ["native-invalid", "[*]CC", "preserve-invalid"]
                )
                workbook.save(mixed_input)
                workbook.close()
                self.erba_mixed_input_path = mixed_input.resolve(
                    strict=True
                )
                self.erba_mixed_outputs_before = {
                    path.resolve()
                    for path in mixed_directory.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                }
                self.erba_mixed_graph_dirs_before = {
                    path.resolve()
                    for path in mixed_directory.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                }
                tab.batch_input_var.set(str(self.erba_mixed_input_path))
                tab.batch_input_display_var.set(
                    self.erba_mixed_input_path.name
                )
                tab.batch_destination_var.set(str(mixed_directory))
                self.erba_mixed_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "eralpha",
                    "partial_unavailable",
                )
                tab.batch_predict_clicked()
                self.erba_mixed_locked_controls = (
                    self._batch_control_states(tab)
                )
                self._require(
                    tab._active_batch_request_id is not None
                    and self.erba_mixed_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERalpha mixed batch did not start with locked controls",
                )
                self._record_batch_new_run("eralpha")
                self.stage = 133
                self._reschedule()
                return
            if self.stage == 133:
                if (
                    tab._active_batch_request_id is not None
                    or not self._wait_dialog(
                        self.erba_mixed_dialog_count
                    )
                ):
                    self._reschedule()
                    return
                mixed_parent = self.erba_mixed_input_path.parent
                mixed_outputs = {
                    path.resolve()
                    for path in mixed_parent.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                } - self.erba_mixed_outputs_before
                self._require(
                    len(mixed_outputs) == 1,
                    "expected one new ERalpha mixed-result workbook, found "
                    f"{len(mixed_outputs)}: {tab.batch_status_var.get()}",
                )
                mixed_output = mixed_outputs.pop()
                self._require(
                    mixed_output.parent == mixed_parent
                    and tab.batch_destination_var.get()
                    == str(mixed_parent),
                    "ERalpha mixed workbook destination is not its input parent",
                )
                mixed_graph_dirs = {
                    path.resolve()
                    for path in mixed_parent.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                } - self.erba_mixed_graph_dirs_before
                self._require(
                    len(mixed_graph_dirs) == 1,
                    "expected one new ERalpha mixed graph directory, found "
                    f"{len(mixed_graph_dirs)}",
                )
                mixed_graph_directory = mixed_graph_dirs.pop()
                mixed_graph_files = {
                    path.name
                    for path in mixed_graph_directory.iterdir()
                    if path.is_file()
                }
                self._require(
                    mixed_graph_directory.parent == mixed_parent
                    and mixed_graph_files
                    == self.ERBA_BATCH_GRAPH_FILES,
                    "ERalpha mixed-batch graphs are incomplete or misplaced",
                )
                workbook = load_workbook(
                    mixed_output,
                    read_only=True,
                    data_only=True,
                )
                try:
                    sheet_names = workbook.sheetnames
                    self._require(
                        sheet_names == list(ERALPHA_WORKBOOK_SHEET_ORDER)
                        and workbook.active.title == "Predictions",
                        "ERalpha mixed workbook is not Predictions-first and active",
                    )
                    prediction_sheet = workbook["Predictions"]
                    prediction_headers = [
                        cell.value for cell in prediction_sheet[1]
                    ]
                    self._require_eralpha_primary_headers(
                        prediction_headers,
                        ("Row_ID", "Analyst_Note"),
                        "ERalpha mixed workbook",
                    )
                    prediction_rows = [
                        dict(zip(prediction_headers, row))
                        for row in prediction_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                    diagnostics_sheet = workbook["Diagnostics"]
                    diagnostics_headers = [
                        cell.value for cell in diagnostics_sheet[1]
                    ]
                    diagnostics_rows = [
                        dict(zip(diagnostics_headers, row))
                        for row in diagnostics_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                    input_sheet = workbook["Input"]
                    input_headers = [
                        cell.value for cell in input_sheet[1]
                    ]
                    input_rows = list(
                        input_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    )
                    guide_rows = list(
                        workbook["Guide"].iter_rows(values_only=True)
                    )
                    metadata_sheet = workbook["Metadata"]
                    metadata_headers = [
                        cell.value for cell in metadata_sheet[1]
                    ]
                    metadata_rows = list(
                        metadata_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    )
                    self._require(
                        len(metadata_rows) == 1,
                        "ERalpha mixed Metadata must contain one "
                        "execution-level row",
                    )
                    metadata = dict(
                        zip(metadata_headers, metadata_rows[0])
                    )
                finally:
                    workbook.close()
                self._require(
                    len(prediction_rows) == 2
                    and prediction_rows[0]["Analyst_Note"]
                    == "preserve-valid"
                    and prediction_rows[1]["Analyst_Note"]
                    == "preserve-invalid",
                    "ERalpha mixed Predictions lost passthrough fields or order",
                )
                self._require_eralpha_diagnostics(
                    diagnostics_headers,
                    diagnostics_rows,
                    prediction_rows,
                    "ERalpha mixed workbook",
                )
                self._require(
                    self._eralpha_prediction_is_consistent(
                        prediction_rows[0]
                    )
                    and prediction_rows[0]["Mol_valid"] is True
                    and prediction_rows[0]["AD"]
                    in {"In-domain", "Out-of-domain"}
                    and prediction_rows[0]["PubChem_CID"] is None
                    and prediction_rows[0]["PubChem_status"] is None
                    and prediction_rows[1]["SMILES"] == "[*]CC"
                    and prediction_rows[1]["Canonical_SMILES"] is None
                    and prediction_rows[1]["Mol_valid"] is True
                    and self._eralpha_prediction_payload_is_blank(
                        prediction_rows[1]
                    )
                    and prediction_rows[1]["PubChem_CID"] is None
                    and prediction_rows[1]["PubChem_status"] is None,
                    "ERalpha mixed predicted/not-predicted row mapping is incorrect",
                )
                self._require(
                    diagnostics_rows[0]["Row_ID"] == "native-valid"
                    and diagnostics_rows[0]["Status_Code"] == "ok"
                    and diagnostics_rows[0]["SMILES_Provenance"]
                    == "direct_input"
                    and diagnostics_rows[0]["Result_Status"]
                    == "Predicted"
                    and diagnostics_rows[1]["Row_ID"]
                    == "native-invalid"
                    and diagnostics_rows[1]["Status_Code"]
                    == "wildcard_smiles"
                    and diagnostics_rows[1]["Status_Message"]
                    and diagnostics_rows[1]["SMILES_Provenance"]
                    == "direct_input"
                    and diagnostics_rows[1]["Result_Status"]
                    == "Not predicted"
                    and diagnostics_rows[1]["Reason_Category"]
                    == "Invalid or ambiguous SMILES"
                    and diagnostics_rows[1]["Reason_Description"]
                    and diagnostics_rows[1]["Recommended_Action"],
                    "ERalpha mixed Diagnostics lost the invalid-row reason "
                    "or stable technical status",
                )
                self._require(
                    input_headers
                    == ["Row_ID", "SMILES", "Analyst_Note"]
                    and input_rows
                    == [
                        (
                            "native-valid",
                            "Oc1ccccc1",
                            "preserve-valid",
                        ),
                        (
                            "native-invalid",
                            "[*]CC",
                            "preserve-invalid",
                        ),
                    ],
                    "ERalpha mixed Input sheet did not preserve exact source rows",
                )
                self._require(
                    any(
                        row[0] == "Not predicted rows"
                        and row[1] == "1"
                        for row in guide_rows
                    )
                    and any(
                        row[0]
                        == "Reason: Invalid or ambiguous SMILES"
                        and row[1] == "1"
                        for row in guide_rows
                    ),
                    "ERalpha mixed Guide does not explain the not-predicted row",
                )
                expected_spec = tab.catalog[
                    (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
                ]
                hex_digits = set("0123456789abcdef")
                self._require(
                    metadata_headers
                    == list(ERBA_CLASSIFICATION_METADATA_COLUMNS)
                    and metadata.get("Excel_Contract_ID")
                    == ERBA_BINDING_CLASSIFICATION_EXCEL_CONTRACT_ID
                    and metadata.get("Workflow") == "erba"
                    and metadata.get("Task") == "classification"
                    and metadata.get("Subtype") == "er_alpha"
                    and metadata.get("Decision_rule")
                    == "binding_probability>=0.5"
                    and metadata.get("Model_ID")
                    == expected_spec.model_id
                    and metadata.get("Model_SHA256")
                    == expected_spec.sha256
                    and metadata.get("Preprocessing_Policy_ID")
                    == CLASSIFICATION_PARENT_POLICY_ID
                    and metadata.get("Performance_Evidence_Scope")
                    == "internal_historically_exposed"
                    and bool(metadata.get("Evidence_Caveat"))
                    and all(
                        isinstance(metadata.get(column), str)
                        and len(metadata[column]) == 64
                        and set(metadata[column]) <= hex_digits
                        for column in self.ERBA_PROVENANCE_COLUMNS
                    ),
                    "ERalpha mixed Metadata lost catalog provenance",
                )
                binding_count = sum(
                    row["Prediction_label"] == "Binding"
                    for row in prediction_rows
                )
                non_binding_count = sum(
                    row["Prediction_label"] == "Non-binding"
                    for row in prediction_rows
                )
                not_predicted_count = sum(
                    row["Result_Status"] != "Predicted"
                    for row in diagnostics_rows
                )
                mixed_summary = tab.batch_result.get(
                    "1.0", "end"
                ).strip()
                self._require(
                    tab.batch_status_var.get()
                    == f"Batch prediction completed: {mixed_output}"
                    and f"Total rows: {len(prediction_rows)}"
                    in mixed_summary
                    and f"Binding: {binding_count}" in mixed_summary
                    and f"Non-binding: {non_binding_count}"
                    in mixed_summary
                    and f"{BATCH_UNAVAILABLE_LABEL}: 1"
                    in mixed_summary
                    and f"Output workbook: {mixed_output}"
                    in mixed_summary
                    and f"Graph files: {len(mixed_graph_files)}"
                    in mixed_summary
                    and f"Graph directory: {mixed_graph_directory}"
                    in mixed_summary,
                    "ERalpha mixed completion summary is incomplete",
                )
                completion_dialog = self._terminal_dialog(
                    self.erba_mixed_dialog_count,
                    kind="info",
                    title="Batch prediction done",
                    context="ERalpha mixed batch partial success",
                )
                restored_controls = self._batch_control_states(tab)
                self._require(
                    f"{BATCH_UNAVAILABLE_LABEL}: 1"
                    in completion_dialog["message"]
                    and restored_controls
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    }
                    and tab.batch_progress_var.get()
                    == "100% - 2/2 - Completed"
                    and int(tab.batch_progress_value.get()) == 100
                    and completion_dialog[
                        "batch_controls_at_dialog"
                    ]["eralpha"]
                    == restored_controls,
                    "ERalpha mixed info dialog or control-restoration contract drift",
                )
                mixed_recognition = self._record_batch_terminal(
                    "eralpha",
                    "partial_unavailable",
                    dialog=completion_dialog,
                    output=mixed_output,
                    output_exists=mixed_output.is_file(),
                    total_count=len(prediction_rows),
                    available_count=(
                        len(prediction_rows) - not_predicted_count
                    ),
                    unavailable_count=not_predicted_count,
                    workbook_unavailable_count=sum(
                        row["Result_Status"] != "Predicted"
                        for row in diagnostics_rows
                    ),
                    fresh_output_count=1,
                )
                self._record(
                    "erba_eralpha_batch_workbook_ad_graphs_and_summary",
                    input=str(self.erba_mixed_input_path),
                    output=str(mixed_output),
                    sheets=sheet_names,
                    prediction_headers=prediction_headers,
                    diagnostics_headers=diagnostics_headers,
                    preserved_input_rows=[
                        list(row) for row in input_rows
                    ],
                    metadata=metadata,
                    graph_directory=str(mixed_graph_directory),
                    graph_files=sorted(mixed_graph_files),
                    binding_count=binding_count,
                    non_binding_count=non_binding_count,
                    not_predicted_count=not_predicted_count,
                    invalid_reason_category=diagnostics_rows[1][
                        "Reason_Category"
                    ],
                    invalid_status_code=diagnostics_rows[1][
                        "Status_Code"
                    ],
                    batch_summary=mixed_summary,
                    completion_dialog=completion_dialog,
                    progress=tab.batch_progress_var.get(),
                    controls_during_run=self.erba_mixed_locked_controls,
                    controls=restored_controls,
                    recognition=mixed_recognition,
                )
                self.erba_mixed_check = self.transcript["checks"][-1]
                self.stage = 134
            if self.stage == 134:
                unavailable_directory = (
                    self.run_root
                    / "all-unavailable-batch"
                    / "eralpha"
                )
                unavailable_directory.mkdir(parents=True, exist_ok=False)
                unavailable_input = (
                    unavailable_directory / "all-unavailable-input.xlsx"
                )
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["Row_ID", "CAS"])
                for index, cas in enumerate(
                    self.MOCK_UNAVAILABLE_CAS_VALUES,
                    1,
                ):
                    sheet.append([f"unavailable-{index}", cas])
                workbook.save(unavailable_input)
                workbook.close()
                self.erba_unavailable_input_path = (
                    unavailable_input.resolve(strict=True)
                )
                self.erba_unavailable_outputs_before = {
                    path.resolve()
                    for path in unavailable_directory.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                }
                self.erba_unavailable_graph_dirs_before = {
                    path.resolve()
                    for path in unavailable_directory.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                }
                tab.batch_input_var.set(
                    str(self.erba_unavailable_input_path)
                )
                tab.batch_input_display_var.set(
                    self.erba_unavailable_input_path.name
                )
                tab.batch_destination_var.set(
                    str(unavailable_directory)
                )
                self.erba_unavailable_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "eralpha",
                    "all_unavailable",
                )
                tab.batch_predict_clicked()
                self.erba_unavailable_locked_controls = (
                    self._batch_control_states(tab)
                )
                self._require(
                    tab._active_batch_request_id is not None
                    and self.erba_unavailable_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERalpha all-unavailable batch did not start with "
                    "locked controls",
                )
                self._record_batch_new_run("eralpha")
                self.stage = 135
                self._reschedule()
                return
            if self.stage == 135:
                if (
                    tab._active_batch_request_id is not None
                    or not self._wait_dialog(
                        self.erba_unavailable_dialog_count
                    )
                ):
                    self._reschedule()
                    return
                unavailable_parent = (
                    self.erba_unavailable_input_path.parent
                )
                unavailable_outputs = {
                    path.resolve()
                    for path in unavailable_parent.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                } - self.erba_unavailable_outputs_before
                self._require(
                    len(unavailable_outputs) == 1,
                    "expected one fresh all-unavailable ERalpha workbook, "
                    f"found {len(unavailable_outputs)}",
                )
                unavailable_output = unavailable_outputs.pop()
                new_graph_directories = {
                    path.resolve()
                    for path in unavailable_parent.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                } - self.erba_unavailable_graph_dirs_before
                self._require(
                    not new_graph_directories,
                    "ERalpha all-unavailable batch created misleading "
                    "graph artifacts",
                )
                workbook = load_workbook(
                    unavailable_output,
                    read_only=True,
                    data_only=True,
                )
                try:
                    unavailable_sheet_names = workbook.sheetnames
                    self._require(
                        unavailable_sheet_names
                        == list(ERALPHA_WORKBOOK_SHEET_ORDER)
                        and workbook.active.title == "Predictions",
                        "ERalpha all-unavailable workbook sheet order drift",
                    )
                    prediction_sheet = workbook["Predictions"]
                    unavailable_headers = [
                        cell.value for cell in prediction_sheet[1]
                    ]
                    self._require_eralpha_primary_headers(
                        unavailable_headers,
                        ("Row_ID",),
                        "ERalpha all-unavailable workbook",
                    )
                    unavailable_rows = [
                        dict(zip(unavailable_headers, row))
                        for row in prediction_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                    diagnostics_sheet = workbook["Diagnostics"]
                    unavailable_diagnostics_headers = [
                        cell.value for cell in diagnostics_sheet[1]
                    ]
                    unavailable_diagnostics_rows = [
                        dict(
                            zip(
                                unavailable_diagnostics_headers,
                                row,
                            )
                        )
                        for row in diagnostics_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    ]
                    unavailable_input_sheet = workbook["Input"]
                    unavailable_input_headers = [
                        cell.value for cell in unavailable_input_sheet[1]
                    ]
                    unavailable_input_rows = list(
                        unavailable_input_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    )
                    unavailable_metadata_sheet = workbook["Metadata"]
                    unavailable_metadata_headers = [
                        cell.value
                        for cell in unavailable_metadata_sheet[1]
                    ]
                    unavailable_metadata_rows = list(
                        unavailable_metadata_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        )
                    )
                    self._require(
                        len(unavailable_metadata_rows) == 1,
                        "ERalpha all-unavailable Metadata must contain "
                        "one execution-level row",
                    )
                    unavailable_metadata = dict(
                        zip(
                            unavailable_metadata_headers,
                            unavailable_metadata_rows[0],
                        )
                    )
                finally:
                    workbook.close()
                unavailable_count = len(
                    self.MOCK_UNAVAILABLE_CAS_VALUES
                )
                self._require_eralpha_diagnostics(
                    unavailable_diagnostics_headers,
                    unavailable_diagnostics_rows,
                    unavailable_rows,
                    "ERalpha all-unavailable workbook",
                )
                self._require(
                    len(unavailable_rows) == unavailable_count
                    and [
                        str(row["CAS"])
                        for row in unavailable_rows
                    ]
                    == list(self.MOCK_UNAVAILABLE_CAS_VALUES)
                    and unavailable_input_headers == ["Row_ID", "CAS"]
                    and unavailable_input_rows
                    == [
                        (f"unavailable-{index}", cas)
                        for index, cas in enumerate(
                            self.MOCK_UNAVAILABLE_CAS_VALUES,
                            1,
                        )
                    ]
                    and all(
                        row["SMILES"] is None
                        and row["Canonical_SMILES"] is None
                        and row["Mol_valid"] is False
                        and self._eralpha_prediction_payload_is_blank(
                            row
                        )
                        and row["PubChem_CID"] is None
                        and row["PubChem_status"] is None
                        for row in unavailable_rows
                    )
                    and all(
                        diagnostic["Row_ID"]
                        == f"unavailable-{index}"
                        and diagnostic["Status_Code"]
                        == "prediction_failed"
                        and diagnostic["Status_Message"]
                        and diagnostic["SMILES_Provenance"]
                        == "pubchem_lookup_failed"
                        and diagnostic["Result_Status"]
                        == "Not predicted"
                        and diagnostic["Reason_Category"]
                        == "PubChem lookup unavailable"
                        and diagnostic["Reason_Description"]
                        and diagnostic["Recommended_Action"]
                        for index, diagnostic in enumerate(
                            unavailable_diagnostics_rows,
                            1,
                        )
                    ),
                    "ERalpha all-unavailable result did not retain its "
                    "reasoned blank-prediction rows",
                )
                expected_spec = tab.catalog[
                    (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
                ]
                self._require(
                    unavailable_metadata_headers
                    == list(ERBA_CLASSIFICATION_METADATA_COLUMNS)
                    and unavailable_metadata.get("Excel_Contract_ID")
                    == ERBA_BINDING_CLASSIFICATION_EXCEL_CONTRACT_ID
                    and unavailable_metadata.get("Model_ID")
                    == expected_spec.model_id
                    and unavailable_metadata.get("Model_SHA256")
                    == expected_spec.sha256,
                    "ERalpha all-unavailable Metadata lost the v3 "
                    "contract or model provenance",
                )
                unavailable_summary = self._batch_result_value(tab)
                unavailable_dialog = self._terminal_dialog(
                    self.erba_unavailable_dialog_count,
                    kind="info",
                    title="Batch prediction done",
                    context="ERalpha all-unavailable batch publication",
                )
                restored_controls = self._batch_control_states(tab)
                self._require(
                    f"{BATCH_UNAVAILABLE_LABEL}: {unavailable_count}"
                    in unavailable_summary
                    and f"{BATCH_UNAVAILABLE_LABEL}: "
                    f"{unavailable_count}"
                    in unavailable_dialog["message"]
                    and tab.batch_progress_var.get()
                    == (
                        f"100% - {unavailable_count}/"
                        f"{unavailable_count} - Completed"
                    )
                    and restored_controls
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    },
                    "ERalpha all-unavailable publication feedback drift",
                )
                unavailable_recognition = self._record_batch_terminal(
                    "eralpha",
                    "all_unavailable",
                    dialog=unavailable_dialog,
                    output=unavailable_output,
                    output_exists=unavailable_output.is_file(),
                    total_count=unavailable_count,
                    available_count=0,
                    unavailable_count=unavailable_count,
                    workbook_unavailable_count=sum(
                        row["Result_Status"] != "Predicted"
                        for row in unavailable_diagnostics_rows
                    ),
                    fresh_output_count=1,
                )
                self.erba_mixed_check["all_unavailable"] = {
                    "input": str(self.erba_unavailable_input_path),
                    "output": str(unavailable_output),
                    "row_count": unavailable_count,
                    "unavailable_count": unavailable_count,
                    "rows": unavailable_rows,
                    "diagnostics": unavailable_diagnostics_rows,
                    "diagnostics_headers": (
                        unavailable_diagnostics_headers
                    ),
                    "metadata": unavailable_metadata,
                    "preserved_input_rows": [
                        list(row)
                        for row in unavailable_input_rows
                    ],
                    "fresh_graph_directories": [],
                    "batch_summary": unavailable_summary,
                    "completion_dialog": unavailable_dialog,
                    "progress": tab.batch_progress_var.get(),
                    "controls_during_run": (
                        self.erba_unavailable_locked_controls
                    ),
                    "controls": restored_controls,
                    "recognition": unavailable_recognition,
                }
                self.stage = 130
            if self.stage == 130:
                failure_input = self._make_corrupt_batch_input("eralpha")
                self.erba_failure_input_path = failure_input
                self.erba_failure_outputs_before = {
                    path.resolve()
                    for path in failure_input.parent.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                }
                self.erba_failure_graph_dirs_before = {
                    path.resolve()
                    for path in failure_input.parent.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                }
                tab.batch_input_var.set(str(failure_input))
                tab.batch_input_display_var.set(failure_input.name)
                tab.batch_destination_var.set(str(failure_input.parent))
                self.erba_failure_dialog_count = len(self.dialogs)
                self._begin_batch_transition_observer(
                    "eralpha",
                    "no_output_failure",
                )
                tab.batch_predict_clicked()
                self.erba_failure_locked_controls = (
                    self._batch_control_states(tab)
                )
                self._require(
                    tab._active_batch_request_id is not None
                    and self.erba_failure_locked_controls
                    == {
                        "input": "disabled",
                        "template": "disabled",
                        "run": "disabled",
                    },
                    "ERalpha deterministic failure did not start with locked controls",
                )
                self._record_batch_new_run("eralpha")
                self.stage = 131
                self._reschedule()
                return
            if self.stage == 131:
                if (
                    tab._active_batch_request_id is not None
                    or not self._wait_dialog(
                        self.erba_failure_dialog_count
                    )
                ):
                    self._reschedule()
                    return
                failure_dialog = self._terminal_dialog(
                    self.erba_failure_dialog_count,
                    kind="error",
                    title="Batch prediction failed",
                    context="ERalpha deterministic batch failure",
                )
                failure_summary = tab.batch_result.get(
                    "1.0", "end"
                ).strip()
                new_outputs = {
                    path.resolve()
                    for path in self.erba_failure_input_path.parent.glob(
                        "ERBA_classification_er_alpha_results*.xlsx"
                    )
                } - self.erba_failure_outputs_before
                new_graph_directories = {
                    path.resolve()
                    for path in self.erba_failure_input_path.parent.glob(
                        "ERBA_classification_er_alpha_results*_graphs*"
                    )
                    if path.is_dir()
                } - self.erba_failure_graph_dirs_before
                restored_controls = self._batch_control_states(tab)
                self._require(
                    restored_controls
                    == {
                        "input": "normal",
                        "template": "normal",
                        "run": "normal",
                    }
                    and tab.batch_progress_var.get()
                    == "100% - 0/0 - Failed"
                    and int(tab.batch_progress_value.get()) == 100
                    and tab.batch_status_var.get().startswith(
                        "Batch prediction failed:"
                    )
                    and failure_summary.startswith(
                        "Batch prediction failed."
                    )
                    and "ERalpha batch prediction completed."
                    not in failure_summary
                    and "Output workbook:" not in failure_summary
                    and not new_outputs
                    and not new_graph_directories
                    and failure_dialog[
                        "batch_controls_at_dialog"
                    ]["eralpha"]
                    == restored_controls,
                    "ERalpha failure exposed false success or did not restore controls",
                )
                failure_recognition = self._record_batch_terminal(
                    "eralpha",
                    "no_output_failure",
                    dialog=failure_dialog,
                    output=None,
                    output_exists=False,
                    fresh_output_count=len(new_outputs),
                )
                self._record(
                    "eralpha_batch_failure_dialog_and_control_recovery",
                    input=str(self.erba_failure_input_path),
                    fixture_kind="intentionally corrupt xlsx bytes",
                    dialog=failure_dialog,
                    controls_during_run=self.erba_failure_locked_controls,
                    controls=restored_controls,
                    progress=tab.batch_progress_var.get(),
                    status=tab.batch_status_var.get(),
                    result=failure_summary,
                    fresh_outputs=[],
                    fresh_graph_directories=[],
                    success_dialog_emitted=False,
                    recognition=failure_recognition,
                )
                self.stage = 14
            if self.stage == 14:
                self._require(
                    app.eralpha_tab._selected_route("single")
                    == (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA),
                    "ERalpha page is not fixed to classification",
                )
                alpha_model = app.eralpha_tab.catalog[
                    (ERBATask.CLASSIFICATION, ERBASubtype.ER_ALPHA)
                ].model_id
                self._require(
                    not hasattr(app, "erbeta_tab")
                    and app.notebook.tabs()
                    and "ERbeta"
                    not in [
                        app.notebook.tab(tab_id, "text")
                        for tab_id in app.notebook.tabs()
                    ],
                    "ERbeta remains exposed in the FDA application",
                )
                self._record(
                    "fixed_eralpha_classification_tab",
                    eralpha_model=alpha_model,
                    erbeta_present=False,
                )
                app.notebook.select(app.erta_tab)
                self._require(app.prediction_summary_var.get() != self.erta_summary or app.active_probability_var.get() != self.erta_probability, "ERTA state was not independently updated after CAS journey")
                self._require(app.last_batch_result is not None and app.notebook.tab(app.notebook.select(), "text") == "ERTA", "ERTA batch state lost after ERBA journey")
                self._record("erta_erba_erta_state_isolation", erta_summary_after_return=app.prediction_summary_var.get(), erta_probability_after_return=app.active_probability_var.get(), erba_result_retained=bool(self._body(app.eralpha_tab)), erta_batch_rows=len(app.last_batch_result))
                self._require(
                    self.shared_default_path is not None
                    and self.shared_default_path.is_file()
                    and _sha256_path(self.shared_default_path)
                    == self.shared_default_sha256,
                    "native QA changed the resolved shared default workbook",
                )
                self._require(
                    self.shared_bundle_path is not None
                    and self.shared_bundle_path.is_file()
                    and _sha256_path(self.shared_bundle_path)
                    == self.shared_bundle_sha256
                    and self.qa_shared_example_path is not None
                    and _sha256_path(self.qa_shared_example_path)
                    == self.shared_bundle_sha256
                    and self.erta_batch_input_path is not None
                    and _sha256_path(self.erta_batch_input_path)
                    == self.shared_bundle_sha256
                    and self.erba_batch_input_path is not None
                    and _sha256_path(self.erba_batch_input_path)
                    == self.shared_bundle_sha256,
                    "bundled test.xlsx or a QA input copy changed during prediction",
                )
                self._record(
                    "read_only_install_and_writable_state",
                    resource_root=str(app.project_root),
                    output_root=str(app.output_root),
                    state_root=str(app.state_root),
                    qa_run_root=str(self.run_root),
                    shared_example={
                        "resolved_default": str(self.shared_default_path),
                        "resolved_default_sha256": self.shared_default_sha256,
                        "resolved_default_bytes_unchanged": True,
                        "bundled_source": str(self.shared_bundle_path),
                        "bundled_sha256": self.shared_bundle_sha256,
                        "bundled_bytes_unchanged": True,
                        "fresh_first_run_copy": str(
                            self.qa_shared_example_path
                        ),
                        "erta_batch_copy": str(
                            self.erta_batch_input_path
                        ),
                        "eralpha_batch_copy": str(
                            self.erba_batch_input_path
                        ),
                        "all_qa_copies_byte_match_bundle": True,
                    },
                )
                self._require(
                    self.batch_recognition_widget_contract is not None,
                    "batch recognition widget contract was not captured",
                )
                self.transcript["batch_recognition_gate"] = (
                    evaluate_batch_recognition_contract(
                        self.batch_recognition_records,
                        self.batch_recognition_widget_contract,
                    )
                )
                check_names = [
                    receipt["check"]
                    for receipt in self.transcript["checks"]
                ]
                required_checks = (
                    self.BASELINE_CHECKS
                    | self.SHARED_EXAMPLE_BATCH_CHECKS
                    | self.BATCH_FEEDBACK_CHECKS
                )
                self._require(
                    required_checks <= set(check_names)
                    and len(self.BASELINE_CHECKS) == 22
                    and len(BATCH_RECOGNITION_CRITERIA) == 9
                    and len(check_names)
                    == 26 + int(bool(self.external_driver_gate)),
                    "native QA check inventory does not retain the 22 baseline "
                    "checks, shared-example/batch-failure contracts, and nine "
                    "recognition criteria",
                )
                self.transcript["check_contract"] = {
                    "baseline_check_count": 22,
                    "baseline_checks": sorted(self.BASELINE_CHECKS),
                    "added_shared_example_batch_check_count": 2,
                    "added_shared_example_batch_checks": sorted(
                        self.SHARED_EXAMPLE_BATCH_CHECKS
                    ),
                    "added_batch_feedback_check_count": 2,
                    "added_batch_feedback_checks": sorted(
                        self.BATCH_FEEDBACK_CHECKS
                    ),
                    "external_driver_check_present": bool(
                        self.external_driver_gate
                    ),
                    "batch_recognition_gate": {
                        "gate_id": BATCH_RECOGNITION_GATE_ID,
                        "criteria_count": len(
                            BATCH_RECOGNITION_CRITERIA
                        ),
                        "criteria_ids": list(
                            BATCH_RECOGNITION_CRITERIA
                        ),
                        "counted_as_additional_check": False,
                    },
                    "total_check_count": len(check_names),
                }
                app.deiconify()
                app.state("normal")
                app.update()
                self._finish(True)
        except Exception:
            self._finish(False, traceback.format_exc())
        finally:
            self._tick_active = False
